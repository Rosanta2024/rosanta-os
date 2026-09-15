import openpyxl, re, pickle, sys, datetime
from collections import defaultdict, Counter
from openpyxl.styles import Font, PatternFill, Alignment
SP, ESP, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
ph,P,fh,Fr = pickle.load(open(SP+"/vivo.pkl","rb"))
wb = openpyxl.load_workbook(ESP, data_only=True)

def clv(v):
    x=str(v if v is not None else "").strip()
    return x[:-2] if x.endswith(".0") else x
def up(s): return re.sub(r"\s+"," ",str(s or "").strip().upper())
def num(v):
    try: return float(v or 0)
    except: return 0.0
def fch(v): return v if isinstance(v,datetime.datetime) else None

ENT={}
def ent(key, nombre, tipo, conf):
    if key not in ENT:
        ENT[key]=dict(key=key,nombre=nombre,tipo=tipo,conf=conf,nit=set(),legales=Counter(),est=Counter(),
            cat_alias=[],cuentas=set(),comercios=Counter(),glosas=Counter(),obs=Counter(),n=0,q=0.0,usd=0.0,
            f1=None,f2=None,fuentes=set(),flags=[])
    return ENT[key]
def mov(e,fuente,cat,q=0.0,usd=0.0,f=None):
    e["n"]+=1; e["q"]+=q; e["usd"]+=usd; e["fuentes"].add(fuente); e["obs"][fuente+":"+(cat or "(vacia)")]+=1
    if f:
        if not e["f1"] or f<e["f1"]: e["f1"]=f
        if not e["f2"] or f>e["f2"]: e["f2"]=f

REC=defaultdict(lambda: defaultdict(lambda:[0,0.0,0.0]))   # REC[fuente][destino]=[filas,Q,USD]
def rec(fuente,destino,q=0.0,usd=0.0):
    r=REC[fuente][destino]; r[0]+=1; r[1]+=q; r[2]+=usd
EXCL=defaultdict(lambda:[0,0.0,0.0]); EFEC=defaultdict(lambda:[0,0.0,0.0]); PORA=[]
MOTIVO={"NOMINA":"equipo / planilla","PROPINAS_AL_EQUIPO":"equipo / planilla","IGSS":"equipo / planilla",
        "DEVOLUCION_INVERSION":"socio / devolucion de inversion","PAGO_TARJETA_CREDITO":"movimiento interno (pago de tarjeta)",
        "PAGO_TARJETA":"movimiento interno (pago de tarjeta)","IMPUESTOS":"impuestos","TRIBUTO":"impuestos",
        "CARGO_FRAUDULENTO":"cargo fraudulento"}
EXC=set(MOTIVO)

RAW=defaultdict(lambda:[0,0.0,0.0])
_f=wb["01_FEL_Maestro"]
for r in range(5,_f.max_row+1):
    if _f.cell(r,5).value is None and _f.cell(r,6).value is None: continue
    RAW["FEL"][0]+=1; RAW["FEL"][1]+=num(_f.cell(r,10).value)
for h,cq in (("04_Banco_BAC",5),("03_Banco_Industrial",4)):
    _s=wb[h]; k="BAC" if "BAC" in h else "BI"
    for r in range(5,_s.max_row+1):
        v=num(_s.cell(r,cq).value)
        if v>0: RAW[k][0]+=1; RAW[k][1]+=v
_t=wb["05_Tarjeta_Credito_BAC"]
for r in range(5,_t.max_row+1):
    q=num(_t.cell(r,3).value); u=num(_t.cell(r,4).value)
    if q>0 or u>0: RAW["TARJETA"][0]+=1; RAW["TARJETA"][1]+=q; RAW["TARJETA"][2]+=u

# ---------------- FEL: una entidad por NIT
fel=wb["01_FEL_Maestro"]
for r in range(5,fel.max_row+1):
    if fel.cell(r,5).value is None and fel.cell(r,6).value is None: continue
    nit=clv(fel.cell(r,5).value); q=num(fel.cell(r,10).value); cat=str(fel.cell(r,14).value or "").strip()
    e=ent(("NIT",nit),None,None,"ALTA")
    e["nit"].add(nit); e["legales"][str(fel.cell(r,6).value).strip()]+=1; e["est"][str(fel.cell(r,7).value).strip()]+=1
    mov(e,"FEL",cat,q=q,f=fch(fel.cell(r,1).value)); rec("FEL","maestro",q)
est2keys=defaultdict(set)
for k,e in ENT.items():
    e["nombre"]= list(e["est"])[0] if len(e["est"])==1 else e["legales"].most_common(1)[0][0]
    for s in e["est"]: est2keys[up(s)].add(k)

# ---------------- CATALOGO VIVO
catv={}
for p in P:
    n=str(p["Proveedor"]).strip()
    if n and up(n) not in catv: catv[up(n)]=p
for u,p in catv.items():
    info=(p["Proveedor"].strip(),p["Categoría_Normalizada"],p["Familia"],p["Es_Personal"])
    keys=est2keys.get(u)
    if keys:
        for k in keys:
            ENT[k]["cat_alias"].append(info)
            if len(keys)>1: ENT[k]["flags"].append("el nombre de catalogo '%s' coincide con %d NIT distintos: el VLOOKUP les pone la misma categoria" % (info[0],len(keys)))
    else:
        tipo="PERSONAL" if (p["Categoría_Normalizada"]=="PERSONAL" or up(p["Es_Personal"]) in ("SI","SÍ")) else "PROVEEDOR"
        e=ent(("CAT",u),info[0],tipo,"MEDIA"); e["cat_alias"].append(info); e["fuentes"].add("CATALOGO")
        e["flags"].append("esta en el catalogo pero no tiene factura FEL en 2026")
    rec("CATALOGO","nombres unicos")
for k,e in ENT.items():
    if k[0]!="NIT": continue
    catalogados={up(c[0]) for c in e["cat_alias"]}
    for s in e["est"]:
        if up(s) not in catalogados: e["flags"].append("establecimiento sin fila en catalogo (sus facturas nuevas entran POR_CLASIFICAR): %s" % s)
    fel_cats={o.split(":",1)[1] for o in e["obs"] if o.startswith("FEL:")}
    if len(fel_cats)>1: e["flags"].append("varias categorias en el FEL: %s" % ", ".join(sorted(fel_cats)))
    cc=[c[1] for c in e["cat_alias"]]; ep=[up(c[3]) for c in e["cat_alias"]]
    fel_major=Counter({o.split(":",1)[1]:v for o,v in e["obs"].items() if o.startswith("FEL:")}).most_common(1)[0][0]
    personal = (cc and all(c=="PERSONAL" for c in cc)) or (ep and all(x in ("SI","SÍ") for x in ep)) or (not cc and fel_major=="PERSONAL")
    e["tipo"]="PERSONAL" if personal else "PROVEEDOR"
    if "24545910" in e["nit"]: e["flags"].append("factura propia de Juanma, como pequeno contribuyente")

# ---------------- BAC
CUENTA_DOC={"902031855":"Andrea Moreno Gil (Producciones Talishte)","902409168":"IGYM, S.A.","902410067":"Doorways, S.A.",
            "903941979":"Felix Menelik Donis Yuman (Tavito)","974954208":"devolucion de inversion"}
bac=wb["04_Banco_BAC"]; cuentas=defaultdict(list)
for r in range(5,bac.max_row+1):
    q=num(bac.cell(r,5).value)
    if q<=0: continue
    d=up(bac.cell(r,4).value); cat=str(bac.cell(r,8).value or "").strip(); f=fch(bac.cell(r,1).value); ref=clv(bac.cell(r,2).value)
    m=re.search(r"TEF A\s*:\s*(\d+)",d)
    if m: cuentas[m.group(1)].append((q,cat,f,ref,d)); continue
    if cat in EXC or "PAGO SAT" in d:
        mot=MOTIVO.get(cat,"impuestos"); EXCL[("BAC",mot,cat)][0]+=1; EXCL[("BAC",mot,cat)][1]+=q; rec("BAC","excluido: "+mot,q)
    elif cat=="COMISIONES_BANCARIAS" or re.search(r"COMISION|CARGO POR SALDO|COBRO EMISI",d):
        e=ent(("BANCO","BAC"),"BANCO BAC (servicios bancarios)","PROVEEDOR","ALTA"); e["glosas"]["BAC: "+d]+=1; mov(e,"BAC",cat,q=q,f=f); rec("BAC","maestro",q)
    else:
        PORA.append(("BAC",f,ref,d,cat,q,0.0)); rec("BAC","pago sin nombre (hoja Pagos_sin_nombre)",q)
for cta,rows in cuentas.items():
    cats=Counter(x[1] for x in rows); tot=sum(x[0] for x in rows)
    if "DEVOLUCION_INVERSION" in cats or cta=="974954208":
        for q,cat,f,ref,d in rows: EXCL[("BAC","socio / devolucion de inversion","cuenta "+cta)][0]+=1; EXCL[("BAC","socio / devolucion de inversion","cuenta "+cta)][1]+=q; rec("BAC","excluido: socio / devolucion de inversion",q)
        continue
    if (cats["NOMINA"]+cats["PROPINAS_AL_EQUIPO"]) >= len(rows)/2:
        for q,cat,f,ref,d in rows: EXCL[("BAC","equipo / planilla","cuentas de personal")][0]+=1; EXCL[("BAC","equipo / planilla","cuentas de personal")][1]+=q; rec("BAC","excluido: equipo / planilla",q)
        continue
    nom="CUENTA BAC %s · %s" % (cta, CUENTA_DOC.get(cta,"titular sin identificar"))
    tipo="PERSONAL" if cats.most_common(1)[0][0]=="PERSONAL" else "PROVEEDOR"
    e=ent(("CUENTA",cta),nom,tipo,"ALTA"); e["cuentas"].add(cta)
    if cta in CUENTA_DOC: e["flags"].append("nombre del titular segun el historico de la banca en linea (pendiente p102)")
    for q,cat,f,ref,d in rows: mov(e,"BAC",cat,q=q,f=f); rec("BAC","maestro",q)

# ---------------- BANCO INDUSTRIAL: las glosas NO crean proveedor; se agrupan en patrones con sugerencia
MES_FULL="ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|AGSOTO|SEPTIEMBRE|SETIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE"
MES_ABR="ENE|FEB|MAR|ABR|JUN|JUL|AGO|SEP|SET|OCT|NOV|DIC"
STOPN={"Y","DE","DEL","LOS","LAS","LA","EL","GT","SA","S","A","SOCIEDAD","ANONIMA","ANÓNIMA","PAGO","COMPLEMENTO"}
def norm_letras(txt):
    t=re.sub(r"[^A-ZÁÉÍÓÚÑ]+"," ",up(txt))
    return "".join(w for w in t.split() if w not in STOPN)
def bi_patron(g):
    t=up(g)
    t=re.sub(r"^PORTAL-BI\s*","",t)
    t=re.sub(r"^S\d{1,2}(\s*Y\s*\d{1,2})?","",t)
    t=re.sub(r"(%s)(?=\d|\s|$)" % MES_FULL," ",t)
    t=re.sub(r"\b(%s)\b" % MES_ABR," ",t)
    t=re.sub(r"\b\d+\s*A\b"," ",t)
    t=re.sub(r"(^|\s)\d+A(?=[A-Z])"," ",t)
    t=re.sub(r"(VARIOS|VARIAS|VARIA)"," ",t)
    t=re.sub(r"\d+"," ",t)
    words=[w for w in re.sub(r"[^A-ZÁÉÍÓÚÑ]+"," ",t).split() if w not in STOPN]
    return " ".join(words), "".join(words)
CONCEPTO={"RENTA","ALQUILER","ALQUILERMKT","MARKETING","EVENTO","EVENTOMARKETING","EVENTOCORSAGA","JARDINERIA","MANTENIMIENTO",
          "REPARACIONMESAS","EQUIPOROSANTA","MERCADO","MERCADOVERDURAS","SEGURO","BASURA","CONTENIDO","INDIVIDUALES","MEJORAS",
          "CARPINTERIA","LAVANDERIA","UNIFORMES","BARRA","ADMIN","ADMINNOMBRAMIENTO","ADMINABOGADA","COMPRACONTARJETA",
          "BANCAELECTRONICA","INDUSTRIAS","NUEVOS","MARISCOS","LICORERA",""}
bi=wb["03_Banco_Industrial"]; BIPAT={}
for r in range(5,bi.max_row+1):
    q=num(bi.cell(r,4).value)
    if q<=0: continue
    g=str(bi.cell(r,3).value or "").strip(); cat=str(bi.cell(r,7).value or "").strip(); f=fch(bi.cell(r,1).value)
    G=up(g)
    if cat in EXC:
        mot=MOTIVO[cat]; EXCL[("BI",mot,cat)][0]+=1; EXCL[("BI",mot,cat)][1]+=q; rec("BI","excluido: "+mot,q)
    elif cat.endswith("_EFECTIVO") or G.startswith("ATM") or G.startswith("RETIRO"):
        k2=("BI",re.sub(r"\d","",G)[:40].strip()); EFEC[k2][0]+=1; EFEC[k2][1]+=q; rec("BI","efectivo sin proveedor",q)
    elif cat=="COMISIONES_BANCARIAS":
        e=ent(("BANCO","BI"),"BANCO INDUSTRIAL (servicios bancarios)","PROVEEDOR","ALTA"); e["glosas"]["BI: "+G]+=1; mov(e,"BI",cat,q=q,f=f); rec("BI","maestro",q)
    else:
        txt,comp=bi_patron(g)
        if not comp: comp="__SIN_TEXTO__"+cat; txt="(sin texto util · %s)" % cat
        p=BIPAT.setdefault(comp,{"txt":txt,"var":Counter(),"cat":Counter(),"n":0,"q":0.0,"f1":None,"f2":None})
        p["var"][g]+=1; p["cat"][cat]+=1; p["n"]+=1; p["q"]+=q
        if f and (not p["f1"] or f<p["f1"]): p["f1"]=f
        if f and (not p["f2"] or f>p["f2"]): p["f2"]=f
        rec("BI","glosa por asignar (hoja BI_glosas)",q)

# ---------------- TARJETA
tc=wb["05_Tarjeta_Credito_BAC"]; filas=[]
for r in range(5,tc.max_row+1):
    q=num(tc.cell(r,3).value); u=num(tc.cell(r,4).value)
    if q<=0 and u<=0: continue
    filas.append((str(tc.cell(r,2).value or ""),q,u,str(tc.cell(r,5).value or "").strip(),str(tc.cell(r,6).value or "").strip(),fch(tc.cell(r,1).value)))
pad={}
for raw,*_ in filas:
    m=re.match(r"^(.*?\S)\s{3,}(\S.*)$",raw)
    if m: pad[up(m.group(1))+" "+up(m.group(2))]=up(m.group(1))
def comercio(raw):
    m=re.match(r"^(.*?\S)\s{3,}(\S.*)$",raw)
    b=up(m.group(1)) if m else pad.get(up(raw),up(raw))
    b=re.sub(r"^FACEBK \*\S+","FACEBK",b)
    b=re.sub(r"\*ADS\d+","*ADS",b)
    b=re.sub(r"^(WIX\.COM)\s+\d+",r"\1",b)
    b=re.sub(r"^(PAYPAL \*|PP\*)(\d*)([A-Z]+)\d*",r"\1\3",b)
    b=re.sub(r"\s\d{6,}\b","",b)
    b=re.sub(r"\s[\d\-\.]{7,}$","",b)
    return re.sub(r"\s+"," ",b).strip()
for raw,q,u,cat,esp,f in filas:
    Ur=up(raw)
    if cat in EXC:
        mot=MOTIVO[cat]; EXCL[("TARJETA",mot,cat)][0]+=1; EXCL[("TARJETA",mot,cat)][1]+=q; EXCL[("TARJETA",mot,cat)][2]+=u; rec("TARJETA","excluido: "+mot,q,u)
    elif cat.endswith("_EFECTIVO") or Ur.startswith("RETIRO") or Ur.startswith("ATM"):
        EFEC[("TARJETA",comercio(raw))][0]+=1; EFEC[("TARJETA",comercio(raw))][1]+=q; EFEC[("TARJETA",comercio(raw))][2]+=u; rec("TARJETA","efectivo sin proveedor",q,u)
    elif cat=="COMISIONES_BANCARIAS" or re.match(r"(INTERES|CARGOS POR|CARGO POR|COBRO ADMTVO)",Ur):
        e=ent(("BANCO","BAC"),"BANCO BAC (servicios bancarios)","PROVEEDOR","ALTA"); e["glosas"]["TARJETA: "+Ur]+=1; mov(e,"TARJETA",cat,q=q,usd=u,f=f); rec("TARJETA","maestro",q,u)
    else:
        k=comercio(raw); e=ent(("TARJETA",k),k,None,"ALTA"); e["comercios"][raw.strip()]+=1
        e.setdefault("pers",Counter())[("P" if (cat=="PERSONAL" or up(esp) in ("SI","SÍ")) else "N")]+=1
        mov(e,"TARJETA",cat,q=q,usd=u,f=f); rec("TARJETA","maestro",q,u)
for e in ENT.values():
    if e["key"][0]=="TARJETA":
        e["tipo"]="PERSONAL" if e["pers"]["P"]>e["pers"]["N"] else "PROVEEDOR"
        tc_cats={o.split(":",1)[1] for o in e["obs"] if o.startswith("TARJETA:")}
        if len(tc_cats)>1: e["flags"].append("varias categorias en la tarjeta: %s" % ", ".join(sorted(tc_cats)))

# ---------------- codigos
orden=sorted(ENT.values(), key=lambda e:(re.sub(r"[^A-Z0-9]","",up(e["nombre"])), str(e["key"])))
for i,e in enumerate(orden,1): e["codigo"]="PRV-%04d" % i

# ---------------- candidatos a fusion (NO se aplican)
STOP={"LA","EL","LOS","LAS","DE","DEL","Y","SA","S","A","SOCIEDAD","ANONIMA","ANÓNIMA","DISTRIBUIDORA","DIST","COMERCIAL","RESTAURANTE",
      "CAFE","HOTEL","SUPER","SUPERMERCADO","TIENDA","CUENTA","BAC","GT","GUATEMALA","ANTIGUA","PAGO","CORPORACION","GRUPO","INVERSIONES",
      "SERVICIOS","PRODUCCIONES","COMPANY","INC","LLC","CO","COM","WWW","TITULAR","SIN","IDENTIFICAR","BANCO","CLINICA","FARMACIA","BAR","ALIMENTOS","BODEGA","FERRETERIA","PANADERIA","AGENCIA","CENTRO","TALLER","LIBRERIA","HOSPITAL","SUPERMERCADOS","MERCADO","VENTAS","VENTA","SALA","CASA","IMPORTADORA","EMPRESA","TIENDAS","RESTAURANTES","MULTIVENTAS","INDUSTRIAS","MATERIALES","DEPOSITO"}
def palabras(e):
    txt=" ".join([e["nombre"]]+list(e["est"])+list(e["legales"])+[c[0] for c in e["cat_alias"]])
    return [w for w in re.findall(r"[A-ZÁÉÍÓÚÑ]{2,}",up(txt).replace("*"," ")) if w not in STOP and not w.isdigit()]
def primera(e):
    w=[x for x in re.findall(r"[A-ZÁÉÍÓÚÑ]{2,}",up(e["nombre"]).replace("*"," ")) if x not in STOP]
    return w[0] if w and len(w[0])>=5 else None
def compacto(s): return re.sub(r"[^A-ZÁÉÍÓÚÑ0-9]","",up(s))
CAND=[]; L=list(ENT.values())
for i in range(len(L)):
    a=L[i]; pa=primera(a); ca=compacto(a["nombre"]); wa=set(palabras(a))
    for j in range(i+1,len(L)):
        b=L[j]; ev=None
        ambos_nit = a["key"][0]=="NIT" and b["key"][0]=="NIT"
        if ca and ca==compacto(b["nombre"]): ev="nombre identico"
        elif ambos_nit: ev=None   # dos NIT distintos son dos proveedores distintos: el NIT es la llave
        elif pa and pa==primera(b): ev="empiezan con la misma palabra: %s" % pa
        elif (a["key"][0]=="CUENTA" or b["key"][0]=="CUENTA"):
            comun={w for w in wa & set(palabras(b)) if len(w)>=6}
            if comun: ev="la cuenta y el proveedor comparten: %s" % ", ".join(sorted(comun))
        if ev: CAND.append((a,b,ev))

# ---------------- llaves de nombre de cada entidad, solo para SUGERIR
def es_persona(nombre): return re.search(r"\S\s*,\s+\S", str(nombre)) is not None and "S.A" not in up(nombre) and "SOCIEDAD" not in up(nombre)
LLAVES={}
for e in orden:
    ks={}
    def add(txt, prefijo, ks=ks):
        k=norm_letras(txt)
        if k: ks[k]=ks.get(k,False) or prefijo
    add(e["nombre"], not es_persona(e["nombre"]))
    for x in e["est"]: add(x, not es_persona(x))
    for x in e["legales"]: add(x, not es_persona(x))
    for c in e["cat_alias"]: add(c[0], not es_persona(c[0]))
    for x in e["comercios"]: add(x, True)
    if e["key"][0]=="CUENTA":
        doc=e["nombre"].split("·",1)[-1]
        for w in re.findall(r"\((.*?)\)",doc): add(w, True)
        add(re.sub(r"\(.*?\)","",doc), False)
    LLAVES[e["codigo"]]=ks
ACRONIMOS={"EEGSA":("326445","EEGSA es la sigla de Empresa Electrica de Guatemala, S.A. (NIT 326445)")}
def sugerir_nombre(comp):
    hits=[]
    for sig,(nit,ev) in ACRONIMOS.items():
        if comp.startswith(sig): hits+=[(e,ev) for e in orden if nit in e["nit"]]
    if len(comp)<4: return hits
    for e in orden:
        if any(e is h[0] for h in hits): continue
        for k,pref in LLAVES[e["codigo"]].items():
            if comp==k: hits.append((e,"nombre identico")); break
            if pref and len(comp)>=5 and k.startswith(comp): hits.append((e,"la glosa es el comienzo de '%s'" % k)); break
            if pref and len(k)>=6 and comp.startswith(k): hits.append((e,"la glosa empieza con '%s'" % k)); break
    return hits
CONFIRMADAS={
 ("TARJETA","PAYPAL *VANEWILCHES"):"Vanessa Wilches · especialista de pauta digital · honorarios profesionales",
 ("NIT","345377"):"distribuidor de licor · la clasificacion de barra la lleva Juanma",
 ("NIT","110989163"):"Migdalia Lico · proveedora de mercado (alimentos) · ya no se le compra",
 ("NIT","52496325"):"Cristina Anona Lico · proveedor actual de frutas y verduras",
 ("NIT","7616325"):"carniceria · alimentos",
 ("NIT","24545910"):"Juanma · factura como pequeno contribuyente · servicios profesionales, no nomina",
 ("NIT","80383955"):"alquiler de la maquina de agua",
 ("NIT","120373785"):"fisioterapia de Juanma · personal",
 ("NIT","120665379"):"2Onzas Mixology · proveedor de cocteleria",
 ("NIT","70408971"):"Jorge Valladares · proveedor de hielo · cocteleria",
}
for e in orden:
    if "ALQUIFIESTAS MEG" in {up(x) for x in e["est"]}: CONFIRMADAS[e["key"]]="Edi Joaquin Gaitan · alquiler de materiales y cristaleria para eventos"

# ---------------- escribir el libro
out=openpyxl.Workbook(); H=Font(bold=True,color="FFFFFF"); F=PatternFill("solid",fgColor="4E6D5A")
def hoja(nombre,cab,filas,anchos):
    ws=out.create_sheet(nombre); ws.append(cab)
    for c in ws[1]: c.font=H; c.fill=F; c.alignment=Alignment(wrap_text=True,vertical="center")
    for f in filas: ws.append(f)
    for i,w in enumerate(anchos,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions
    return ws
fmt=lambda d:d.strftime("%Y-%m-%d") if d else ""
M=[]
for e in orden:
    M.append([e["codigo"],e["nombre"],CONFIRMADAS.get(e["key"],""),e["tipo"],e["conf"],", ".join(sorted(e["nit"])),
      " | ".join(e["legales"]), " | ".join(e["est"]),
      " | ".join("%s (%s%s)" % (c[0],c[1] or "sin categoria",", "+c[2] if c[2] else "") for c in e["cat_alias"]),
      ", ".join(sorted(e["cuentas"])), " | ".join(e["comercios"]), " | ".join(list(e["glosas"])[:6]),
      ", ".join("%s x%d"%(k,v) for k,v in e["obs"].most_common()), e["n"], round(e["q"],2), round(e["usd"],2),
      fmt(e["f1"]), fmt(e["f2"]), ", ".join(sorted(e["fuentes"])), " · ".join(e["flags"])])
hoja("Maestro",["CODIGO","NOMBRE","IDENTIDAD CONFIRMADA POR JUANMA","TIPO","CONFIANZA","NIT","NOMBRE LEGAL (FEL)","ESTABLECIMIENTOS (FEL)","FILA EN CATALOGO (categoria, familia)",
   "CUENTAS BAC","TEXTO EN TARJETA","GLOSAS EN BANCOS","CATEGORIAS OBSERVADAS (fuente:categoria xN)","MOVIMIENTOS","TOTAL Q","TOTAL US$",
   "PRIMERA","ULTIMA","FUENTES","OBSERVACIONES"],M,[10,34,40,11,11,12,30,30,34,12,30,30,40,11,13,11,11,11,16,60])
hoja("Candidatos_fusion",["CODIGO A","NOMBRE A","FUENTE A","CODIGO B","NOMBRE B","FUENTE B","EVIDENCIA","SE FUSIONA? (Juanma)"],
   [[a["codigo"],a["nombre"],", ".join(sorted(a["fuentes"])),b["codigo"],b["nombre"],", ".join(sorted(b["fuentes"])),ev,""] for a,b,ev in sorted(CAND,key=lambda x:(x[2],x[0]["codigo"]))],
   [10,34,14,10,34,14,40,18])
def sugerir_cat(cat):
    c=[e for e in orden if any(o.split(":",1)[1].replace("_"," ").upper()==cat.replace("_"," ").upper() for o in e["obs"]) and e["tipo"]!="PERSONAL" and e["key"][0] in ("NIT","CUENTA")]
    return ", ".join("%s %s"%(e["codigo"],e["nombre"]) for e in c) if 0<len(c)<=2 else ""
BIF=[]; BI_CON=0; BI_SUG=0; BI_SIN=0; QS=[0.0,0.0,0.0]
for comp,p in sorted(BIPAT.items(), key=lambda kv:-kv[1]["q"]):
    cat_top=p["cat"].most_common(1)[0][0]
    if comp in CONCEPTO or comp.startswith("__SIN_TEXTO__"):
        tipo="concepto sin nombre"; sug=sugerir_cat(cat_top); ev="sugerencia solo por categoria (%s)" % cat_top if sug else ""
        BI_CON+=1
    else:
        tipo="con nombre"; hits=sugerir_nombre(comp)
        sug=" | ".join("%s %s"%(e["codigo"],e["nombre"]) for e,_ in hits[:3]) + (" | ..." if len(hits)>3 else "")
        ev=" | ".join(x for _,x in hits[:3])
        if not hits:
            sug=sugerir_cat(cat_top); ev=("sin coincidencia de nombre; sugerencia solo por categoria (%s)" % cat_top) if sug else ""
    if sug: BI_SUG+=1; QS[0]+=p["q"]
    else: BI_SIN+=1; QS[1]+=p["q"]
    BIF.append([p["txt"] or "(sin texto util)",tipo," | ".join(list(p["var"])[:4]),", ".join("%s x%d"%(k,v) for k,v in p["cat"].most_common()),
                p["n"],round(p["q"],2),fmt(p["f1"]),fmt(p["f2"]),sug,ev,""])
hoja("BI_glosas",["PATRON DE GLOSA","TIPO","TEXTOS ORIGINALES","CATEGORIAS","MOVIMIENTOS","MONTO Q","PRIMERA","ULTIMA",
     "CODIGO SUGERIDO (no aplicado)","EVIDENCIA","CODIGO (Juanma)"],BIF,[28,18,44,30,11,13,11,11,50,40,14])
PSN=[]
for x in sorted(PORA,key=lambda x:(x[0],x[1] or datetime.datetime(1900,1,1))):
    hits=sugerir_nombre(norm_letras(x[3])) if "ACH" not in up(x[3]) else []
    sug=" | ".join("%s %s"%(e["codigo"],e["nombre"]) for e,_ in hits[:3]) or sugerir_cat(x[4])
    PSN.append([x[0],fmt(x[1]),x[2],x[3],x[4],round(x[5],2),sug,""])
hoja("Pagos_sin_nombre",["FUENTE","FECHA","REFERENCIA","GLOSA","CATEGORIA","MONTO Q","SUGERENCIA (no aplicada)","CODIGO (Juanma)"],
     PSN,[9,11,13,34,24,12,50,14])
hoja("Excluidos",["FUENTE","MOTIVO","CATEGORIA / GRUPO","MOVIMIENTOS","MONTO Q","MONTO US$"],
   [[k[0],k[1],k[2],v[0],round(v[1],2),round(v[2],2)] for k,v in sorted(EXCL.items())],[10,36,30,12,14,12])
hoja("Efectivo",["FUENTE","RETIRO / CAJERO","MOVIMIENTOS","MONTO Q","MONTO US$"],
   [[k[0],k[1],v[0],round(v[1],2),round(v[2],2)] for k,v in sorted(EFEC.items(),key=lambda kv:-kv[1][1])],[10,44,12,14,12])
C=[]
for fu,d in REC.items():
    tf=sum(v[0] for v in d.values()); tq=sum(v[1] for v in d.values()); tu=sum(v[2] for v in d.values())
    for dest,v in sorted(d.items()): C.append([fu,dest,v[0],round(v[1],2),round(v[2],2)])
    C.append([fu,"TOTAL DE LA FUENTE",tf,round(tq,2),round(tu,2)])
hoja("Conciliacion",["FUENTE","DESTINO","FILAS","MONTO Q","MONTO US$"],C,[12,44,10,16,12])
tipos=Counter((e["tipo"],e["conf"]) for e in orden)
ws=out.active; ws.title="LEEME"
for line in [
 "MAESTRO DE PROVEEDORES · EXTRACCION CON CODIGO INTERNO",
 "Generado el 14 de septiembre de 2026.",
 "",
 "DE DONDE SALE",
 "Movimientos: espejo del maestro del 14-sep 15:51 (FEL, Banco Industrial, BAC y tarjeta). Las correcciones hechas a mano en el FEL despues de esa hora no se ven en 'CATEGORIAS OBSERVADAS'.",
 "Catalogo: 00_Proveedores leido EN VIVO despues de las correcciones de Juanma del 14-sep.",
 "",
 "COMO SE FORMO CADA PROVEEDOR (solo llaves duras, nunca por parecido de nombre)",
 "FEL: una fila por NIT. Si el NIT factura con un solo establecimiento, ese es el nombre; si tiene varios, el nombre legal.",
 "Catalogo: se une al NIT solo si el nombre es IDENTICO al establecimiento del FEL, que es exactamente lo que hace el VLOOKUP. Si no coincide con ninguno, queda como proveedor propio.",
 "BAC: una fila por numero de cuenta de destino. Las cuentas de planilla y la de devolucion de inversion no entran.",
 "Tarjeta: una fila por comercio, con el texto del estado de cuenta. Solo se juntan variantes que son exactamente el mismo texto con otro espaciado, y se quitan los numeros de transaccion de Facebook, Google Ads y Wix.",
 "Banco Industrial: la glosa es texto escrito a mano y NO identifica al proveedor. No genera codigo: se agrupa en patrones (hoja BI_glosas) con un codigo SUGERIDO y su evidencia, para que Juanma confirme.",
 "",
 "CODIGOS",
 "PRV-0001 en adelante, en orden alfabetico. Un codigo no se reutiliza: si dos filas resultan ser el mismo proveedor, una absorbe a la otra y el codigo absorbido queda retirado.",
 "",
 "LO QUE FALTA DECIDIR",
 "Candidatos_fusion: pares que PODRIAN ser el mismo proveedor (misma palabra inicial, nombre identico, cuenta con nombre parecido). Ninguno esta fusionado.",
 "BI_glosas: cada patron de glosa del Banco Industrial con su codigo sugerido. Pagos_sin_nombre: transferencias del BAC sin destinatario identificable (ACH) y otros pagos sin llave.",
 "Excluidos y Efectivo: no son proveedores (planilla, socios, impuestos, pagos de tarjeta, retiros de cajero). Estan para que la conciliacion cuadre.",
 "Conciliacion: cada fila de cada fuente cae en un solo destino. El total de cada fuente debe ser igual a la suma de sus destinos.",
 "",
 "RESUMEN",
]+["  %s · confianza %s: %d" % (t,c,n) for (t,c),n in sorted(tipos.items())]+[
 "  total de codigos: %d · candidatos a fusion: %d · patrones de glosa BI: %d · pagos sin nombre: %d" % (len(orden),len(CAND),len(BIPAT),len(PORA))]:
    ws.append([line])
ws.column_dimensions["A"].width=150; ws["A1"].font=Font(bold=True,size=14,color="4E6D5A")
for r in (4,8,15,18,24): ws.cell(r,1).font=Font(bold=True)
out.save(OUT)

print("  archivo:", OUT.split("/")[-1])
print("  codigos: %d · candidatos a fusion: %d · patrones BI: %d (con sugerencia %d, sin %d) · pagos sin nombre: %d" % (len(orden),len(CAND),len(BIPAT),BI_SUG,BI_SIN,len(PORA)))
for (t,c),n in sorted(tipos.items()): print("    %-9s %-5s %d" % (t,c,n))
print("  por origen:", dict(Counter(e["key"][0] for e in orden)))
print()
print("  conciliacion (filas / Q / US$):")
for fu,d in REC.items():
    print("    %-9s" % fu, {k:(v[0],round(v[1])) for k,v in d.items()})
print()
codigos=[e["codigo"] for e in orden]; assert len(codigos)==len(set(codigos)), "codigos repetidos"
ubic=Counter()
for e in orden:
    for c in e["cat_alias"]: ubic[up(c[0])]+=1
malos=[n for n,v in ubic.items() if v!=1 and n not in est2keys]; assert not malos, "catalogo mal ubicado: %s" % malos[:5]
assert set(ubic)==set(catv), "nombres de catalogo sin ubicar: %s" % list(set(catv)-set(ubic))[:5]
for fu in ("FEL","BAC","BI","TARJETA"):
    tf=sum(v[0] for v in REC[fu].values()); tq=sum(v[1] for v in REC[fu].values()); tu=sum(v[2] for v in REC[fu].values())
    assert tf==RAW[fu][0] and abs(tq-RAW[fu][1])<0.01 and abs(tu-RAW[fu][2])<0.01, "no cuadra %s: %s vs %s" % (fu,(tf,tq,tu),RAW[fu])
    print("  cuadra %-8s filas %4d · Q%12s · US$%9s" % (fu, tf, f"{tq:,.2f}", f"{tu:,.2f}"))
print("  verificaciones: codigos unicos OK · catalogo ubicado completo OK · 4 fuentes cuadradas OK")
print("  BI: Q%s con sugerencia · Q%s sin sugerencia" % (f"{QS[0]:,.0f}", f"{QS[1]:,.0f}"))
pickle.dump({"n":len(orden)}, open(SP+"/extraccion_ok.pkl","wb"))
