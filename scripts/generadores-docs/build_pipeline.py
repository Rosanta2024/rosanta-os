# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Pipeline de Contratación (5 etapas)",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.5 · Un proceso, no una corazonada",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
etapas=[
 ("Etapa 1 · PUBLICAR (Día 1)","Escribe una vacante específica: qué es el puesto, qué buscas, qué ofreces. Incluye un FILTRO de atención al detalle ('incluye la palabra carisma en la primera línea de tu aplicación') y descarta a quien no lo cumpla. Dónde publicar en Guatemala: grupos de empleo en Facebook (Antigua/Sacatepéquez), Instagram @rosanta, Computrabajo, Tecoloco, INTECAP/escuelas gastronómicas, y referidos del equipo (bono por referido que se queda 90+ días)."),
 ("Etapa 2 · FILTRAR (Días 2-5)","No leas aplicaciones completas. Escanea 3 no-negociables del puesto (ver hoja 'No-negociables'). Si no cumple los 3, no entrevistas."),
 ("Etapa 3 · ENTREVISTA (Días 5-7)","20 minutos. Las mismas 4 preguntas siempre, cada respuesta se puntúa 1-5 (ver hoja 'Entrevista'). Promedio 4+ pasa a prueba; menos de 4, se descarta."),
 ("Etapa 4 · PRUEBA (Días 8-9)","Un turno de prueba PAGADO junto al equipo. Observa: ¿llegó a tiempo? ¿cómo trató al equipo? ¿preguntó? ¿puede hacer el trabajo físicamente? ¿tu equipo actual querría trabajar con esa persona? Pide feedback al equipo después del turno."),
 ("Etapa 5 · INDUCCIÓN (Días 10-14)","Día 1: tour, presentaciones, recorrido por intranet/SOPs, acompañar a un miembro fuerte, sin trabajo independiente. Días 2-3: con mentor, tareas con guía. Días 4-5: trabajo independiente con check-in al inicio y fin de turno. Fin de semana 1: sentada de 15 min ('¿cómo vas? ¿qué te confunde? ¿qué necesitas?')."),
]
for tit,desc in etapas:
    wr(r,tit,12,True,VERDE); r+=1
    wr(r,desc,10,h=58); r+=1
    r+=1
wr(r,"LO ESENCIAL",12,True,COBRE); r+=1
for t in [
 "• El filtro de atención al detalle elimina 60-70% de los aplicantes de bajo esfuerzo antes de leer un solo CV.",
 "• Pide feedback a tu equipo actual tras el turno de prueba: ellos saben con quién quieren trabajar.",
 "• Mismas preguntas para todos, puntuadas 1-5: contratar deja de ser corazonada y se vuelve comparable.",
]:
    wr(r,t,10,bg=CREMA,h=26); r+=1

# ===== PIPELINE TRACKER =====
pt=wb.create_sheet("📋 Pipeline"); pt.sheet_view.showGridLines=False
widths=[22,20,22,16,14,12,14,30]
for i,w in enumerate(widths,start=1): pt.column_dimensions[get_column_letter(i)].width=w
pt.merge_cells("A1:H1")
t=pt.cell(1,1,"ROSANTA · Seguimiento de candidatos"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); pt.row_dimensions[1].height=24
hdr=["Candidato","Puesto","Fuente","Etapa","Puntaje entrevista (1-5)","Filtro OK","Fecha de prueba","Notas"]
for c,h in enumerate(hdr,start=1):
    cell=pt.cell(3,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
pt.row_dimensions[3].height=36
ej=["Ej. Ana López","Mesero","Facebook empleos","Entrevista",4.5,"Sí","","Buena actitud, disponible fin de semana"]
for c,v in enumerate(ej,start=1):
    cell=pt.cell(4,c,v); cell.font=F(9,italic=True,color="6B5A45"); cell.fill=fill(CREMA); cell.border=border; cell.alignment=wrap
for rr in range(5,25):
    for c in range(1,9):
        cell=pt.cell(rr,c); cell.border=border; cell.alignment=wrap; cell.font=F(9)
    pt.row_dimensions[rr].height=22
dv_stage=DataValidation(type="list",formula1='"Publicado,Filtrado,Entrevista,Prueba,Inducción,Contratado,Descartado"',allow_blank=True); pt.add_data_validation(dv_stage); dv_stage.add("D4:D24")
dv_f=DataValidation(type="list",formula1='"Sí,No"',allow_blank=True); pt.add_data_validation(dv_f); dv_f.add("F4:F24")

# ===== NO-NEGOCIABLES =====
nn=wb.create_sheet("📋 No-negociables"); nn.sheet_view.showGridLines=False
for col,w in [("A",4),("B",26),("C",70)]: nn.column_dimensions[col].width=w
nn.merge_cells("A1:C1")
t=nn.cell(1,1,"ROSANTA · 3 no-negociables por puesto (para filtrar rápido)"); t.font=F(13,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); nn.row_dimensions[1].height=24
nn.cell(3,2,"Puesto").font=F(9,True,"FFFFFF"); nn.cell(3,2).fill=fill(TIERRA); nn.cell(3,2).border=border
nn.cell(3,3,"Los 3 no-negociables (si no cumple los 3, no entrevistas)").font=F(9,True,"FFFFFF"); nn.cell(3,3).fill=fill(TIERRA); nn.cell(3,3).border=border; nn.cell(3,3).alignment=wrapc
nn.row_dimensions[3].height=20
roles=[
 ("Mesero","1+ año de experiencia en servicio · disponibilidad fin de semana (vie-dom) · incluyó la palabra filtro"),
 ("Coctelero / Mixólogo","experiencia en barra/coctelería · disponibilidad fin de semana · incluyó la palabra filtro"),
 ("Cocina (línea)","experiencia en cocina · manejo de tiempos e inocuidad · incluyó la palabra filtro"),
 ("Prep / Lavaplatos","disponibilidad de horario · actitud y limpieza comprobables · incluyó la palabra filtro"),
 ("Anfitrión / Host","trato al cliente y presentación · disponibilidad fin de semana · incluyó la palabra filtro"),
]
rr=4
for rol,crit in roles:
    a=nn.cell(rr,2,rol); a.font=F(9,True,VERDE); a.alignment=wrap; a.border=border; a.fill=fill(CREMA)
    b=nn.cell(rr,3,crit); b.font=F(9); b.alignment=wrap; b.border=border
    nn.row_dimensions[rr].height=32; rr+=1
nn.cell(rr+1,2,"Ajusta los criterios a cada vacante concreta.").font=F(9,italic=True,color="6B5A45")

# ===== ENTREVISTA / FILTRO / INDUCCIÓN =====
iq=wb.create_sheet("📋 Entrevista e inducción"); iq.sheet_view.showGridLines=False
for col,w in [("A",4),("B",44),("C",40),("D",30)]: iq.column_dimensions[col].width=w
iq.merge_cells("A1:D1")
t=iq.cell(1,1,"ROSANTA · Entrevista, filtro e inducción"); t.font=F(13,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); iq.row_dimensions[1].height=24
# preguntas
iq.cell(3,1,"PREGUNTAS DE ENTREVISTA (mismas para todos · puntúa 1-5)").font=F(11,True,VERDE)
for c,h in enumerate(["#","Pregunta","Qué buscas","Guía de puntaje"],start=1):
    cell=iq.cell(4,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=wrapc; cell.border=border
iq.row_dimensions[4].height=20
qs=[
 ("Cuéntame de un turno que se salió de control. ¿Qué hiciste?","Resolución de problemas bajo presión","1=Pobre · 3=Aceptable · 5=Excelente"),
 ("¿Qué hace un gran [puesto] que uno promedio no hace?","Autoconciencia y estándar propio","1=Pobre · 3=Aceptable · 5=Excelente"),
 ("¿Por qué dejas tu trabajo actual?","Banderas rojas: culpa, drama, inestabilidad","1=Pobre · 3=Aceptable · 5=Excelente"),
 ("¿Qué preguntas tienes para mí?","Interés genuino en el puesto y Rosanta","1=Pobre · 3=Aceptable · 5=Excelente"),
]
rr=5
for i,(q,look,guide) in enumerate(qs,start=1):
    iq.cell(rr,1,i).font=F(9,True); iq.cell(rr,1).alignment=ctr; iq.cell(rr,1).border=border
    for c,v in [(2,q),(3,look),(4,guide)]:
        cell=iq.cell(rr,c,v); cell.font=F(9); cell.alignment=wrap; cell.border=border
    iq.row_dimensions[rr].height=34; rr+=1
iq.cell(rr,2,"Promedio 4+ → pasa a turno de prueba. Menos de 4 → se descarta.").font=F(9,italic=True,color="6B5A45"); rr+=2
# filtro
iq.cell(rr,1,"FILTRO DE ATENCIÓN AL DETALLE").font=F(11,True,VERDE); rr+=1
iq.cell(rr,2,"Palabra filtro:").font=F(10,True); c=iq.cell(rr,3,"carisma"); c.font=F(10,True,COBRE); c.fill=fill(AMBAR); c.border=border; rr+=1
iq.cell(rr,2,"Dónde incluirla:").font=F(10,True); iq.cell(rr,3,"En la PRIMERA línea de la aplicación. Ejemplo: 'Incluye la palabra carisma al inicio de tu mensaje.'").font=F(9); iq.cell(rr,3).alignment=wrap; iq.row_dimensions[rr].height=30; rr+=2
# induccion dia 1
iq.cell(rr,1,"CHECKLIST DE INDUCCIÓN · DÍA 1").font=F(11,True,VERDE); rr+=1
for c,h in enumerate(["#","Tarea","Responsable","Notas"],start=1):
    cell=iq.cell(rr,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=wrapc; cell.border=border
iq.row_dimensions[rr].height=18; rr+=1
tasks=[
 ("Tour: sala, barra, cocina, jardín/huerto, cava, chimenea, baños","Jefe de área"),
 ("Presentación con el equipo del turno","Jefe de área"),
 ("Recorrido por intranet y SOPs (Manual de Servicio, SOP Sala/Cocina)","Jefe de área"),
 ("Explicar el concepto gastrococtelería y la marca (Cocina con Carisma, 14·91)","Jefe de área"),
 ("Acompañar (shadow) a un miembro fuerte del equipo · sin trabajo independiente","Mentor asignado"),
 ("Uniforme, casillero, horarios y política de propinas","Jefe de área"),
 ("Repaso de menú y alérgenos básicos","Chef / Jefe de sala"),
 ("Cierre del día: dudas y expectativas","Jefe de área"),
]
for i,(task,owner) in enumerate(tasks,start=1):
    iq.cell(rr,1,i).font=F(9,True); iq.cell(rr,1).alignment=ctr; iq.cell(rr,1).border=border
    iq.cell(rr,2,task).font=F(9); iq.cell(rr,2).alignment=wrap; iq.cell(rr,2).border=border
    iq.cell(rr,3,owner).font=F(9); iq.cell(rr,3).alignment=wrap; iq.cell(rr,3).border=border
    iq.cell(rr,4).border=border; iq.cell(rr,4).fill=fill(CREMA)
    iq.row_dimensions[rr].height=28; rr+=1

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Pipeline_Contratacion.xlsx"
wb.save(out); print("guardado:",out)
