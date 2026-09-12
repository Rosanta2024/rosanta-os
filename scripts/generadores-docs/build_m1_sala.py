# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"
FVERDE=PatternFill("solid",fgColor=VERDE)
FCREMA=PatternFill("solid",fgColor=CREMA)
FCOBRE=PatternFill("solid",fgColor="F7E9CC")
FSUB=PatternFill("solid",fgColor="EDE4D0")
thin=Side(style="thin",color="D8CDB5")
B=Border(left=thin,right=thin,top=thin,bottom=thin)
TITLE=Font(name="Arial",size=16,bold=True,color="FFFFFF")
SUBT=Font(name="Arial",size=10,italic=True,color="FFFFFF")
H=Font(name="Arial",size=11,bold=True,color=VERDE)
HW=Font(name="Arial",size=10,bold=True,color="FFFFFF")
SEC=Font(name="Arial",size=10,bold=True,color=MARRON)
TXT=Font(name="Arial",size=10,color=MARRON)
SM=Font(name="Arial",size=9,color="6B5A45")
wrap=Alignment(wrap_text=True,vertical="top")
ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
left=Alignment(horizontal="left",vertical="center",wrap_text=True)

wb=openpyxl.Workbook()

# ---------- Hoja 1: Rol e instrucciones ----------
ws=wb.active; ws.title="M1 Sala · Rol"
ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=2
ws.column_dimensions["B"].width=30
ws.column_dimensions["C"].width=78
ws.merge_cells("B1:C1"); c=ws["B1"]; c.value="ROSANTA · M1 de Sala"; c.font=TITLE; c.fill=FVERDE; c.alignment=Alignment(vertical="center",indent=1)
ws.merge_cells("B2:C2"); c=ws["B2"]; c.value="Instrumento diario del jefe de turno · Cocina con Carisma"; c.font=SUBT; c.fill=FVERDE; c.alignment=Alignment(vertical="center",indent=1)
ws.row_dimensions[1].height=26; ws.row_dimensions[2].height=16

rows=[
 ("¿Qué es el M1?","El M1 es el primer nivel de gestión: el jefe de turno que dirige a los especialistas de la sala y corre el scorecard diario sobre ellos. No es un puesto nuevo; es un rol que alguien ocupa cada turno."),
 ("¿Quién es el M1?","Fin de semana: José (jefe de sala); su segundo es Efraín.  Días bajos (lun–mié): el coctelero, que además atiende la barra."),
 ("¿A quién gestiona?","A los especialistas de sala del turno: mesero(s) y coctelero. En días bajos gestiona 2 personas (él como coctelero + 1 mesero)."),
 ("Los 2 SOPs que usa","SOP de Sala (mesero + recepción) y SOP de Barra (coctelero). De ahí salen las tareas verificables del Paso 1."),
 ("Cómo corre el día","1) Antes de abrir: asigna tareas de apertura y repasa el feature.  2) Durante: vigila tiempos, calidad y maridajes.  3) Al cerrar: llena el scorecard de cada especialista y define el plan de mañana."),
 ("Método (RAA)","Resultado → Análisis → Acción. El resultado se toma de una fuente independiente (POS, reservas WIX, quejas), no de la opinión del empleado."),
 ("Los 7 pasos","1 Cantidad · 2 Calidad · 3 Horas · 4 Resultados operativos · 5 Resultados financieros · 6 Categoría (1–5) · 7 Plan de acción para mañana."),
 ("Cortesías/quejas","Toda cortesía o descuento se registra en el POS con motivo. Las quejas se manejan con el protocolo CERA (Cálmate y escucha, Empatiza, Resuelve, Agradece)."),
]
r=4
for k,v in rows:
    ws.merge_cells(f"B{r}:B{r}"); a=ws[f"B{r}"]; a.value=k; a.font=SEC; a.fill=FCREMA; a.alignment=wrap; a.border=B
    b=ws[f"C{r}"]; b.value=v; b.font=TXT; b.alignment=wrap; b.border=B
    ws.row_dimensions[r].height=14+ (len(v)//60)*13 + 26
    r+=1

# ---------- Hoja 2: Listas de tareas (Paso 1) ----------
def tasklist(title, subtitle, blocks):
    s=wb.create_sheet(title)
    s.sheet_view.showGridLines=False
    s.column_dimensions["A"].width=2
    s.column_dimensions["B"].width=6
    s.column_dimensions["C"].width=74
    s.column_dimensions["D"].width=12
    s.merge_cells("B1:D1"); c=s["B1"]; c.value=subtitle; c.font=TITLE; c.fill=FVERDE; c.alignment=Alignment(vertical="center",indent=1)
    s.merge_cells("B2:D2"); c=s["B2"]; c.value="Tareas verificables del Paso 1 (Cantidad) · derivadas del SOP"; c.font=SUBT; c.fill=FVERDE; c.alignment=Alignment(vertical="center",indent=1)
    s.row_dimensions[1].height=24; s.row_dimensions[2].height=15
    # header
    r=4
    for col,txt in (("B","#"),("C","Tarea"),("D","¿Hecha?")):
        c=s[f"{col}{r}"]; c.value=txt; c.font=HW; c.fill=PatternFill("solid",fgColor=TIERRA); c.alignment=ctr; c.border=B
    r+=1
    dv=DataValidation(type="list",formula1='"Sí,No,N/A"',allow_blank=True); s.add_data_validation(dv)
    n=1
    for sec,items in blocks:
        s.merge_cells(f"B{r}:D{r}"); c=s[f"B{r}"]; c.value=sec; c.font=SEC; c.fill=FSUB; c.alignment=Alignment(vertical="center",indent=1); c.border=B
        r+=1
        for it in items:
            s[f"B{r}"].value=n; s[f"B{r}"].font=Font(name="Arial",size=9,bold=True,color=VERDE); s[f"B{r}"].alignment=ctr; s[f"B{r}"].border=B
            s[f"C{r}"].value=it; s[f"C{r}"].font=TXT; s[f"C{r}"].alignment=wrap; s[f"C{r}"].border=B
            s[f"D{r}"].border=B; s[f"D{r}"].alignment=ctr; dv.add(s[f"D{r}"])
            s.row_dimensions[r].height=14+(len(it)//60)*13+16
            n+=1; r+=1
    return s

# --- MESERO ---
tasklist("Tareas Mesero","MESERO · Sala",[
 ("Apertura (antes de abrir)",[
   "Revisar reservas y notas de grupos del día en WIX",
   "Montar y pulir mesas; revisar copas antes de sacarlas",
   "Surtir estaciones: cubiertos, servilletas, pan, agua",
   "Verificar ambiente: música, luz, jardín/pérgola/chimenea, baños",
   "Repasar el feature del día y su maridaje (del Huddle)",
   "Uniforme y presentación personal",
 ]),
 ("Durante el servicio",[
   "Recepción en ≤15s a quien esté más cerca de la puerta (no hay host)",
   "Presentar la gastrococtelería, el menú y el feature del día",
   "Tomar el pedido en el POS: numerar puestos, confirmar alergias y avisar a cocina",
   "Servir bebidas en 3–5 min; controlar tiempos de entrantes y fuertes con maridaje",
   "Check-back poco después de los fuertes",
   "Ofrecer postres o un segundo coctel",
   "Verificar la cuenta antes de cerrarla y entregarla correcta y a tiempo",
   "Preguntar por factura (FEL); despedir con calidez e invitar a reseña con el QR",
 ]),
 ("Cierre",[
   "Side work por estación: limpiar, resurtir y guardar",
   "Ordenar el ambiente: chimenea, música y luces exteriores",
   "Cuadre de propinas",
   "Dejar todo montado para el siguiente turno",
 ]),
])

# --- COCTELERO ---
tasklist("Tareas Coctelero","COCTELERO · Barra",[
 ("Apertura de barra",[
   "Revisar reservas y eventos del día (WIX) para estimar volumen",
   "Encender y verificar equipo: refrigeración, hielo, licuadora, cristalería",
   "Revisar niveles par de licores, mixers, jugos, garnishes y hielo",
   "Reponer desde bodega y anotar faltantes para el pedido",
   "Montar la estación: jiggers, coladores, tabla, cuchillo, trapos",
   "Repasar el coctel/feature del día y su maridaje",
   "Higiene y uniforme",
 ]),
 ("Pedidos e insumos",[
   "Comparar existencias contra los niveles par",
   "Registrar faltantes en la requisición",
   "Enviar el pedido a compras en el día/hora acordados",
   "Al recibir: verificar cantidad, calidad y fechas; guardar con rotación PEPS",
   "Registrar mermas y roturas; avisar al supervisor",
 ]),
 ("Preparaciones",[
   "Batching de cocteles de alta rotación, etiquetado y fechado",
   "Garnishes, jarabes, infusiones y jugos frescos",
   "Preparar hielo según formato",
   "Cumplir receta estándar con jigger (sin 'a ojo')",
 ]),
 ("Servicio en barra",[
   "Ordenar tickets del POS por tiempo y maridaje",
   "Preparar con receta estándar; cuidar presentación y garnish",
   "Coordinar tiempos con cocina (maridajes salen con su plato)",
   "Mantener consistencia: mismo sabor, medida y presentación",
 ]),
 ("Cierre de barra",[
   "Guardar y etiquetar preparaciones; desechar lo vencido",
   "Limpieza profunda: barra, pozo de hielo, refrigeración, piso",
   "Cuadre de existencias clave (premium) y mermas del turno",
   "Cuadre de propinas de barra",
   "Dejar la requisición/faltantes lista para el día siguiente",
   "Dejar la estación montada para el siguiente turno",
 ]),
])

# ---------- Hoja 4: Scorecard diario M1 ----------
sc=wb.create_sheet("Scorecard Diario")
sc.sheet_view.showGridLines=False
widths={"A":2,"B":18,"C":13,"D":13,"E":15,"F":22,"G":22,"H":12,"I":30}
for k,v in widths.items(): sc.column_dimensions[k].width=v
sc.merge_cells("B1:I1"); c=sc["B1"]; c.value="ROSANTA · Scorecard Diario del M1 · Sala"; c.font=TITLE; c.fill=FVERDE; c.alignment=Alignment(vertical="center",indent=1)
sc.merge_cells("B2:I2"); c=sc["B2"]; c.value="Resultado → Análisis → Acción · el resultado viene de fuente independiente (POS, WIX, quejas)"; c.font=SUBT; c.fill=FVERDE; c.alignment=Alignment(vertical="center",indent=1)
sc.row_dimensions[1].height=24; sc.row_dimensions[2].height=15
# meta row
sc["B4"].value="Fecha:"; sc["B4"].font=SEC
sc["D4"].value="M1 (jefe de turno):"; sc["D4"].font=SEC
sc.merge_cells("D4:E4")
sc["F4"].value="Tipo de día:"; sc["F4"].font=SEC
sc["G4"].value="(bajo / fin de semana)"; sc["G4"].font=SM
# header
hdr=["Especialista","1 · Cantidad","2 · Calidad","3 · Horas (prog/real)","4 · Result. operativos","5 · Result. financieros","6 · Categoría","7 · Plan de acción (mañana)"]
r=6
for i,txt in enumerate(hdr):
    col=chr(ord("B")+i); c=sc[f"{col}{r}"]; c.value=txt; c.font=HW; c.fill=PatternFill("solid",fgColor=TIERRA); c.alignment=ctr; c.border=B
sc.row_dimensions[r].height=30
dv_eval=DataValidation(type="list",formula1='"Excelente,Bueno,Malo,Terrible"',allow_blank=True); sc.add_data_validation(dv_eval)
dv_cat=DataValidation(type="list",formula1='"1 · A-player,2 · Bueno/flojo,3 · Aprendiendo,4 · No mejora,5 · Flojo"',allow_blank=True); sc.add_data_validation(dv_cat)
for r in range(7,15):
    for i in range(8):
        col=chr(ord("B")+i); cell=sc[f"{col}{r}"]; cell.border=B; cell.alignment=wrap; cell.font=TXT
    dv_eval.add(sc[f"C{r}"]); dv_eval.add(sc[f"D{r}"]); dv_cat.add(sc[f"H{r}"])
    sc.row_dimensions[r].height=34
# hint row
sc.merge_cells("B16:I16"); c=sc["B16"]; c.value="Paso 1 (Cantidad): califica contra las listas de las hojas 'Tareas Mesero' y 'Tareas Coctelero'.  Paso 6: categoría según pasos 1–5.  Paso 7: una acción concreta y medible para mañana."; c.font=SM; c.alignment=wrap
sc.row_dimensions[16].height=30

wb.save("/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_M1_Sala.xlsx")
print("guardado M1 Sala")
