# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"
AMBAR="FBE6C7"; AZUL="DBE7F1"  # ofensivo / defensivo tints
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Clasificación de Jugadores (Ofensivo vs. Defensivo)",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.2 · Manéjalos distinto o pierdes a los dos tipos",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"LOS DOS TIPOS (ninguno es mejor)",12,True,COBRE); r+=1
wr(r,"OFENSIVO: sugiere mejoras, toma iniciativa fuera de su rol, se aburre con la rutina, es variable turno a turno, a veces se desvía del SOP. Responde a autonomía, propiedad y retos nuevos. Riesgo: aburrimiento — si no lo retas, se va.",10,bg=AMBAR,h=44); r+=1
wr(r,"DEFENSIVO: consistente turno a turno, sigue el SOP exacto, confiable y predecible, no propone cambios, incómodo fuera de su rol. Responde a estabilidad, claridad y reconocimiento por la consistencia. Riesgo: estrés — si lo sobre-exiges, se quiebra o renuncia.",10,bg=AZUL,h=44); r+=1
r+=1
wr(r,"CÓMO MANEJAR A CADA UNO",12,True,VERDE); r+=1
wr(r,"Ofensivos: dales propiedad de una métrica o proyecto (no solo tareas), rétalos con problemas a resolver, autonomía con guardarraíl (resultado fijo, método libre), revísalos por resultados no por actividad, promuévelos rápido o te superan.",10,h=44); r+=1
wr(r,"Defensivos: expectativas claras y rutinas consistentes, reconoce su consistencia en público, no los sorprendas con cambios de rol sin preparación, entrénalos despacio antes de ampliar su alcance. Son la columna vertebral: trátalos así.",10,h=44); r+=1
r+=1
wr(r,"ERRORES COMUNES",12,True,VERDE); r+=1
for t in [
 "1. Promover a un defensivo a gerencia solo por antigüedad. La antigüedad no es liderazgo.",
 "2. Dejar que los ofensivos hagan lo que quieran. Autonomía no es anarquía: el resultado es fijo, el método flexible.",
 "3. Ignorar a los defensivos. Como no exigen atención, un día se van por sentirse invisibles. Reconoce la consistencia: es rara y valiosa.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"OJO — no confundir con el scorecard",12,True,VERDE); r+=1
wr(r,"El TIPO de jugador (cómo manejarlo) es distinto de las 5 CATEGORÍAS del scorecard de 7 pasos (que miden desempeño). Un defensivo puede ser categoría 1 (jugador A). Se complementan.",10,bg=CREMA,h=30)

# ===== CLASIFICACIÓN =====
cl=wb.create_sheet("📋 Clasificación"); cl.sheet_view.showGridLines=False
widths=[4,20,22,11,11,11,12,10,12,44]
for i,w in enumerate(widths,start=1): cl.column_dimensions[get_column_letter(i)].width=w
cl.merge_cells("A1:J1")
t=cl.cell(1,1,"ROSANTA · Clasificación de Jugadores"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); cl.row_dimensions[1].height=24
cl.merge_cells("A2:J2")
cl.cell(2,1,"Respuestas cargadas de tu formulario. O=señal Ofensiva, D=señal Defensiva. 3+ de un tipo define al jugador. Celdas editables con lista.").font=F(9,italic=True,color="6B5A45")
hdr=["#","Nombre","Puesto","1. Sugiere mejoras","2. Toma iniciativa","3. Se frustra c/ rutina","4. Consistencia","5. Sigue SOPs","Tipo","Acción de manejo"]
for c,h in enumerate(hdr,start=1):
    cell=cl.cell(4,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
cl.row_dimensions[4].height=44
# datos: nombre, puesto, [q1..q5] O/D, accion
S={"O":"Ofensivo","D":"Defensivo"}
data=[
 ("Jeffry López","Jefe de Cocina",["O","D","D","D","O"],
  "Defensivo que lidera cocina: reconoce su apego a fichas técnicas (clave para CMV) y dale rutina clara. Como sí propone y a veces se desvía, ábrele un canal formal (una idea de menú al mes) con guardarraíl: la receta estándar manda. POTENCIAL OFENSIVO: tiene capacidad de ser ofensivo con el estímulo adecuado. Dale un reto acotado (p. ej. liderar el menu engineering trimestral con guardarraíl de CMV) y mide si responde."),
 ("Nadia del Águila","2º de Cocina",["D","D","D","D","D"],
  "Columna vertebral de cocina. Expectativas claras, rutina estable y reconocimiento público de su consistencia. Entrénala despacio antes de ampliar su alcance."),
 ("Fernanda Pedroza","3º de Cocina",["D","D","D","D","D"],
  "Defensiva de medio tiempo. Claridad y rutina; reconoce su fiabilidad en horas pico. No la sorprendas con cambios."),
 ("Yazmin","Lavaplatos",["D","D","D","D","D"],
  "Backbone silencioso. Reconoce que la cocina siempre está lista y limpia. No la vuelvas invisible."),
 ("José Mazate","Jefe de Sala / Mixólogo",["O","O","D","D","O"],
  "TU ÚNICO OFENSIVO. Dale propiedad de métricas de sala/barra (rotación, ticket, CMV barra), rétalo con problemas, autonomía con guardarraíl (resultado fijo, método libre), revísalo por resultados. Promuévelo rápido: es tu candidato natural a Jefe de Operaciones/GM (#1 de tu escalera). Al subirlo, planifica quién cubre sala."),
 ("Efraín Taquez","2º de Sala / Mixólogo",["D","D","D","D","D"],
  "2º de sala confiable. Rutina clara y reconocimiento de consistencia. Es tu ancla operativa de sala cuando José crezca."),
 ("Marvin Pamal","Mesero",["O","O","D","D","D"],
  "Confiable con chispa de iniciativa (propone y toma iniciativa). Reconoce su consistencia y dale pequeñas responsabilidades extra, con preparación. POTENCIAL OFENSIVO: dale una micro-métrica propia (upsell de postre/coctel de su sección) y más reto gradual; puede crecer hacia líder de sala."),
 ("Eddy Cariño","Mesero",["D","D","D","D","D"],
  "Defensivo. Rutina clara y reconocimiento por su fiabilidad en el servicio."),
 ("Melvin Cojolón","Mesero",["D","D","D","D","D"],
  "Defensivo de medio tiempo. Rutina, claridad y reconocimiento. Entrénalo despacio si amplías su rol."),
 ("Mario Valle","Jardín / Mantenimiento",["D","D","D","D","D"],
  "Backbone de mantenimiento. Expectativas claras y reconoce que nada se cae en servicio. No lo dejes invisible."),
 ("Vanessa Wilches","Pauta (externo)",["O","D","D","D","D"],
  "Propone mejoras pero ejecuta sobre lo dado. Briefs claros y consistentes; reconoce el cumplimiento de metas de pauta."),
 ("Daniel Flores","Video / Foto (externo)",["O","D","D","D","D"],
  "Propone pero sigue el brief. Dale lineamientos claros y reconoce entregas consistentes de contenido."),
]
row=5
for i,(nom,pue,qs,acc) in enumerate(data,start=1):
    o=qs.count("O"); tipo="Ofensivo" if o>=3 else "Defensivo"
    cl.cell(row,1,i).font=F(9,True); cl.cell(row,1).alignment=ctr; cl.cell(row,1).border=border
    cl.cell(row,2,nom).font=F(9,True,VERDE); cl.cell(row,2).alignment=wrap; cl.cell(row,2).border=border
    cl.cell(row,3,pue).font=F(9); cl.cell(row,3).alignment=wrap; cl.cell(row,3).border=border
    for j,ans in enumerate(qs):
        c=cl.cell(row,4+j,S[ans]); c.font=F(9); c.alignment=ctr; c.border=border
        c.fill=fill(AMBAR if ans=="O" else AZUL)
    tc=cl.cell(row,9,tipo); tc.font=F(9,True,MARRON); tc.alignment=ctr; tc.border=border
    tc.fill=fill(AMBAR if tipo=="Ofensivo" else AZUL)
    ac=cl.cell(row,10,acc); ac.font=F(9); ac.alignment=wrap; ac.border=border; ac.fill=fill(CREMA)
    cl.row_dimensions[row].height=64; row+=1
nb=cl.cell(row,2,"Potencial ofensivo a desarrollar (defensivos hoy, con capacidad si se les da el estímulo): Jeffry López y Marvin Pamal.")
cl.merge_cells(start_row=row,start_column=2,end_row=row,end_column=10)
nb.font=F(9,italic=True,color="6B5A45"); nb.alignment=wrap; cl.row_dimensions[row].height=26
row+=1
# validaciones
dvq=DataValidation(type="list",formula1='"Ofensivo,Defensivo"',allow_blank=True); cl.add_data_validation(dvq); dvq.add(f"D5:H{row-1}")
dvt=DataValidation(type="list",formula1='"Ofensivo,Defensivo"',allow_blank=True); cl.add_data_validation(dvt); dvt.add(f"I5:I{row-1}")

# ===== DIAGNÓSTICO =====
dg=wb.create_sheet("🔎 Diagnóstico"); dg.sheet_view.showGridLines=False
for col,w in [("A",3),("B",32),("C",74)]: dg.column_dimensions[col].width=w
dg.row_dimensions[1].height=26
dg.cell(1,2,"ROSANTA · Qué revela la clasificación").font=F(15,True,VERDE)
dg.cell(2,2,"Resultado: 1 Ofensivo (José Mazate) · 11 Defensivos, de los cuales 2 con potencial ofensivo (Jeffry, Marvin).").font=F(11,True,COBRE)
findings=[
 ("Banca de gerentes delgada","Tu único jugador ofensivo es José. Los mejores futuros gerentes salen de los ofensivos, así que él es tu candidato a Jefe de Operaciones/GM (#1 de tu escalera). Riesgo de un solo punto: si lo subes, ¿quién dirige sala? Buena noticia: tienes banca de desarrollo — Jeffry (cocina) y Marvin (sala) son defensivos HOY pero con capacidad ofensiva si les das el estímulo adecuado. Desarróllalos con retos acotados y prepara a Efraín como ancla de sala."),
 ("Cocina liderada por un defensivo","Jeffry es defensivo: excelente para consistencia, apego a fichas y CMV, pero no esperes que impulse innovación. El cambio en cocina lo empujas tú o el sistema (recetario, scorecard), no él."),
 ("Equipo muy defensivo = arma de doble filo","Fortaleza: consistencia, apego a SOP, bajo drama, servicio predecible. Riesgo: al introducir cambios, hazlo despacio, claro y con preparación, o los estresas."),
 ("Reconoce la consistencia (11 personas)","Los defensivos no piden atención y un día se van por sentirse invisibles. Mete el reconocimiento de la consistencia en tu rutina semanal (conecta con el scorecard: suelen ser categoría 1 sin exigir nada)."),
 ("Ofensivo con guardarraíl","A José dale autonomía pero con resultado fijo (SOP de servicio, estándar de marca). Autonomía no es anarquía."),
 ("No es desempeño","Defensivo no significa bajo rendimiento. Nadia, Efraín o Mario pueden ser jugadores A (categoría 1 del scorecard) siendo defensivos."),
]
rr=4
dg.cell(rr,2,"Hallazgo").font=F(9,True,"FFFFFF"); dg.cell(rr,2).fill=fill(TIERRA); dg.cell(rr,2).border=border; dg.cell(rr,2).alignment=wrapc
dg.cell(rr,3,"Qué hacer").font=F(9,True,"FFFFFF"); dg.cell(rr,3).fill=fill(TIERRA); dg.cell(rr,3).border=border; dg.cell(rr,3).alignment=wrapc
dg.row_dimensions[rr].height=20; rr+=1
for h,d in findings:
    a=dg.cell(rr,2,h); a.font=F(9,True,VERDE); a.alignment=wrap; a.border=border; a.fill=fill(CREMA)
    b=dg.cell(rr,3,d); b.font=F(9); b.alignment=wrap; b.border=border
    dg.row_dimensions[rr].height=60; rr+=1

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Clasificacion_Jugadores.xlsx"
wb.save(out); print("guardado:",out)
