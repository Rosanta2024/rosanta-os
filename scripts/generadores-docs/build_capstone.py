# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Capstone del Curso 2 (verificación)",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Management OS · Confirma que todo esté construido, completo y conectado",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"QUÉ ES EL CAPSTONE",12,True,COBRE); r+=1
wr(r,"No es una lección nueva: es el paso final de verificación. Revisa que cada entregable del curso exista, esté completo y bien conectado con los demás. La hoja de al lado lista los 16 entregables, su archivo, estado y el pendiente real de cada uno (las celdas ámbar que quedaron por definir).",10,h=44); r+=1
r+=1
wr(r,"CÓMO SE CONECTA TODO (el sistema completo)",12,True,VERDE); r+=1
for t in [
 "El SCORECARD de 7 pasos genera el dato diario → los M1 lo suben en formato RAA al SYNC SEMANAL → el SISTEMA CONTROLADOR verifica esos resultados de forma independiente (Daily Flash del cierre, Weekly Scorecard, Monthly Review).",
 "La CLASIFICACIÓN de jugadores + la ESCALERA de reemplazo + la RUTA de gerentes alimentan el ORGANIGRAMA FUTURO (José y Jeffry hacia GM; Efraín de banca).",
 "La AUDITORÍA de presencia física define qué va OVERSEAS; el PLAYBOOK lo ejecuta (Vanessa ya es el caso probado).",
 "Los 5 MAPAS DE PROCESO documentan la operación; el CONTROLADOR verifica que de verdad ocurran.",
]:
    wr(r,t,10,bg=CREMA,h=40); r+=1

# ===== CHECKLIST =====
cs=wb.create_sheet("📋 Verificación"); cs.sheet_view.showGridLines=False
for col,w in [("A",4),("B",34),("C",34),("D",14),("E",34)]: cs.column_dimensions[col].width=w
cs.merge_cells("A1:E1")
t=cs.cell(1,1,"ROSANTA · Verificación de entregables — Curso 2"); t.font=F(13,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); cs.row_dimensions[1].height=24
cs.merge_cells("A2:E2")
cs.cell(2,1,"16 entregables, todos construidos. 'Pendiente' = lo que falta que TÚ definas (celdas ámbar de cada archivo).").font=F(9,italic=True,color="6B5A45")
for c,h in enumerate(["#","Entregable","Archivo","Estado","Pendiente / próximo paso"],start=1):
    cell=cs.cell(4,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
cs.row_dimensions[4].height=26
items=[
 ("Scorecard diario de 7 pasos","Rosanta_Scorecard_Diario_7_Pasos","Listo","Fijar metas y correrlo a diario 1 semana"),
 ("Organigrama actual + H.O.T. Canvas","Rosanta_Organigrama_HOT_Canvas","Listo","Nombres ya completos"),
 ("Clasificación de jugadores","Rosanta_Clasificacion_Jugadores","Listo","1 ofensivo (José) + 2 con potencial"),
 ("Escalera de reemplazo","Rosanta_Escalera_de_Reemplazo","Listo","Poner tus horas y costos reales"),
 ("Auditoría de presencia física","Rosanta_Auditoria_Presencia_Fisica","Listo","Sumar horas/sem; confirmar candidato #1"),
 ("Sistema Controlador (Panel 3 señales)","Rosanta_Sistema_Controlador_Panel","Listo","Definir metas y límites de alerta"),
 ("Playbook de contratación overseas","Rosanta_Playbook_Overseas","Listo","Salario, horas, plataforma, filtro"),
 ("Organigrama futuro (tiered)","Rosanta_Organigrama_Futuro","Listo","Costos y fechas de los 3 tiers"),
 ("Ritmos de reunión","Rosanta_Ritmos_de_Reunion","Listo","Fijar horarios de Huddle/Sync/Mensual"),
 ("Ruta de desarrollo de gerentes","Rosanta_Ruta_Desarrollo_Gerentes","Listo","Fechas y 1a conversación mensual"),
 ("Auditoría de mapas de proceso","Rosanta_Auditoria_Mapas_Proceso","Listo","Herramienta de planeación (corregida)"),
 ("Mapa: Apertura y Cierre","Rosanta_Apertura_Cierre","Listo","Teléfonos de servicio; límite de caja"),
 ("Mapa: Flujo de Servicio (FOH)","Rosanta_Flujo_Servicio_FOH","Listo","Política de propina de grupo; imprimir tarjeta"),
 ("Mapa: Flujo de Producción (BOH)","Rosanta_Flujo_Produccion_BOH","Listo","Pars; fotos de emplatado en el pase"),
 ("Mapa: Pipeline de Contratación","Rosanta_Pipeline_Contratacion","Listo","Confirmar palabra filtro"),
 ("Mapa: Mapeo Avanzado","Rosanta_Mapeo_Avanzado","Listo","Montos de autorización de cortesías"),
]
rr=5
for i,(ent,arch,est,pend) in enumerate(items,start=1):
    cs.cell(rr,1,i).font=F(9,True); cs.cell(rr,1).alignment=ctr; cs.cell(rr,1).border=border
    cs.cell(rr,2,ent).font=F(9,True,VERDE); cs.cell(rr,2).alignment=wrap; cs.cell(rr,2).border=border
    cs.cell(rr,3,arch).font=F(9); cs.cell(rr,3).alignment=wrap; cs.cell(rr,3).border=border
    stc=cs.cell(rr,4,est); stc.font=F(9,True); stc.alignment=ctr; stc.border=border; stc.fill=fill("E1F0E4")
    pc=cs.cell(rr,5,pend); pc.font=F(9); pc.alignment=wrap; pc.border=border; pc.fill=fill(AMBAR)
    cs.row_dimensions[rr].height=28; rr+=1
dv=DataValidation(type="list",formula1='"Listo,En ajuste,Pendiente"',allow_blank=True); cs.add_data_validation(dv); dv.add(f"D5:D{rr-1}")
rr+=1
cs.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
cs.cell(rr,1,"RESULTADO: los 16 entregables del Curso 2 están construidos y adaptados a Rosanta. Lo que sigue no es construir, es OPERAR: fijar los ámbar, correr el scorecard y las reuniones esta semana, y dejar que el sistema haga el trabajo que antes hacías tú.").font=F(10,True,VERDE)
cs.cell(rr,1).alignment=wrap; cs.row_dimensions[rr].height=44

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Capstone_Curso2.xlsx"
wb.save(out); print("guardado:",out)
