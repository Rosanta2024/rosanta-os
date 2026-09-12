# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Auditoría de Mapas de Proceso",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.12 · Convierte el conocimiento que vive en las personas en propiedad del negocio",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"MAPA DE PROCESO vs. SOP",12,True,COBRE); r+=1
wr(r,"El mapa de proceso es la vista de pájaro: QUÉ pasa y en qué orden (un flujo con puntos de decisión y responsables). El SOP es el detalle debajo: CÓMO se hace cada paso. El mapa da la foto grande; el SOP da la ejecución. Los dos son necesarios.",10,h=44); r+=1
r+=1
wr(r,"POR QUÉ IMPORTA",12,True,VERDE); r+=1
wr(r,"Si tu mejor cocinero renuncia, ¿cuánto tardas en entrenar al reemplazo? Si el conocimiento vive en las personas y no en un sistema, la respuesta es 'demasiado'. Matemática del ROI: 2-3 horas documentar un proceso una vez → corre 5-7 veces por semana, 52 semanas al año, con equipo rotativo. Sin mapa: preguntas, inconsistencia, baja calidad. Con mapa: repetible, entrenable, escalable.",10,h=58); r+=1
r+=1
wr(r,"CÓMO CONSTRUIRLO (y la prueba)",12,True,VERDE); r+=1
for t in [
 "Mapa: observa el proceso de punta a punta, identifica etapas y puntos de decisión, dibújalo (inicio → pasos → bifurcaciones → fin) y marca responsables.",
 "SOP: por cada etapa, pasos numerados (una acción por paso), criterios en las decisiones ('si X, haz Y'), y pruébalo con alguien que nunca lo ha hecho.",
 "LA PRUEBA es el paso clave: si alguien que nunca hizo la tarea sigue el mapa y el SOP y produce un resultado aceptable, sirve. Donde se atore, ahí falta detalle.",
]:
    wr(r,t,10,h=32); r+=1
r+=1
wr(r,"ERRORES COMUNES",12,True,VERDE); r+=1
for t in [
 "1. Hacerlo muy complicado: es una lista de verificación con contexto, no un manual. Si tardas más en leerlo que en hacerlo, simplifícalo.",
 "2. Escribirlo y no probarlo: un mapa sin probar es una adivinanza. Dáselo al más nuevo y observa dónde se traba.",
 "3. Meterlo en un fólder: los fólderes juntan polvo. Vive en la intranet, buscable y accesible desde cualquier dispositivo.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"NOTA ROSANTA",12,True,VERDE); r+=1
wr(r,"Ya tienes bastante documentado (Manual de Servicio, SOP Sala, Manual de Supervisores, recetario con fichas técnicas, y el pipeline de contratación que armamos). La auditoría de al lado marca qué existe hoy y qué falta convertir en mapa vivo. Confirma tú el estado Sí/No/Parcial de cada uno.",10,bg=CREMA,h=40)

# ===== AUDITORÍA =====
au=wb.create_sheet("📋 Auditoría de procesos"); au.sheet_view.showGridLines=False
for col,w in [("A",4),("B",26),("C",16),("D",12),("E",34),("F",34)]: au.column_dimensions[col].width=w
au.merge_cells("A1:F1")
t=au.cell(1,1,"ROSANTA · Auditoría de los 5 mapas de proceso"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); au.row_dimensions[1].height=24
au.merge_cells("A2:F2")
au.cell(2,1,"Los 5 mapas del curso ya están construidos (ver columna archivo). Urgencia 1-5 = referencia de impacto.").font=F(9,italic=True,color="6B5A45")
hdr=["#","Mapa de proceso","¿Documentado hoy?","Urgencia (1-5)","Qué existe en Rosanta hoy","Nota / mayor problema si es inconsistente"]
for c,h in enumerate(hdr,start=1):
    cell=au.cell(4,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
au.row_dimensions[4].height=42
rows=[
 ("Apertura y Cierre","Sí",5,"Construido: Rosanta_Apertura_Cierre","El de mayor tasa de error; ya documentado con puntos de decisión."),
 ("Flujo de Servicio (FOH)","Sí",3,"Construido: Rosanta_Flujo_Servicio_FOH (+ tarjeta y QR de reseñas)","Los 8 momentos + excepciones + tarjeta de bolsillo."),
 ("Flujo de Producción (BOH)","Sí",4,"Construido: Rosanta_Flujo_Produccion_BOH","Prep, mise en place, ejecución y control de calidad (en °C)."),
 ("Pipeline de Contratación (Hiring)","Sí",2,"Construido: Rosanta_Pipeline_Contratacion","5 etapas + no-negociables + entrevista + inducción."),
 ("Mapeo Avanzado (procesos únicos)","Sí",2,"Construido: Rosanta_Mapeo_Avanzado","Auditoría de preguntas + ejemplo (cortesías/descuentos). Es el 5º mapa del curso, no inventario."),
]
rr=5
for i,(mapa,doc,urg,exist,nota) in enumerate(rows,start=1):
    au.cell(rr,1,i).font=F(9,True); au.cell(rr,1).alignment=ctr; au.cell(rr,1).border=border
    au.cell(rr,2,mapa).font=F(9,True,VERDE); au.cell(rr,2).alignment=wrap; au.cell(rr,2).border=border
    dc=au.cell(rr,3,doc); dc.font=F(9,True); dc.alignment=ctr; dc.border=border
    dc.fill=fill("E1F0E4" if doc=="Sí" else AMBAR)
    uc=au.cell(rr,4,urg); uc.font=F(9,True); uc.alignment=ctr; uc.border=border; uc.fill=fill(AMBAR)
    au.cell(rr,5,exist).font=F(9); au.cell(rr,5).alignment=wrap; au.cell(rr,5).border=border
    au.cell(rr,6,nota).font=F(9); au.cell(rr,6).alignment=wrap; au.cell(rr,6).border=border
    au.row_dimensions[rr].height=46; rr+=1
dv_doc=DataValidation(type="list",formula1='"Sí,No,Parcial"',allow_blank=True); au.add_data_validation(dv_doc); dv_doc.add("C5:C9")
dv_u=DataValidation(type="list",formula1='"1,2,3,4,5"',allow_blank=True); au.add_data_validation(dv_u); dv_u.add("D5:D9")

rr+=1
au.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
au.cell(rr,1,"ORDEN DE PRIORIDAD SUGERIDO (construir primero lo de mayor urgencia y menor cobertura):").font=F(10,True,VERDE); au.row_dimensions[rr].height=18; rr+=1
prio=[
 ("OK","Los 5 mapas del curso (Apertura/Cierre, FOH, BOH, Hiring, Avanzado) ya están CONSTRUIDOS."),
 ("Nota","Inventario y Recepción NO es uno de los 5 mapas del curso, pero conviene documentarlo por el CMV: hoy vive en el Sistema Controlador (merma/CMV) y en el BOH (pesaje/waste log). Opcional como 6º mapa propio."),
 ("Sigue","Operar: probar cada mapa con la persona más nueva y subirlos a la intranet."),
]
for p,txt in prio:
    au.cell(rr,1,p).font=F(9,True,COBRE); au.cell(rr,1).alignment=ctr; au.cell(rr,1).border=border
    au.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=6)
    c=au.cell(rr,2,txt); c.font=F(9); c.alignment=wrap; c.border=border; c.fill=fill(CREMA)
    au.row_dimensions[rr].height=22; rr+=1

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Auditoria_Mapas_Proceso.xlsx"
wb.save(out); print("guardado:",out)
