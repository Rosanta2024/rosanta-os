# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Mapa de Apertura y Cierre",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.14 · El proceso #1 de tu auditoría",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"POR QUÉ IMPORTA",12,True,COBRE); r+=1
wr(r,"Apertura y cierre son los procesos con MÁS errores cuando no están documentados: pasan todos los días y con equipo rotativo. Los puntos de decisión importan tanto como los pasos: cuando algo sale mal, el mapa debe decir exactamente qué hacer.",10,h=44); r+=1
r+=1
wr(r,"CÓMO USARLO",12,True,VERDE); r+=1
for t in [
 "1. Ajusta las tareas, equipos y contactos a Rosanta (celdas ámbar = teléfonos y límites que defines tú).",
 "2. Pruébalo con tu persona más nueva: que siga el mapa de apertura mientras observas. Donde se atore, ahí falta detalle.",
 "3. Imprímelo y pégalo en la estación del jefe de turno; y súbelo a la intranet como documento vivo.",
]:
    wr(r,t,10,h=24); r+=1
r+=1
wr(r,"CONEXIÓN CON EL SISTEMA CONTROLADOR",12,True,VERDE); r+=1
wr(r,"El último paso del cierre — enviar el resumen al dueño/GM (ventas, comensales, anulaciones, incidencias) — es justo el Daily Flash del Sistema Controlador. Así el pulso diario queda registrado por una fuente que no es la persona medida.",10,bg=CREMA,h=40)

def build(sheetname,title,rows,decisions):
    s=wb.create_sheet(sheetname); s.sheet_view.showGridLines=False
    for col,w in [("A",5),("B",46),("C",26),("D",22),("E",30)]: s.column_dimensions[col].width=w
    s.merge_cells("A1:E1")
    t=s.cell(1,1,title); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); s.row_dimensions[1].height=24
    for c,h in enumerate(["#","Tarea","Responsable","Verificación","Notas"],start=1):
        cell=s.cell(3,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
    s.row_dimensions[3].height=18
    rr=4
    for i,(task,resp,ver) in enumerate(rows,start=1):
        s.cell(rr,1,i).font=F(9,True); s.cell(rr,1).alignment=ctr; s.cell(rr,1).border=border
        s.cell(rr,2,task).font=F(9); s.cell(rr,2).alignment=wrap; s.cell(rr,2).border=border
        s.cell(rr,3,resp).font=F(9,True,VERDE); s.cell(rr,3).alignment=wrap; s.cell(rr,3).border=border
        s.cell(rr,4,ver).font=F(9); s.cell(rr,4).alignment=wrap; s.cell(rr,4).border=border
        s.cell(rr,5).border=border; s.cell(rr,5).fill=fill(CREMA)
        s.row_dimensions[rr].height=28; rr+=1
    rr+=1
    s.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
    s.cell(rr,1,"PUNTOS DE DECISIÓN (qué hacer cuando algo sale mal)").font=F(10,True,COBRE); s.cell(rr,1).fill=fill(CREMA); s.row_dimensions[rr].height=18; rr+=1
    for c,h in enumerate(["","Situación","Acción","Contacto / límite",""],start=1):
        if h:
            cell=s.cell(rr,c,h); cell.font=F(9,True,MARRON); cell.border=border; cell.alignment=wrapc
    s.row_dimensions[rr].height=16; rr+=1
    for sit,act,contact in decisions:
        s.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=2)
        s.cell(rr,2,sit).font=F(9,True,VERDE); s.cell(rr,2).alignment=wrap; s.cell(rr,2).border=border
        s.cell(rr,3,act).font=F(9); s.cell(rr,3).alignment=wrap; s.cell(rr,3).border=border
        cc=s.cell(rr,4,contact); cc.font=F(9); cc.alignment=wrap; cc.border=border; cc.fill=fill(AMBAR)
        s.merge_cells(start_row=rr,start_column=4,end_row=rr,end_column=5)
        s.row_dimensions[rr].height=30; rr+=1
    return s

# APERTURA
open_rows=[
 ("Desarmar alarma, abrir, encender luces","Jefe de turno","Visual"),
 ("Ventilación/HVAC a temperatura; abrir jardín y patio","Jefe de turno","Termostato"),
 ("Recorrido: limpieza, mesas, baños, jardín, cava","Jefe de turno","Checklist"),
 ("POS (Posfile) encendido, conexión OK, menú del día cargado","José / mesero de turno","Transacción de prueba"),
 ("Revisar reservas y notas de grupos","Host / Sala","SonTickets (tablet/impreso)"),
 ("Barra: montaje, hielo, garnish, inventario rápido de licores","Coctelero","Visual / inventario"),
 ("Cocina: equipos encendidos, temperaturas verificadas, sacar mise en place","Cocinero de apertura","Log de temperaturas"),
 ("Revisar 86 (agotados) y actualizar POS","Jeffry / responsable de cocina","POS"),
 ("Pre-shift huddle (5 min): meta, foco, feature + maridaje","Jefe de turno","Formato del Huddle"),
 ("Música, señalización, chimenea/fogata (si aplica), ambiente","Sala","Visual"),
 ("Abrir puertas a la hora exacta","Jefe de turno","Reloj"),
]
open_dec=[
 ("Equipo no llega a temperatura en 15 min","Llamar a la empresa de servicio","Tel: [en Hub]"),
 ("POS no conecta","Reiniciar router; llamar a soporte Posfile","Tel: [en Hub]"),
 ("No-show de personal","Activar protocolo de ausencia (mínimo: 2 sala / 3 cocina entre semana)","[protocolo en Hub]"),
]
build("☀️ Apertura","ROSANTA · Apertura",open_rows,open_dec)

# CIERRE
close_rows=[
 ("Última hora de sentada / última orden","Jefe de turno","SonTickets / POS"),
 ("Side work de cierre por estación (sala, barra, cocina)","Todo el equipo","Checklist"),
 ("Cocina: desmontaje, etiquetar y guardar prep; limpieza","Cocina","Recorrido"),
 ("Cierre de barra: inventario de licores, limpieza","Coctelero","Checklist"),
 ("Reporte de fin de día del POS (Posfile)","Jefe de turno","Impreso"),
 ("Cuadre de caja (efectivo)","Jefe de turno","Log de caja"),
 ("Cierre de lote de tarjeta (BAC)","Jefe de turno","Reporte de lote"),
 ("Cuadre de propinas (passthrough, no es costo)","Jefe de turno","Registro de propinas"),
 ("Recorrido: equipos apagados, luces, chimenea apagada, baños, jardín","Jefe de turno","Físico"),
 ("Armar alarma y cerrar todas las puertas","Jefe de turno","Físico"),
 ("Enviar resumen de cierre al dueño/GM: ventas, comensales, anulaciones, incidencias","Jefe de turno","WhatsApp/correo → Daily Flash del Controlador"),
]
close_dec=[
 ("Varianza de caja mayor a Q[monto]","Documentar y avisar al dueño/GM","Límite: Q[monto]"),
 ("Falla de equipo","Registrar en el tracker de mantenimiento (Mario Valle)","[tracker en Hub]"),
 ("Riesgo de inocuidad","Descartar, documentar y avisar","[protocolo en Hub]"),
]
build("🌙 Cierre","ROSANTA · Cierre",close_rows,close_dec)

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Apertura_Cierre.xlsx"
wb.save(out); print("guardado:",out)
