# -*- coding: utf-8 -*-
import pandas as pd, numpy as np
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
SRC='/sessions/clever-hopeful-pascal/mnt/uploads/ReporteVentasProductos_8_11_2026_204344.xlsx'
OUT='/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Ingenieria_Menu_2026.xlsx'
df=pd.read_excel(SRC)
df.columns=[c.strip() for c in df.columns]
for col in ['Cantidad','Precio_Venta','Descuento','Precio_Compra','Total']:
    df[col]=pd.to_numeric(df[col],errors='coerce').fillna(0)
# quitar fila total (producto vacío)
df=df[df['Producto'].notna()].copy()
df['Prod']=df['Producto'].astype(str).str.split('|').str[0].str.strip()
df=df[df['Prod'].str.lower()!='nan']
df['fecha']=pd.to_datetime(df['fecha_Emision'],dayfirst=True,errors='coerce')
meses=max(1,(df['fecha'].max()-df['fecha'].min()).days/30.44)
tot=df['Total'].sum()

def grupo(cat):
    c=str(cat)
    if c in ['Fuerte','Entrante','Guarniciones','Postre','Brunch Dulce','Brunch Cocteleria']: return 'Comida'
    if c.startswith('Coctel') or c.startswith('Licor') or c in ['Cerveza','Espumante','Shot','Digestivo','Refresco & Agua','Café & Té','Promo Cocteles'] or c.startswith('Vino') or c.startswith('Vinos'): return 'Bebida'
    if c in ['Menú Especial','Eventos','Especial Día','Adicionales']: return 'Evento/Especial'
    return 'Otro'

g=df.groupby('Prod').agg(unid=('Cantidad','sum'),ingreso=('Total','sum'),
    cat=('Categoria',lambda s:s.mode().iat[0] if len(s.mode()) else ''),
    precio=('Precio_Venta',lambda s:s[s>0].median() if (s>0).any() else 0)).reset_index()
g['grupo']=g['cat'].map(grupo)
g['pct_ing']=g['ingreso']/tot*100
g['unid_mes']=g['unid']/meses
g=g.sort_values('ingreso',ascending=False).reset_index(drop=True)
g['cum']=g['ingreso'].cumsum()/tot*100

# clasificacion (solo items con ingreso>0), medianas de referencia
core=g[(g['ingreso']>0)&(g['grupo'].isin(['Comida','Bebida']))]
mu=core['unid'].median(); mp=core['precio'].median()
def clasi(r):
    if r['grupo'] not in ['Comida','Bebida']: return '—'
    if r['ingreso']<=0: return 'Sin venta'
    pop=r['unid']>=mu; car=r['precio']>=mp
    if pop and car: return 'Estrella'
    if pop and not car: return 'Popular (precio bajo)'
    if not pop and car: return 'Premium / nicho'
    return 'Baja rotación'
g['clase']=g.apply(clasi,axis=1)

# ---------- construir xlsx ----------
VERDE="4A6741";TIERRA="A0785A";COBRE="C9923F";MARRON="3D2B1F"
FV=PatternFill("solid",fgColor=VERDE);FT=PatternFill("solid",fgColor=TIERRA)
thin=Side(style="thin",color="D8CDB5");B=Border(left=thin,right=thin,top=thin,bottom=thin)
TITLE=Font(name="Arial",size=15,bold=True,color="FFFFFF");SUB=Font(name="Arial",size=10,italic=True,color="FFFFFF")
HW=Font(name="Arial",size=9,bold=True,color="FFFFFF");SEC=Font(name="Arial",size=10,bold=True,color=VERDE)
TXT=Font(name="Arial",size=9,color=MARRON);ctr=Alignment(horizontal="center",vertical="center")
CLR={'Estrella':'E7F0E1','Popular (precio bajo)':'EAF0DE','Premium / nicho':'F0E7D3','Baja rotación':'F0D6D3'}
wb=openpyxl.Workbook()

# Resumen
ws=wb.active; ws.title="Resumen"; ws.sheet_view.showGridLines=False
for col,w in zip("ABCD",[34,18,14,14]): ws.column_dimensions[col].width=w
ws.merge_cells("A1:D1"); ws["A1"].value="ROSANTA · Ingeniería de Menú 2026"; ws["A1"].font=TITLE; ws["A1"].fill=FV; ws["A1"].alignment=Alignment(vertical="center",indent=1)
ws.merge_cells("A2:D2"); ws["A2"].value=f"Ventas por producto · {df['fecha'].min().date()} a {df['fecha'].max().date()} · fuente: PosFile"; ws["A2"].font=SUB; ws["A2"].fill=FV; ws["A2"].alignment=Alignment(vertical="center",indent=1)
ws.row_dimensions[1].height=24;ws.row_dimensions[2].height=15
res=[("Ingreso total del período", f"Q {tot:,.0f}"),
 ("Tickets (transacciones)", f"{df['Doc ID'].nunique():,}"),
 ("Unidades vendidas", f"{int(df['Cantidad'].sum()):,}"),
 ("Ticket promedio", f"Q {tot/df['Doc ID'].nunique():,.0f}"),
 ("Productos distintos vendidos", f"{g['Prod'].nunique()}"),
 ("Top 10 productos = del ingreso", f"{g.head(10)['ingreso'].sum()/tot*100:.0f}%"),
 ("Top 20 productos = del ingreso", f"{g.head(20)['ingreso'].sum()/tot*100:.0f}%"),
 ("Productos que suman el 80% del ingreso", f"{int((g['cum']<=80).sum())+1} de {len(g)}"),
 ("Productos con < 1 venta/mes (cola larga)", f"{int((g['unid_mes']<1).sum())}"),
 ("Nota", "Sin costo cargado en el POS: falta el margen. Se agrega al activar costos/recetas."),
]
r=4
for k,v in res:
    ws.cell(row=r,column=1,value=k).font=SEC; ws.cell(row=r,column=1).border=B
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    ws.cell(row=r,column=2,value=v).font=TXT; ws.cell(row=r,column=2).border=B; ws.cell(row=r,column=2).alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[r].height=22; r+=1
# leyenda clases
r+=1; ws.cell(row=r,column=1,value="Clasificación (popularidad × precio, proxy sin margen):").font=SEC; r+=1
for name,desc in [("Estrella","alta rotación + precio alto → cuidar y destacar"),("Popular (precio bajo)","alta rotación + precio bajo → subir precio o el ticket"),("Premium / nicho","poca rotación + precio alto → mantener/impulsar"),("Baja rotación","poca rotación + precio bajo → revisar o quitar")]:
    ws.cell(row=r,column=1,value=name).font=Font(name="Arial",size=9,bold=True,color=MARRON); ws.cell(row=r,column=1).fill=PatternFill("solid",fgColor=CLR.get(name,"FFFFFF"))
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); ws.cell(row=r,column=2,value=desc).font=TXT; r+=1

def tabla(ws,cols,rows,start=4,colorcol=None):
    for i,(t,w) in enumerate(cols):
        ws.column_dimensions[chr(65+i)].width=w
        c=ws.cell(row=start,column=i+1,value=t);c.font=HW;c.fill=FT;c.alignment=ctr;c.border=B
    rr=start+1
    for row in rows:
        for i,val in enumerate(row):
            c=ws.cell(row=rr,column=i+1,value=val);c.border=B;c.font=TXT
            if isinstance(val,(int,float)) and i>0: c.alignment=ctr
        if colorcol is not None:
            cl=row[colorcol]; fill=CLR.get(cl)
            if fill:
                for i in range(len(row)): ws.cell(row=rr,column=i+1).fill=PatternFill("solid",fgColor=fill)
        rr+=1
    return ws

# Ingeniería (todos)
ws2=wb.create_sheet("Ingeniería de Menú"); ws2.sheet_view.showGridLines=False
ws2.merge_cells("A1:H1"); ws2["A1"].value="Ingeniería de Menú · por producto"; ws2["A1"].font=TITLE; ws2["A1"].fill=FV; ws2["A1"].alignment=Alignment(vertical="center",indent=1); ws2.row_dimensions[1].height=22
cols=[("Producto",34),("Grupo",14),("Categoría",16),("Unid",8),("Unid/mes",9),("Ingreso Q",12),("% ing",8),("Precio",9),("Clasificación",20)]
rows=[]
for _,x in g.iterrows():
    rows.append([x['Prod'],x['grupo'],str(x['cat']),int(x['unid']),round(x['unid_mes'],1),round(x['ingreso'],0),round(x['pct_ing'],1),round(x['precio'],0),x['clase']])
# header en fila 3
for i,(t,w) in enumerate(cols):
    ws2.column_dimensions[chr(65+i)].width=w
    c=ws2.cell(row=3,column=i+1,value=t);c.font=HW;c.fill=FT;c.alignment=ctr;c.border=B
rr=4
for row in rows:
    for i,val in enumerate(row):
        c=ws2.cell(row=rr,column=i+1,value=val);c.border=B;c.font=TXT
        if i>=3: c.alignment=ctr
    fill=CLR.get(row[8])
    if fill: ws2.cell(row=rr,column=9).fill=PatternFill("solid",fgColor=fill)
    rr+=1
ws2.freeze_panes="A4"

# Por categoría
ws3=wb.create_sheet("Por Categoría"); ws3.sheet_view.showGridLines=False
ws3.merge_cells("A1:D1"); ws3["A1"].value="Ingreso por categoría"; ws3["A1"].font=TITLE; ws3["A1"].fill=FV; ws3["A1"].alignment=Alignment(vertical="center",indent=1); ws3.row_dimensions[1].height=22
cc=df.groupby('Categoria').agg(ingreso=('Total','sum'),unid=('Cantidad','sum')).sort_values('ingreso',ascending=False).reset_index()
crows=[[r['Categoria'],int(r['unid']),round(r['ingreso'],0),round(r['ingreso']/tot*100,1)] for _,r in cc.iterrows()]
tabla(ws3,[("Categoría",26),("Unidades",12),("Ingreso Q",14),("% del ingreso",13)],crows,start=3)
ws3.cell(row=3,column=1);

# Candidatos a retirar (baja rotacion, comida/bebida, <1/mes)
ws4=wb.create_sheet("Candidatos a revisar"); ws4.sheet_view.showGridLines=False
ws4.merge_cells("A1:E1"); ws4["A1"].value="Candidatos a revisar / retirar (cola larga)"; ws4["A1"].font=TITLE; ws4["A1"].fill=FV; ws4["A1"].alignment=Alignment(vertical="center",indent=1); ws4.row_dimensions[1].height=22
ws4.merge_cells("A2:E2"); ws4["A2"].value="Comida y bebida con menos de 1 venta al mes en el período. Simplifican menú y catálogo del POS."; ws4["A2"].font=Font(name="Arial",size=9,italic=True,color="6B5A45"); ws4.row_dimensions[2].height=14
cand=g[(g['grupo'].isin(['Comida','Bebida']))&(g['unid_mes']<1)].sort_values('unid')
crows=[[x['Prod'],str(x['cat']),int(x['unid']),round(x['unid_mes'],2),round(x['ingreso'],0)] for _,x in cand.iterrows()]
tabla(ws4,[("Producto",34),("Categoría",18),("Unid (período)",13),("Unid/mes",10),("Ingreso Q",12)],crows,start=3)

wb.save(OUT)
print("Guardado:",OUT)
print("Ingreso real:",round(tot,0),"| items comida+bebida clasificados:",len(core),"| candidatos a revisar:",len(cand))
print("medianas ref → unid:",mu,"precio:",mp)
PY_MARKER=1
