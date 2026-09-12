# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Flujo de Servicio (FOH)",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.15 · 8 momentos que definen la experiencia del cliente",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"LA IDEA",12,True,COBRE); r+=1
wr(r,"Toda visita sigue la misma secuencia de 8 momentos. Cuando uno se salta o es inconsistente, la experiencia se rompe. La tarjeta de estándar de servicio (hoja 3) es el recordatorio de bolsillo para que cada mesero esté en la misma página.",10,h=44); r+=1
r+=1
wr(r,"VOZ DE MARCA EN EL SERVICIO",12,True,VERDE); r+=1
wr(r,"80% de tus clientes son turistas angloparlantes: saluda y recomienda en inglés o español según el cliente. El cliente es el protagonista (háblale de 'tu mesa', 'tu noche'). En el momento 3, recomienda el FEATURE del día con su MARIDAJE (gastrococtelería) — el mismo que se define en el Huddle. El upsell de postre/coctel/maridaje es una métrica de tu scorecard. Y no olvides el QR de reseñas: preséntalo con la cuenta e invita a escanearlo en la despedida cuando la visita fue buena — alimenta el panel-reseñas (Weekly Scorecard del Controlador).",10,h=72); r+=1
r+=1
wr(r,"PROTOCOLO DE QUEJAS: LAST",12,True,VERDE); r+=1
wr(r,"Escucha · Reconoce · Resuelve · Agradece. Escucha completo sin interrumpir, reconoce lo que sintió el cliente, resuelve de inmediato (rehacer/descuento según la política de cortesías) y agradece que te lo dijo. Cortesías se registran en POS (control del Sistema Controlador).",10,bg=CREMA,h=44); r+=1
r+=1
wr(r,"CÓMO USARLO",12,True,VERDE); r+=1
for t in [
 "1. Ajusta tiempos, estándares y guiones a Rosanta (celdas ámbar = política de propina de grupo y montos que defines tú).",
 "2. Imprime y lamina la Tarjeta de Estándar de Servicio para cada mesero.",
 "3. Entrena al equipo de sala con la tarjeta esta semana.",
]:
    wr(r,t,10,h=24); r+=1

# ===== 8 TOUCHPOINTS =====
fl=wb.create_sheet("📋 8 momentos"); fl.sheet_view.showGridLines=False
for col,w in [("A",4),("B",22),("C",34),("D",14),("E",40)]: fl.column_dimensions[col].width=w
fl.merge_cells("A1:E1")
t=fl.cell(1,1,"ROSANTA · Los 8 momentos del servicio"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); fl.row_dimensions[1].height=24
for c,h in enumerate(["#","Momento","Estándar Rosanta","Tiempo máx","Guion / acción"],start=1):
    cell=fl.cell(3,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
fl.row_dimensions[3].height=18
tps=[
 ("Bienvenida","Reconocido en 15 seg, sonrisa, contacto visual","15 seg","'Bienvenido a Rosanta / Welcome to Rosanta'"),
 ("Acomodo","Sentado en 2 min (o tiempo de espera exacto), menús + agua","2 min","Menús bilingües; ofrecer el mejor lugar (jardín/chimenea)"),
 ("Primer contacto","Mesero se presenta, recomienda el feature + su maridaje","3 min tras sentarse","'Hoy te recomiendo [plato] con [coctel] — maridan increíble'"),
 ("Orden","Orden exacta, confirmar modificaciones, sugerir upgrade","Cuando esté listo","Confirmar alergias; sugerir maridaje o entrada"),
 ("Entrega de alimentos","Platos correctos al asiento correcto, mesero confirma","Dentro del tiempo de pase","'Aquí está tu [plato], buen provecho'"),
 ("Check-back","'¿Cómo está todo?' Resolver de inmediato","2 min tras la comida","Si hay problema, aplicar LAST"),
 ("Cierre","Ofrecer postre o segundo coctel; cuenta pronta con el QR de reseñas visible","3 min tras retirar","'¿Te tentamos con un postre o un último coctel?' — entrega la cuenta con el QR de reseñas"),
 ("Despedida","'Gracias. Nos encantaría verte de nuevo.' + invitar a dejar reseña con el QR","2 min tras el pago","Despedir por su nombre; si la visita fue buena, invita a escanear el QR de reseñas (Google) que está en mesa/cuenta"),
]
rr=4
for i,(mom,std,tm,guion) in enumerate(tps,start=1):
    fl.cell(rr,1,i).font=F(9,True); fl.cell(rr,1).alignment=ctr; fl.cell(rr,1).border=border
    fl.cell(rr,2,mom).font=F(9,True,VERDE); fl.cell(rr,2).alignment=wrap; fl.cell(rr,2).border=border
    fl.cell(rr,3,std).font=F(9); fl.cell(rr,3).alignment=wrap; fl.cell(rr,3).border=border
    fl.cell(rr,4,tm).font=F(9); fl.cell(rr,4).alignment=ctr; fl.cell(rr,4).border=border
    c=fl.cell(rr,5,guion); c.font=F(9,italic=True); c.alignment=wrap; c.border=border; c.fill=fill(CREMA)
    fl.row_dimensions[rr].height=34; rr+=1
rr+=1
fl.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
fl.cell(rr,1,"MANEJO DE EXCEPCIONES").font=F(10,True,COBRE); fl.cell(rr,1).fill=fill(CREMA); fl.row_dimensions[rr].height=18; rr+=1
exc=[
 ("Tiempo de espera excedido","El jefe de turno se acerca, se disculpa y da un tiempo estimado. Si son 10+ min de más, ofrecer un coctel o entrada de cortesía (registrar en POS)."),
 ("Problema de calidad de alimento","Disculpa, retirar el plato, ofrecer rehacer o reemplazo. Cortesía mayor a Q[monto] la autoriza el jefe de turno."),
 ("Queja del cliente","Protocolo LAST: Escucha, Reconoce, Resuelve, Agradece."),
 ("Grupo grande (6+)","Política de propina/servicio [definir], mesero dedicado, revisión del jefe de turno a media comida."),
]
for sit,act in exc:
    fl.cell(rr,2,sit).font=F(9,True,VERDE); fl.cell(rr,2).alignment=wrap; fl.cell(rr,2).border=border
    fl.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=5)
    c=fl.cell(rr,3,act); c.font=F(9); c.alignment=wrap; c.border=border
    if '[monto]' in act or '[definir]' in act: c.fill=fill(AMBAR)
    fl.row_dimensions[rr].height=32; rr+=1

# ===== TARJETA =====
cd=wb.create_sheet("💳 Tarjeta de servicio"); cd.sheet_view.showGridLines=False
cd.column_dimensions["A"].width=3; cd.column_dimensions["B"].width=70
cd.merge_cells("B2:B2")
c=cd.cell(2,2,"ROSANTA · Tarjeta de Estándar de Servicio"); c.font=F(13,True,"FFFFFF"); c.fill=fill(VERDE); c.alignment=wrapc; cd.row_dimensions[2].height=26
cd.cell(3,2,"Imprime y lamina · una por mesero").font=F(9,italic=True,color="6B5A45")
card=[
 "1. Saluda en 15 seg — sonrisa y contacto visual",
 "2. Sienta en 2 min — agua + menús",
 "3. Primer contacto en 3 min — recomienda el feature + su maridaje",
 "4. Toma la orden — confirma modificaciones, sugiere upgrade",
 "5. Entrega los platos — confirma que estén correctos",
 "6. Check-back en 2 min — '¿cómo está todo?'",
 "7. Retira → ofrece postre/coctel → cuenta",
 "8. Agradece por su nombre → despedida → invita a escanear el QR de reseñas",
]
rr=5
for line in card:
    cc=cd.cell(rr,2,line); cc.font=F(11); cc.alignment=wrap; cc.fill=fill(CREMA); cc.border=border
    cd.row_dimensions[rr].height=26; rr+=1
cd.cell(rr+1,2,"Si algo sale mal → LAST: Escucha · Reconoce · Resuelve · Agradece.").font=F(10,True,COBRE)

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Flujo_Servicio_FOH.xlsx"
wb.save(out); print("guardado:",out)
