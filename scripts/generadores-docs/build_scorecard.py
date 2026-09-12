#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"
FONT="Arial"
def font(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=openpyxl.Workbook()
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=105

def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=font(sz,b,color,italic)
    c.alignment=Alignment(wrap_text=True,vertical="top")
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
    return c

r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Sistema de 7 Pasos — Scorecard Diario",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Herramienta de gestión para jefes de turno (M1) · Lección 2.5",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"QUÉ ES ESTO",12,True,VERDE); r+=1
wr(r,"Delegar sin sistema es solo esperar que alguien haga lo que tú hacías. Este scorecard es el motor que convierte la delegación en responsabilidad. Es una disciplina DIARIA (no semanal): cada día, para cada especialista, ~5–7 minutos por persona. Es la diferencia entre gestión profesional y adivinar.",10,h=58); r+=1
r+=1
wr(r,"QUIÉN LO LLENA",12,True,VERDE); r+=1
wr(r,"Tus M1 (jefes de turno / supervisores de Sala y de Cocina) lo llenan todos los días, un renglón por especialista de su equipo. Hay una hoja para SALA (meseros, coctelero, host) y otra para COCINA (línea, prep). Duplica la hoja del día para el día siguiente (clic derecho en la pestaña › Mover o copiar › Crear una copia) y cambia la fecha.",10,h=58); r+=1
r+=1
wr(r,"LA REGLA DE ORO: RAA — Resultado · Análisis · Acción",12,True,COBRE); r+=1
wr(r,"Primero el RESULTADO (qué pasó), luego el ANÁLISIS (por qué pasó), luego la ACCIÓN (qué se hará mañana). Ese orden no se negocia. Y lo más importante: el RESULTADO nunca lo reporta la persona que estás midiendo — debe venir de una fuente independiente (POS, hoja de reservas SonTickets, log de merma, reseñas), 100% verificable.",10,h=58); r+=1
r+=1
wr(r,"LOS 7 PASOS",12,True,VERDE); r+=1
pasos=[
"1. Cantidad — Compara las tareas asignadas ayer contra las realmente completadas. Califica: Excelente / Bueno / Malo / Terrible. Si no es Excelente, escribe 1–2 frases del porqué.",
"2. Calidad — Separa cantidad de calidad. Se pueden terminar todas las tareas pero mal hechas. Ejemplo: 'terminó todo (cantidad Excelente), pero el mise en place quedó desordenado y el corte disparejo (calidad Mala)'.",
"3. Horas (Prog vs Real) — Horas programadas vs horas realmente productivas. Anota desfases. Ejemplo: 'Programado 8h, trabajó 7h; 1h en el celular en la parte de atrás'.",
"4. Resultados Operativos — Resultados NO financieros del turno. Siempre Plan vs Real. Sala: rotación de mesa, precisión de orden, satisfacción, upsell de postre/coctel/maridaje, # mesas. Cocina: tiempo de salida del plato, % merma, mise en place, consistencia de receta/emplatado.",
"5. Resultados Financieros — Aporte financiero del turno. Plan vs Real. Sala: ventas del turno, ticket promedio, ventas de coctelería. Cocina: CMV del turno (obj. barra 20% / cocina 30%), merma en Q.",
"6. Categoría (1–5) — Con base en los pasos 1–5, clasifica al especialista para decidir la estrategia de gestión.",
"7. Plan de Acción para Mañana — El paso más importante. Traduce el análisis en un plan concreto y específico. Resultado → Análisis → Acción.",
]
for p in pasos: wr(r,p,10,h=44 if len(p)>150 else 30); r+=1
r+=1
wr(r,"LAS 5 CATEGORÍAS DE EMPLEADO",12,True,VERDE); r+=1
cats=[
("1","Buenos resultados + Trabajador — Tus jugadores A. Protégelos y retenlos.",VERDE),
("2","Buenos resultados + Flojo — Talento que no da el 100%. Déjalo que te traiga utilidad.","6B7A3A"),
("3","Malos resultados + Trabajador + Aprendiendo — Nuevo o sin experiencia con buena actitud. Invierte en capacitarlo.",COBRE),
("4","Malos resultados + Trabajador + No mejora — Se esfuerza pero no mejora pese al coaching. Déjalo ir.",TIERRA),
("5","Malos resultados + Flojo — Despide de inmediato.","9B3B2E"),
]
for num,txt,col in cats:
    c=ws.cell(row=r,column=2,value=f"{num} = {txt}"); c.font=font(10,True,"FFFFFF"); c.fill=fill(col)
    c.alignment=Alignment(wrap_text=True,vertical="center"); ws.cell(row=r,column=1).fill=fill(col)
    ws.row_dimensions[r].height=30; r+=1
r+=1
wr(r,"CÓMO INSTALAR EL SISTEMA (4 semanas)",12,True,VERDE); r+=1
sem=[
"Semana 1: Capacita a tus M1 en los 7 pasos. Recorre un ejemplo en vivo con un especialista real, juntos.",
"Semana 2: Los M1 corren el sistema solos. Tú revisas su documentación a diario y coacheas la calidad del análisis.",
"Semana 3: Los M1 presentan su categorización y planes de acción en el Sync Semanal. Coacheas patrones del equipo.",
"Semana 4+: Revisas semanal, no diario. El foco pasa a que los M1 mejoren en los pasos 6 y 7 (categorizar y planear).",
]
for s in sem: wr(r,s,10,h=30); r+=1
r+=1
wr(r,"NOTAS ROSANTA",12,True,VERDE); r+=1
notas=[
"• Staffing mínimo entre semana: 2 en Sala (coctelero + mesero) y 3 en Cocina por turno. Fin de semana: 3 Sala y 4 Cocina.",
"• Fuentes independientes de resultado: POS/ventas, hoja de reservas SonTickets, log de merma, reseñas Google, tablero maestro.",
"• Objetivos de costo (CMV): Barra 20% · Cocina 30%. Ticket promedio de referencia: ~Q259.",
"• Rosanta NO hace delivery: no apliquen métricas de reparto.",
"• Errores comunes: saltarse pasos, no documentar ('si no está escrito, no pasó'), fijarse solo en resultados e ignorar el esfuerzo, y tolerar Categoría 4 y 5.",
]
for n in notas: wr(r,n,10,bg=CREMA,h=28); r+=1

def build_scorecard(title,area,headers,example_row):
    s=wb.create_sheet(title); s.sheet_view.showGridLines=False
    ncols=len(headers); widths=[4,26,12,12,16,34,30,12,40]
    for i,w in enumerate(widths,start=1): s.column_dimensions[get_column_letter(i)].width=w
    s.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncols)
    t=s.cell(1,1,f"ROSANTA · Scorecard Diario — {area}"); t.font=font(15,True,"FFFFFF"); t.fill=fill(VERDE)
    t.alignment=Alignment(horizontal="left",vertical="center"); s.row_dimensions[1].height=26
    s.cell(3,1,"Fecha:").font=font(10,True,VERDE)
    s.cell(3,2).fill=fill("FFF9EE"); s.cell(3,2).border=border
    s.cell(3,4,"Jefe de turno (M1):").font=font(10,True,VERDE)
    s.merge_cells(start_row=3,start_column=5,end_row=3,end_column=6)
    s.cell(3,5).fill=fill("FFF9EE"); s.cell(3,5).border=border
    s.cell(3,8,"Turno:").font=font(10,True,VERDE)
    s.cell(3,9).fill=fill("FFF9EE"); s.cell(3,9).border=border
    hr=5
    for c,h in enumerate(headers,start=1):
        cell=s.cell(hr,c,h); cell.font=font(9,True,"FFFFFF"); cell.fill=fill(TIERRA)
        cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); cell.border=border
    s.row_dimensions[hr].height=40
    er=hr+1
    for c,v in enumerate(example_row,start=1):
        cell=s.cell(er,c,v); cell.font=font(9,italic=True,color="6B5A45"); cell.fill=fill(CREMA)
        cell.alignment=Alignment(wrap_text=True,vertical="top"); cell.border=border
    s.row_dimensions[er].height=44
    first=er+1; last=first+9
    for i,row in enumerate(range(first,last+1),start=1):
        s.cell(row,1,i).font=font(9,True); s.cell(row,1).alignment=Alignment(horizontal="center",vertical="center")
        for c in range(1,ncols+1):
            cell=s.cell(row,c); cell.border=border; cell.alignment=Alignment(wrap_text=True,vertical="top"); cell.font=font(9)
            if c==2: cell.fill=fill("FFFFFF")
        s.row_dimensions[row].height=34
    dv_cal=DataValidation(type="list",formula1='"Excelente,Bueno,Malo,Terrible"',allow_blank=True)
    dv_cal.prompt="Excelente / Bueno / Malo / Terrible"; s.add_data_validation(dv_cal)
    dv_cal.add(f"C{first}:C{last}"); dv_cal.add(f"D{first}:D{last}")
    dv_cat=DataValidation(type="list",formula1='"1,2,3,4,5"',allow_blank=True)
    dv_cat.prompt="1=A-Player · 2=Bueno/Flojo · 3=Aprendiendo · 4=No mejora · 5=Despedir"; s.add_data_validation(dv_cat)
    dv_cat.add(f"H{first}:H{last}")
    lr=last+2; s.merge_cells(start_row=lr,start_column=1,end_row=lr,end_column=ncols)
    leg=s.cell(lr,1,"Llena las celdas en blanco. Cantidad y Calidad tienen lista (Excelente/Bueno/Malo/Terrible). Categoría es lista 1–5. La fila crema es solo un ejemplo — no la borres, cópiala. Recuerda: el Resultado (pasos 4 y 5) viene de fuente independiente, nunca del empleado.")
    leg.font=font(8,italic=True,color="6B5A45"); leg.alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[lr].height=40
    return s

headers=["#","Especialista (rol)","1. Cantidad","2. Calidad","3. Horas (Prog vs Real)",
"4. Resultados Operativos (Plan vs Real)","5. Resultados Financieros (Plan vs Real)","6. Categoría (1–5)","7. Plan de Acción para Mañana"]
ej_sala=["Ej.","María — mesera","Bueno","Excelente","8 prog / 7 real",
"Rotación 62 min (meta 45) · precisión OK · upsell postre 1 de 8 mesas","Ventas turno Q980 (plan Q1,200) · ticket Q245","3",
"Acompaña 2h al mesero más rápido para técnicas de eficiencia y sugerencia de coctel."]
build_scorecard("📋 Scorecard SALA","SALA (FOH)",headers,ej_sala)
ej_cocina=["Ej.","Carlos — línea caliente","Excelente","Malo","9 prog / 9 real",
"Tiempo salida 22 min (meta 14) · merma 6% (meta 3%) · emplatado disparejo","CMV turno 34% (meta 30%) · merma Q180","3",
"Repasa fichas técnicas de emplatado del top 20; jefe verifica 3 platos al azar en el pase."]
build_scorecard("📋 Scorecard COCINA","COCINA (BOH)",headers,ej_cocina)

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Scorecard_Diario_7_Pasos.xlsx"
wb.save(out); print("guardado:",out)
