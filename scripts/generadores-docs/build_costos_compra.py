# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
OUT='/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Costos_Compra_Barra.xlsx'

VERDE="4A6741";TIERRA="A0785A";COBRE="C9923F";MARRON="3D2B1F";CREMA="F5EFE0"
FV=PatternFill("solid",fgColor=VERDE);FT=PatternFill("solid",fgColor=TIERRA)
FYEL=PatternFill("solid",fgColor="FFF3C4")
FCAR=PatternFill("solid",fgColor="E7F0E1")   # carta 2027
FBACK=PatternFill("solid",fgColor="EFE7DA")  # back bar
FING=PatternFill("solid",fgColor="F0E7D3")   # ingrediente coctel
thin=Side(style="thin",color="D8CDB5");B=Border(left=thin,right=thin,top=thin,bottom=thin)
TITLE=Font(name="Arial",size=14,bold=True,color="FFFFFF")
SUB=Font(name="Arial",size=9,italic=True,color="FFFFFF")
HW=Font(name="Arial",size=9,bold=True,color="FFFFFF")
SEC=Font(name="Arial",size=10,bold=True,color=VERDE)
TXT=Font(name="Arial",size=9,color=MARRON)
BOLD=Font(name="Arial",size=9,bold=True,color=MARRON)
ctr=Alignment(horizontal="center",vertical="center")
lft=Alignment(horizontal="left",vertical="center")
MQ='"Q"#,##0.00'

wb=openpyxl.Workbook()

# ---------- DESTILADOS Y LICORES ----------
# (Categoria, Producto, Distribuidor, Costo botella Q, ml, Clasificacion)
# Costo = lista "Actualizado"; ml del cierre; fallback cierre donde no hay actualizado.
C='Carta 2027'; BK='Back bar / extra'; IN='Ingrediente cóctel'
DEST=[
 ("Ron","Zacapa 23","Licorera",300.0,750,C),
 ("Ron","Zacapa XO","Licorera",827.0,750,C),
 ("Ron","Botran 12","Licorera",84.0,750,C),
 ("Ron","Botran Reserva Blanca (Blanco)","Licorera",84.0,750,C),
 ("Ron","Colonial","Licorera",150.0,750,C),
 ("Ron","Botran Oro","Licorera",75.0,1000,BK),
 ("Ron","Quezalteca","—",50.0,1000,BK),
 ("Whisky","Jameson","Tavito",188.0,750,C),
 ("Whisky","Old Parr 12","Marte",321.46,750,C),
 ("Whisky","Buchanan's 12","Marte",377.88,750,C),
 ("Whisky","Buchanan's 12 (litro)","Marte",487.26,1000,C),
 ("Whisky","Buchanan's 18","Marte",831.40,750,C),
 ("Whisky","Glenmorangie Original","Marte",452.76,700,C),
 ("Whisky","Glenlivet 12","Marcas Mundiales",750.0,750,C),
 ("Whisky","Macallan 12","Alcazaren",700.0,700,C),
 ("Whisky","Wild Turkey 101","Primium",195.0,750,C),
 ("Whisky","Wild Turkey 81","Primium",170.0,750,C),
 ("Whisky","Johnnie Walker Red Label","Marte",185.0,750,BK),
 ("Whisky","Johnnie Walker Black Label","Marte",318.93,750,BK),
 ("Whisky","Johnnie Walker Double Black","Marte",411.32,1000,BK),
 ("Whisky","Johnnie Walker Gold Label","Marte",570.05,750,BK),
 ("Whisky","Akashi","Noble Life",350.76,500,BK),
 ("Whisky","Pig Nose","Noble Life",143.76,750,BK),
 ("Whisky","Evan Williams","Alcazaren",124.31,750,BK),
 ("Vodka","Smirnoff","Marte",93.05,750,C),
 ("Vodka","Grey Goose","Marcas Mundiales",397.0,1000,C),
 ("Vodka","Ciroc","Marte",416.38,750,C),
 ("Vodka","Sky","—",85.0,750,BK),
 ("Gin","Bombay","(alta pendiente — aún no se compra)","",750,C),
 ("Gin","Hendricks","Marcas Mundiales",372.0,750,IN),
 ("Gin","Puerto de Indias","Noble Life",171.0,750,IN),
 ("Gin","Puerto de Indias Rose","Noble Life",143.35,750,IN),
 ("Gin","Gordons","Marte",72.5,700,IN),
 ("Gin","Gibson","—",91.75,700,BK),
 ("Gin","Xibal Akbal","—",155.0,700,BK),
 ("Tequila","1800 Cristalino","Tavito",305.0,700,C),
 ("Tequila","Jimador Blanco","Marte",134.36,750,C),
 ("Tequila","Don Julio Blanco","Marte",418.41,750,C),
 ("Tequila","Don Julio Reposado","Marte",528.24,750,C),
 ("Tequila","Don Julio 70","Marte",497.44,750,C),
 ("Mezcal","Mezcal Artesanal · Espadín","Ilegal",300.0,700,C),
 ("Mezcal","Mezcal Artesanal · Tobasiche","Ilegal",300.0,700,C),
 ("Mezcal","Mezcal Artesanal · Cupreata","Ilegal",300.0,700,C),
 ("Mezcal","Mezcal Verde","Primium",265.56,700,BK),
 ("Aperitivo/Digestivo","Grand Marnier","Primium",314.75,700,C),
 ("Aperitivo/Digestivo","Campari","Tavito",145.0,750,IN),
 ("Aperitivo/Digestivo","Aperol","Alcazaren",131.0,700,IN),
 ("Aperitivo/Digestivo","Cointreau","Marte",288.44,750,IN),
 ("Aperitivo/Digestivo","Licor 43","Tavito",214.0,750,IN),
 ("Aperitivo/Digestivo","Kahlúa","Tavito",145.0,1000,IN),
 ("Aperitivo/Digestivo","Amaretto Cremoso","Marte",304.54,700,IN),
 ("Aperitivo/Digestivo","Cynar","Alcazaren",166.07,700,IN),
 ("Aperitivo/Digestivo","Cinzano Rosso","Tavito",72.0,750,IN),
 ("Aperitivo/Digestivo","Cinzano Bianco","Tavito",72.0,750,IN),
 ("Aperitivo/Digestivo","Yzaguirre Rojo","Elite Marcas",122.0,1000,IN),
 ("Aperitivo/Digestivo","Yzaguirre Reserva","Elite Marcas",150.0,1000,IN),
 ("Aperitivo/Digestivo","Villa Cardea","Primium",84.0,700,IN),
 ("Aperitivo/Digestivo","Zoco (Pacharán)","Alcazaren",202.5,700,IN),
 ("Aperitivo/Digestivo","Licor de Saúco (Elderflower)","Alcazaren",78.33,700,IN),
]

ws=wb.active; ws.title="Destilados y Licores"; ws.sheet_view.showGridLines=False
ws.merge_cells("A1:H1"); ws["A1"]="ROSANTA · Costos de compra — Destilados y licores"
ws["A1"].font=TITLE; ws["A1"].fill=FV; ws["A1"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[1].height=24
ws.merge_cells("A2:H2"); ws["A2"]="Precio de compra por botella (lista actualizada de barra). Costo por onza y por trago para cargar al POS."
ws["A2"].font=SUB; ws["A2"].fill=FV; ws["A2"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[2].height=14
# assumption cell
ws["F3"]="Onzas por trago:"; ws["F3"].font=BOLD; ws["F3"].alignment=Alignment(horizontal="right",vertical="center")
ws["G3"]=1.5; ws["G3"].fill=FYEL; ws["G3"].font=BOLD; ws["G3"].border=B; ws["G3"].alignment=ctr
ws["H3"]="⟵ editable"; ws["H3"].font=Font(name="Arial",size=8,italic=True,color="8A7A5C")
cols=[("Categoría",16),("Producto",34),("Distribuidor",18),("Costo botella",13),("ml",7),("Costo / onza",12),("Costo / trago",12),("Clasificación",18)]
hr=4
for i,(t,w) in enumerate(cols):
    ws.column_dimensions[chr(65+i)].width=w
    c=ws.cell(row=hr,column=i+1,value=t); c.font=HW; c.fill=FT; c.alignment=ctr; c.border=B
r=hr+1
for cat,prod,dist,costo,ml,clas in DEST:
    ws.cell(row=r,column=1,value=cat).border=B; ws.cell(row=r,column=1).font=TXT
    ws.cell(row=r,column=2,value=prod).border=B; ws.cell(row=r,column=2).font=TXT
    ws.cell(row=r,column=3,value=dist).border=B; ws.cell(row=r,column=3).font=TXT
    cc=ws.cell(row=r,column=4,value=(costo if costo!="" else None)); cc.border=B; cc.font=TXT; cc.number_format=MQ; cc.alignment=ctr
    cm=ws.cell(row=r,column=5,value=ml); cm.border=B; cm.font=TXT; cm.alignment=ctr
    o=ws.cell(row=r,column=6); o.border=B; o.font=TXT; o.number_format=MQ; o.alignment=ctr
    tg=ws.cell(row=r,column=7); tg.border=B; tg.font=TXT; tg.number_format=MQ; tg.alignment=ctr
    if costo!="":
        o.value=f"=D{r}/(E{r}/29.5735)"
        tg.value=f"=F{r}*$G$3"
    cl=ws.cell(row=r,column=8,value=clas); cl.border=B; cl.font=TXT; cl.alignment=ctr
    fill={C:FCAR,BK:FBACK,IN:FING}.get(clas)
    if fill: cl.fill=fill
    r+=1
ws.freeze_panes="A5"
# leyenda
r+=1
ws.cell(row=r,column=1,value="Leyenda:").font=SEC
for name,fill,desc in [(C,FCAR,"va en la carta 2027 (cárgalo con costo)"),(IN,FING,"no se vende solo; es ingrediente de coctel"),(BK,FBACK,"botella extra / back bar, fuera de carta")]:
    r+=1
    a=ws.cell(row=r,column=1,value=name); a.font=BOLD; a.fill=fill; a.border=B
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=5)
    ws.cell(row=r,column=2,value=desc).font=TXT
r+=2
ws.cell(row=r,column=1,value="Nota: 'Costo/trago' = costo/onza × onzas por trago (celda G3, editable). Bombay aún no se compra (alta pendiente). Fuente: Inventario Barra 2026, hoja 'Dest. lic. vinos Actualizados' + Cierre Julio 03.08.26.").font=Font(name="Arial",size=8,italic=True,color="6B5A45")
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)

# ---------- VINOS ----------
VIN=[
 ("Tinto","La Linda","Malbec","Cofradía",115.0),
 ("Tinto","LAN Crianza","Blend","Cofradía",160.0),
 ("Tinto","LAN Reserva","Blend","Cofradía",245.0),
 ("Tinto","Pulenta","Red Blend","Cofradía",275.0),
 ("Tinto","Mariluna","Tempranillo Bobal","Elite Marcas",110.0),
 ("Tinto","Bohigas Garnatxa","Garnatxa","Elite Marcas",132.0),
 ("Tinto","Acustic","Cariñena Garnatxa","Elite Marcas",180.0),
 ("Tinto","Almodí","Garnatxa Peluda","Elite Marcas",142.0),
 ("Tinto","Fable","Syrah","La Clack",170.0),
 ("Tinto","Tous Cousins","Syrah","La Clack",160.0),
 ("Tinto","Nashroom","Lemberger","Franconia",156.0),
 ("Tinto","Feldhase","Trollinger","Franconia",128.0),
 ("Blanco","Mariluna","Verdejo Macabeo","Elite Marcas",103.5),
 ("Blanco","La Val","Albariño","Elite Marcas",143.0),
 ("Blanco","Killka","Chardonnay","Cofradía",155.0),
 ("Blanco","LAN Verdejo","Verdejo","Cofradía",165.0),
 ("Blanco","Santiago Ruiz","Rías Baixas","Cofradía",250.0),
 ("Blanco","Txakoli","Hondarrabi Zuri","Elite Marcas",147.0),
 ("Blanco","Caprice","Viognier Grenache","La Clack",185.0),
 ("Blanco","Ilercavonia","Garnatxa Blanca","Elite Marcas",142.0),
 ("Blanco","Julian Hart","Riesling","Franconia",180.0),
 ("Blanco","Turk","Grüner Veltliner","Franconia",132.0),
 ("Blanco","Pulpe","Blend","La Clack",170.0),
 ("Rosado","La Flor","Malbec","Cofradía",155.0),
 ("Rosado","Pasión","Bobal","Elite Marcas",98.5),
 ("Cava","Bohigas","Xarel·lo","Elite Marcas",126.0),
]
ws=wb.create_sheet("Vinos"); ws.sheet_view.showGridLines=False
ws.merge_cells("A1:F1"); ws["A1"]="ROSANTA · Costos de compra — Vinos"; ws["A1"].font=TITLE; ws["A1"].fill=FV; ws["A1"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[1].height=24
ws.merge_cells("A2:F2"); ws["A2"]="Precio de compra por botella (750 ml). Costo por copa para cargar al POS."; ws["A2"].font=SUB; ws["A2"].fill=FV; ws["A2"].alignment=Alignment(vertical="center",indent=1)
ws["D3"]="Copas por botella:"; ws["D3"].font=BOLD; ws["D3"].alignment=Alignment(horizontal="right",vertical="center")
ws["E3"]=5; ws["E3"].fill=FYEL; ws["E3"].font=BOLD; ws["E3"].border=B; ws["E3"].alignment=ctr
ws["F3"]="⟵ editable"; ws["F3"].font=Font(name="Arial",size=8,italic=True,color="8A7A5C")
cols=[("Tipo",12),("Casa",22),("Cepa",22),("Distribuidor",16),("Costo botella",13),("Costo / copa",12)]
for i,(t,w) in enumerate(cols):
    ws.column_dimensions[chr(65+i)].width=w
    c=ws.cell(row=4,column=i+1,value=t); c.font=HW; c.fill=FT; c.alignment=ctr; c.border=B
r=5
for tipo,casa,cepa,dist,costo in VIN:
    ws.cell(row=r,column=1,value=tipo).border=B; ws.cell(row=r,column=1).font=TXT
    ws.cell(row=r,column=2,value=casa).border=B; ws.cell(row=r,column=2).font=TXT
    ws.cell(row=r,column=3,value=cepa).border=B; ws.cell(row=r,column=3).font=TXT
    ws.cell(row=r,column=4,value=dist).border=B; ws.cell(row=r,column=4).font=TXT
    cc=ws.cell(row=r,column=5,value=costo); cc.border=B; cc.font=TXT; cc.number_format=MQ; cc.alignment=ctr
    cp=ws.cell(row=r,column=6,value=f"=E{r}/$E$3"); cp.border=B; cp.font=TXT; cp.number_format=MQ; cp.alignment=ctr
    r+=1
ws.freeze_panes="A5"

# ---------- CERVEZAS Y BEBIDAS ----------
CER=[
 ("Cerveza","Gallo (bot. retornable)","La Nueva S.A.",24,155.0),
 ("Cerveza","Gallo lata 12 oz","La Nueva S.A.",24,188.0),
 ("Cerveza","Cabro Reserva (desechable)","La Nueva S.A.",24,245.70),
 ("Cerveza","Monte Carlo","La Nueva S.A.",24,245.0),
 ("Cerveza","Antigua Sin Novia","Industrias VSG",24,450.0),
 ("Cerveza","Antigua Muy Noble","Industrias VSG",24,450.0),
 ("Cerveza artesanal","Stangen","—",6,107.5),
 ("Cerveza artesanal","Weissbier","—",6,107.5),
 ("Cerveza artesanal","IPA","—",6,107.5),
 ("Cerveza artesanal","Golden Ale","—",6,107.5),
 ("Bebida","Coca-Cola Regular","—",24,130.0),
 ("Bebida","Coca-Cola Cero","—",24,90.0),
 ("Bebida","Club Soda","—",24,85.0),
 ("Bebida","Tónica","—",24,85.0),
]
ws=wb.create_sheet("Cervezas y bebidas"); ws.sheet_view.showGridLines=False
ws.merge_cells("A1:E1"); ws["A1"]="ROSANTA · Costos de compra — Cervezas y bebidas"; ws["A1"].font=TITLE; ws["A1"].fill=FV; ws["A1"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[1].height=24
ws.merge_cells("A2:E2"); ws["A2"]="Precio por caja y costo por unidad para cargar al POS."; ws["A2"].font=SUB; ws["A2"].fill=FV; ws["A2"].alignment=Alignment(vertical="center",indent=1)
cols=[("Categoría",16),("Producto",26),("Distribuidor",16),("Unid / caja",10),("Costo caja",12)]
for i,(t,w) in enumerate(cols):
    ws.column_dimensions[chr(65+i)].width=w
    c=ws.cell(row=4,column=i+1,value=t); c.font=HW; c.fill=FT; c.alignment=ctr; c.border=B
ws.column_dimensions['F'].width=12
c=ws.cell(row=4,column=6,value="Costo unidad"); c.font=HW; c.fill=FT; c.alignment=ctr; c.border=B
r=5
for cat,prod,dist,un,caja in CER:
    ws.cell(row=r,column=1,value=cat).border=B; ws.cell(row=r,column=1).font=TXT
    ws.cell(row=r,column=2,value=prod).border=B; ws.cell(row=r,column=2).font=TXT
    ws.cell(row=r,column=3,value=dist).border=B; ws.cell(row=r,column=3).font=TXT
    cu=ws.cell(row=r,column=4,value=un); cu.border=B; cu.font=TXT; cu.alignment=ctr
    cc=ws.cell(row=r,column=5,value=caja); cc.border=B; cc.font=TXT; cc.number_format=MQ; cc.alignment=ctr
    cx=ws.cell(row=r,column=6,value=f"=E{r}/D{r}"); cx.border=B; cx.font=TXT; cx.number_format=MQ; cx.alignment=ctr
    r+=1
ws.freeze_panes="A5"

wb.save(OUT)
print("Guardado:",OUT)
