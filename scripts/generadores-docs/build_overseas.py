# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center")
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Playbook de Contratación en el Extranjero",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.7 · Candidato #1: Asistente de Marketing y Contenido (remoto)",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"YA TIENES OVERSEAS FUNCIONANDO",12,True,COBRE); r+=1
wr(r,"Vanessa (Pauta) trabaja desde Colombia. El modelo ya te funciona: talento de LatAm, mismo idioma, casi la misma zona horaria (Colombia GMT-5 vs Guatemala GMT-6, 1h de diferencia). Este playbook formaliza ese modelo para tu siguiente contratación, que trabajaría junto a Vanessa. Ella incluso puede ayudarte a filtrar y entrevistar, porque conoce el trabajo.",10,h=58); r+=1
r+=1
wr(r,"LAS 5 FASES",12,True,VERDE); r+=1
for t in [
 "Fase 1 · DEFINIR — Escribe el Role Scorecard: misión, resultados a 90 días, competencias, horas y compensación.",
 "Fase 2 · PUBLICAR — Redacta la vacante con un filtro de atención al detalle. Descarta a quien no lo cumpla (elimina 60-70% de bajo esfuerzo).",
 "Fase 3 · FILTRAR — Test task PAGADO ($20-30) de 1-2 horas que imite el trabajo real. Filtra a los no serios y respeta su tiempo.",
 "Fase 4 · ENTREVISTA — Videollamada de 30 min con cámara. Conversación natural, repasar cómo pensó el test task, fijar expectativas.",
 "Fase 5 · ONBOARDING — Rampa de 30 días con videos Loom. Semana 1 fundación, 2 supervisado, 3 graduado, 4 full.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"PLATAFORMAS (recomendado para LatAm)",12,True,VERDE); r+=1
for t in [
 "• Workana — LatAm (ideal: español nativo y misma zona horaria que Vanessa).",
 "• OnlineJobs.ph — Filipinas (VAs, contabilidad; más barato pero inglés y zona horaria lejana).",
 "• Upwork — global (bueno para el test task por proyecto). · LinkedIn — perfiles con experiencia.",
]:
    wr(r,t,10,h=24); r+=1
r+=1
wr(r,"RANGOS DE COSTO (referencia del curso)",12,True,VERDE); r+=1
for t in [
 "• Asistente de marketing / redes: $500-1,200/mes.  • Asistente virtual / admin: $400-800/mes.  • Verificación financiera: $500-1,500/mes.",
 "• Métodos de pago: Wise, PayPal, Payoneer o la propia plataforma.",
]:
    wr(r,t,10,bg=CREMA,h=24); r+=1
r+=1
wr(r,"ERRORES COMUNES",12,True,VERDE); r+=1
for t in [
 "1. Saltarse el test task pagado: es el mejor predictor del desempeño real (la entrevista revela personalidad; el test task, capacidad).",
 "2. No grabar Looms: son tu infraestructura de entrenamiento. Grábalos una vez y todo futuro hire usa los mismos.",
 "3. Microgestionar después de la semana 2: la rampa de 30 días existe para soltar de forma progresiva.",
]:
    wr(r,t,10,h=30); r+=1

# ===== PLAYBOOK =====
pb=wb.create_sheet("📋 Playbook"); pb.sheet_view.showGridLines=False
pb.column_dimensions["A"].width=3
pb.column_dimensions["B"].width=30
pb.column_dimensions["C"].width=78
def sec(title):
    global rr
    pb.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=3)
    c=pb.cell(rr,1,title); c.font=F(12,True,"FFFFFF"); c.fill=fill(VERDE); c.alignment=Alignment(vertical="center")
    pb.row_dimensions[rr].height=22; rr+=1
def kv(k,v,amber=False,h=None):
    global rr
    a=pb.cell(rr,2,k); a.font=F(9,True,VERDE); a.alignment=wrap; a.border=border
    b=pb.cell(rr,3,v); b.font=F(9); b.alignment=wrap; b.border=border
    if amber: b.fill=fill(AMBAR)
    else: b.fill=fill(CREMA)
    pb.row_dimensions[rr].height=h or 26; rr+=1
def note(v,h=None):
    global rr
    pb.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=3)
    c=pb.cell(rr,2,v); c.font=F(9,italic=True,color="6B5A45"); c.alignment=wrap
    pb.row_dimensions[rr].height=h or 22; rr+=1
def blank():
    global rr; rr+=1

pb.merge_cells("A1:C1")
t=pb.cell(1,1,"ROSANTA · Playbook — Asistente de Marketing y Contenido (remoto)"); t.font=F(13,True,"FFFFFF"); t.fill=fill(TIERRA); t.alignment=Alignment(vertical="center"); pb.row_dimensions[1].height=24
rr=3

sec("FASE 1 · ROLE SCORECARD")
kv("Rol","Asistente de Marketing y Contenido (remoto, LatAm)")
kv("Reporta a","Vanessa Wilches (Pauta) / Dirección")
kv("Misión (1 frase)","Ejecutar el calendario de contenido y apoyar la pauta para multiplicar alcance y reservas, de modo que Vanessa y Daniel se enfoquen en estrategia y producción.",h=40)
kv("Horas/semana","20-30 h (empezar part-time)",amber=True)
kv("Horario","Solapamiento con hora Guatemala (GMT-6); Colombia es GMT-5, 1h de diferencia. Franja núcleo 9:00-13:00 GT.",h=32)
kv("Compensación","$600-800/mes para empezar (rango $500-1,200)",amber=True)
note("Outcomes — resultados medibles esperados a 90 días:")
kv("Outcome 1","Calendario de contenido publicado al 100% cada semana, sin huecos (definir # de piezas/semana).",amber=True,h=30)
kv("Outcome 2","Respuesta a comentarios y DMs en menos de 4 h dentro del horario.")
kv("Outcome 3","Hook rate / engagement sostenido o al alza vs. baseline (apoyo a Vanessa).",h=30)
kv("Outcome 4","Banco de captions y hashtags de marca creado y mantenido, reutilizable.")
note("Competencias — no-negociables:")
kv("Competencia 1","Español nativo + inglés funcional (80% de clientes son turistas angloparlantes).",h=30)
kv("Competencia 2","Manejo de Meta Business Suite y programador de publicaciones.")
kv("Competencia 3","Redacción para redes (captions con gancho) y edición básica (Canva / CapCut).",h=30)
kv("Competencia 4","Atención al detalle y respeto a la voz de marca (cliente protagonista, 2a persona, sin 'nosotros').",h=30)
blank()

sec("FASE 2 · JOB LISTING (vacante)")
kv("Título","Asistente de Marketing y Contenido (remoto) — Restaurante en Antigua Guatemala")
kv("Sobre Rosanta","Rosanta es un restaurante de gastrococtelería en Antigua Guatemala: cocina natural de temporada maridada con coctelería moderna. Equipo pequeño y remoto-amigable (ya trabajamos con talento en Colombia).",h=40)
kv("Descripción del rol","Ejecutar el calendario de contenido, redactar captions bilingües, programar publicaciones, responder la comunidad y apoyar los reportes de pauta.",h=40)
kv("Requisitos","Español nativo + inglés; experiencia en redes de marca; Meta Business Suite; Canva/CapCut; buena redacción.",h=30)
kv("Compensación + horas","$600-800/mes · 20-30 h/sem · solapamiento con hora Guatemala",amber=True)
kv("Filtro de atención al detalle","En la PRIMERA línea de tu aplicación escribe exactamente: 'Cocino con carisma'. Descarta toda aplicación que no lo incluya.",h=32)
kv("Plataforma(s)","Workana (LatAm, recomendada) · alternativas: Upwork, LinkedIn",amber=True)
blank()

sec("FASE 3 · TEST TASK PAGADO ($20-30)")
kv("Tarea","Redacta una semana de contenido para Rosanta: 5 piezas con gancho + caption + hashtags (bilingüe donde aplique), alineadas a la voz de marca. Entrega un doc + 1 pieza montada en Canva.",h=44)
kv("Pago","$25 por la tarea (respeta su tiempo y filtra a los no serios)",amber=True)
kv("Qué evaluar","Gancho, voz de marca (cliente protagonista, 2a persona, sin 'nosotros', sin em dashes), ortografía, y si sigue TODAS las instrucciones.",h=32)
blank()

sec("FASE 4 · ENTREVISTA (30 min, video, cámara)")
kv("Paso 1","5 min de conversación natural (verificar fluidez de español e inglés).")
kv("Paso 2","Repasar cómo pensó el test task (el proceso, no solo el resultado).")
kv("Paso 3","Fijar expectativas: horas, pago, cadencia de comunicación, herramientas.")
note("Tip: Vanessa puede acompañar esta entrevista; conoce el trabajo y detecta nivel real.")
blank()

sec("FASE 5 · ONBOARDING · rampa de 30 días")
kv("Semana 1 · Fundación","Configurar accesos por videollamada. Grabar Looms de cada proceso recurrente (calendario, programación, respuesta a comunidad, reporte). La persona observa y replica.",h=40)
kv("Semana 2 · Supervisado","Ejecuta sola. Revisas el 100% de su trabajo a diario. Check-in de 15 min.",h=26)
kv("Semana 3 · Graduado","Ejecución independiente. Revisas el 50% al azar. Check-in día por medio.")
kv("Semana 4 · Full","Es dueña del trabajo. Revisión semanal.")
blank()

sec("COMUNICACIÓN Y PAGO")
kv("Update diario (async)","Máx 5 bullets: qué hizo, qué queda pendiente, bloqueos.")
kv("Sync semanal","Videollamada 15-30 min: revisar trabajo y prioridades de la semana.")
kv("Protocolo de urgencia","Definir desde el inicio qué cuenta como urgente y cómo contactarte.",amber=True)
kv("Método de pago","Wise, PayPal, Payoneer o la plataforma (mismo método que ya usas con Vanessa).",amber=True)

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Playbook_Overseas.xlsx"
wb.save(out); print("guardado:",out)
