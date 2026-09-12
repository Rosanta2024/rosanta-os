# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Ruta de Desarrollo de Gerentes",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.10 · No puedes contratar para salir de un problema de gestión: hay que desarrollar",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"LAS 4 ETAPAS",12,True,COBRE); r+=1
etapas=[
 ("Etapa 1 · Miembro de equipo","Señales para promover: llega puntual (90+ días), hace su side work sin que se lo pidan, recibe feedback sin ponerse a la defensiva, los compañeros le preguntan a él."),
 ("Etapa 2 · Líder de turno (M1) · desarrollo 90 días","Aprende: formato del Huddle, resolución básica de conflictos, leer el ritmo del turno, decisiones a mitad de turno, una métrica financiera simple (# comensales o ventas por hora). Señales: resuelve 80%+ sin escalar, otros siguen su dirección, se hace dueño del resultado del turno."),
 ("Etapa 3 · Jefe de departamento (M2) · desarrollo 90-180 días","Aprende: el Sistema de 7 Pasos (scorecard diario de cada especialista), P&L de su departamento, horarios y manejo de mano de obra, contratar y entrevistar, conversaciones de desempeño. Señales: cumple los KPI de su depto, entrenó a alguien que lo cubra, trae soluciones al Sync, no problemas."),
 ("Etapa 4 · Gerente General (M3) · desarrollo 6-12 meses","Aprende: propiedad total del P&L, desarrollo de equipo (coachear a otros gerentes), relaciones con proveedores, pensamiento estratégico, listo para expansión. Señales: corre su área sin el dueño 30+ días seguidos, pasa la auditoría del Sistema Controlador solo, gestiona gerentes, no solo miembros."),
]
for tit,desc in etapas:
    wr(r,tit,11,True,VERDE); r+=1
    wr(r,desc,10,h=44); r+=1
r+=1
wr(r,"LA CONVERSACIÓN MENSUAL DE DESARROLLO (20 min)",12,True,VERDE); r+=1
for t in [
 "1. 'Aquí estás en la ruta' — muéstrale el mapa de etapas.",
 "2. 'Esto vi este mes' — una fortaleza y un área a desarrollar.",
 "3. 'La siguiente habilidad a enfocar' — una capacidad concreta de la etapa siguiente.",
 "4. '¿Qué apoyo necesitas de mí?'",
]:
    wr(r,t,10,bg=CREMA,h=22); r+=1
wr(r,"20 minutos, una vez al mes. Es la herramienta de gestión más barata y de mayor ROI que tienes.",10,italic=True,h=22); r+=1
r+=1
wr(r,"ERRORES COMUNES",12,True,VERDE); r+=1
for t in [
 "1. Saltarse etapas: promover a Jefe de depto sin pasar por Líder de turno = se ahoga gestionando un departamento.",
 "2. No hacer visible la ruta: si el equipo no ve el camino, no sabe que existe. Compártelo abiertamente.",
 "3. Desarrollar solo a los ofensivos: los defensivos pueden crecer a líderes de turno confiables e incluso jefes de departamento con el apoyo correcto. Justo por eso Jeffry (defensivo con potencial) va en esta ruta.",
]:
    wr(r,t,10,h=30); r+=1

# ===== PLANES =====
pl=wb.create_sheet("📋 Planes de desarrollo"); pl.sheet_view.showGridLines=False
for col,w in [("A",5),("B",34),("C",30),("D",24),("E",30),("F",14)]: pl.column_dimensions[col].width=w
rr=1
STAGES='"1. Miembro de equipo,2. Líder de turno,3. Jefe de departamento,4. Gerente General"'
STATUS='"No iniciado,En progreso,Logrado"'
dv_stage=DataValidation(type="list",formula1=STAGES,allow_blank=True); pl.add_data_validation(dv_stage)
dv_st=DataValidation(type="list",formula1=STATUS,allow_blank=True); pl.add_data_validation(dv_st)

def person_header(title):
    global rr
    pl.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
    c=pl.cell(rr,1,title); c.font=F(13,True,"FFFFFF"); c.fill=fill(VERDE); c.alignment=Alignment(vertical="center"); pl.row_dimensions[rr].height=24; rr+=1
def kv2(k1,v1,k2,v2,amber=True):
    global rr
    pl.cell(rr,1,k1).font=F(9,True,VERDE); pl.cell(rr,1).alignment=wrap
    pl.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=3)
    a=pl.cell(rr,2,v1); a.font=F(9); a.alignment=wrap; a.border=border; a.fill=fill(CREMA)
    pl.cell(rr,4,k2).font=F(9,True,VERDE); pl.cell(rr,4).alignment=wrap
    b=pl.cell(rr,5,v2); b.font=F(9); b.alignment=wrap; b.border=border; b.fill=fill(AMBAR if amber else CREMA)
    pl.row_dimensions[rr].height=24; rr+=1
def plan_hdr():
    global rr
    pl.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
    c=pl.cell(rr,1,"PLAN DE 90 DÍAS"); c.font=F(10,True,COBRE); c.fill=fill(CREMA); pl.row_dimensions[rr].height=18; rr+=1
    for cc,h in enumerate(["#","Habilidad a desarrollar","Cómo (método)","Métrica que posee","Señal de promoción","Estado"],start=1):
        cell=pl.cell(rr,cc,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=wrapc; cell.border=border
    pl.row_dimensions[rr].height=28; rr+=1
def prow(n,skill,how,metric,signal):
    global rr
    pl.cell(rr,1,n).font=F(9,True); pl.cell(rr,1).alignment=ctr; pl.cell(rr,1).border=border
    for cc,v in [(2,skill),(3,how),(4,metric),(5,signal)]:
        cell=pl.cell(rr,cc,v); cell.font=F(9); cell.alignment=wrap; cell.border=border
    stc=pl.cell(rr,6,"No iniciado"); stc.font=F(9); stc.alignment=ctr; stc.border=border; stc.fill=fill(CREMA)
    dv_st.add(f"F{rr}")
    pl.row_dimensions[rr].height=40; rr+=1
def sched(label):
    global rr
    pl.cell(rr,1,label).font=F(9,True,VERDE); pl.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=3)
    b=pl.cell(rr,4); b.fill=fill(AMBAR); b.border=border; pl.merge_cells(start_row=rr,start_column=4,end_row=rr,end_column=6)
    pl.row_dimensions[rr].height=20; rr+=1
def gap():
    global rr; rr+=1

# ---- JOSÉ ----
person_header("PLAN DE DESARROLLO · PERSONA 1")
kv2("Nombre:","José Mazate","Tipo de jugador (2.2):","Ofensivo",amber=False)
kv2("Rol actual:","Jefe de Sala / Mixólogo (M2)","Etapa actual:","3. Jefe de departamento")
dv_stage.add(f"E{rr-1}")
kv2("Etapa objetivo:","4. Gerente General","Fecha objetivo:","[fecha]")
dv_stage.add(f"B{rr-1}")
plan_hdr()
prow(1,"Propiedad del P&L de sala/barra (leerlo y explicarlo)","Sesiones con el maestro + coaching mensual","Margen de contribución de sala","Explica su P&L sin ayuda")
prow(2,"Correr el scorecard de 7 pasos con TODO su equipo a diario","Shadowing + presentar en el Sync","Cumplimiento diario del scorecard","Lo corre sin que se lo pidas")
prow(3,"Coachear a otros (desarrollar a Efraín como su reemplazo)","Delegarle turnos completos","Efraín cubre sala 1 turno/semana","Entrena a quien lo cubra")
prow(4,"Proveedores de barra + pensamiento estratégico","Incluirlo en negociaciones y planeación","CMV barra 20%","Corre sala sin ti 30 días; pasa la auditoría del Sistema Controlador")
sched("Primera conversación mensual:")
gap()

# ---- JEFFRY ----
person_header("PLAN DE DESARROLLO · PERSONA 2")
kv2("Nombre:","Jeffry López","Tipo de jugador (2.2):","Defensivo (con potencial ofensivo)",amber=False)
kv2("Rol actual:","Jefe de Cocina (M2)","Etapa actual:","3. Jefe de departamento")
dv_stage.add(f"E{rr-1}")
kv2("Etapa objetivo:","4. Gerente General","Fecha objetivo:","[fecha]")
dv_stage.add(f"B{rr-1}")
plan_hdr()
prow(1,"Propiedad del P&L de cocina y del CMV","Sesiones con recetario + maestro","CMV cocina 30%","Mantiene el CMV y lo explica")
prow(2,"Correr el scorecard de 7 pasos con cocina a diario","Shadowing + presentar en el Sync","Cumplimiento diario del scorecard","Lo corre solo")
prow(3,"Reto ofensivo: liderar el menu engineering trimestral (estímulo a su potencial)","Proyecto acotado con guardarraíl de CMV","1 mejora de menú con impacto medible","Propone e implementa una mejora sin que se la pidas")
prow(4,"Entrenar a Nadia como backup + conversaciones de desempeño","Delegar y acompañar","Nadia cubre cocina 1 turno/semana","Trae soluciones al Sync; corre cocina sin ti 30 días")
sched("Primera conversación mensual:")
gap()
pl.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
pl.cell(rr,1,"Banca de sala: cuando José suba a GM, Efraín Taquez es el candidato natural a Jefe de Sala (M2). Desarróllalo en paralelo con la misma ruta (Etapa 2 → 3).").font=F(9,italic=True,color="6B5A45")
pl.cell(rr,1).alignment=wrap; pl.row_dimensions[rr].height=30

pl.column_dimensions["A"].width=5
out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Ruta_Desarrollo_Gerentes.xlsx"
wb.save(out); print("guardado:",out)
