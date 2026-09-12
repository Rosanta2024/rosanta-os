# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Ritmos de Reunión",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.9 · Tres reuniones, ni una más",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"EL PRINCIPIO: 95% MAESTRO",12,True,COBRE); r+=1
wr(r,"Las reuniones existen para enseñar, coachear y resolver problemas. Todo lo demás es un memo. En cada reunión tu rol #1 es enseñar. Si nadie aprendió algo, la reunión fue un desperdicio. Las reuniones son el vehículo del Sistema de 7 Pasos: sin ritmo, el sistema muere.",10,h=44); r+=1
r+=1
wr(r,"ENSEÑA EL ORGANIGRAMA DE MÉTRICAS",12,True,VERDE); r+=1
wr(r,"No revises números: enseña a leerlos. La métrica madre es la UTILIDAD → luego ventas y gastos → luego lo que influye en ventas (rotación, ticket, upsell, reservas) y lo que influye en gastos (CMV cocina 30%, CMV barra 20%, merma, planilla). Conecta cada acción con la utilidad, como en el fútbol cada jugada se conecta con ganar.",10,h=58); r+=1
r+=1
wr(r,"TRES REUNIONES, TRES PROPÓSITOS",12,True,VERDE); r+=1
for t in [
 "Huddle Diario = Alineación. Todos salen sabiendo la meta, el foco y el feature del día. No es reunión, es secuencia de lanzamiento.",
 "Sync Semanal = Coaching. Aquí el dato diario del scorecard de 7 pasos se convierte en patrones semanales. Jeffry y José presentan; tú coacheas sobre patrones, no sobre personas.",
 "Revisión Mensual = Estrategia. Un paso atrás para ver tendencias del P&L (maestro) y el desarrollo del equipo.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"ERRORES COMUNES",12,True,VERDE); r+=1
for t in [
 "1. Saltarse el Huddle por 'estar muy ocupados'. 5 min de alineación ahorran 30 de confusión en servicio. Entre más ocupado, más lo necesitas.",
 "2. Dejar que el Sync se vuelva sesión de quejas. Los M1 presentan DATO del scorecard, no sentimientos. Nada de desahogo sin dato.",
 "3. No coachear. Si tus jefes salen sabiendo lo mismo que al entrar, la reunión falló. Enseña una cosa cada vez.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"NOTA ROSANTA",12,True,VERDE); r+=1
for t in [
 "• El cierre semanal automático (domingos 18:00) y el dashboard del lunes ya te preparan los datos del Sync.",
 "• Con un equipo mayormente defensivo, el bloque de reconocimiento del Sync no es opcional: reconoce la consistencia en público.",
]:
    wr(r,t,10,bg=CREMA,h=26); r+=1

# ===== REUNIONES =====
mt=wb.create_sheet("📋 Reuniones"); mt.sheet_view.showGridLines=False
for col,w in [("A",5),("B",46),("C",12),("D",26),("E",26)]: mt.column_dimensions[col].width=w
mt.merge_cells("A1:E1")
t=mt.cell(1,1,"ROSANTA · Formatos y horarios de reunión"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); mt.row_dimensions[1].height=24
rr=3
def block_title(txt,sub):
    global rr
    mt.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
    c=mt.cell(rr,1,txt); c.font=F(12,True,"FFFFFF"); c.fill=fill(TIERRA); c.alignment=Alignment(vertical="center"); mt.row_dimensions[rr].height=22; rr+=1
    mt.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
    s=mt.cell(rr,1,sub); s.font=F(9,italic=True,color="6B5A45"); mt.row_dimensions[rr].height=16; rr+=1
def hdr():
    global rr
    for c,h in enumerate(["#","Punto","Tiempo","Responsable","Notas / fuente"],start=1):
        cell=mt.cell(rr,c,h); cell.font=F(9,True,MARRON); cell.fill=fill(CREMA); cell.alignment=wrapc; cell.border=border
    mt.row_dimensions[rr].height=16; rr+=1
def item(n,txt,tm,owner,note):
    global rr
    mt.cell(rr,1,n).font=F(9,True); mt.cell(rr,1).alignment=ctr; mt.cell(rr,1).border=border
    mt.cell(rr,2,txt).font=F(9); mt.cell(rr,2).alignment=wrap; mt.cell(rr,2).border=border
    mt.cell(rr,3,tm).font=F(9); mt.cell(rr,3).alignment=ctr; mt.cell(rr,3).border=border
    mt.cell(rr,4,owner).font=F(9); mt.cell(rr,4).alignment=wrap; mt.cell(rr,4).border=border
    mt.cell(rr,5,note).font=F(9); mt.cell(rr,5).alignment=wrap; mt.cell(rr,5).border=border
    mt.row_dimensions[rr].height=28; rr+=1
def sched(label):
    global rr
    a=mt.cell(rr,2,label); a.font=F(9,True,VERDE); a.alignment=wrap
    b=mt.cell(rr,4); b.fill=fill(AMBAR); b.border=border
    mt.merge_cells(start_row=rr,start_column=4,end_row=rr,end_column=5)
    mt.row_dimensions[rr].height=20; rr+=1
def gap():
    global rr; rr+=1

# HUDDLE
block_title("HUDDLE DIARIO","5-10 min · al inicio de cada turno · lo lidera el jefe de turno")
hdr()
item(1,"Meta de ventas de hoy","30 seg","Jefe de turno","Del maestro / dashboard")
item(2,"Foco operativo del día: una sola prioridad","30 seg","Jefe de turno","")
item(3,"Feature del día: un plato + su maridaje de coctel, con guion de venta","60 seg","Jefe de turno","Gastrococtelería: cliente protagonista")
item(4,"86 (agotados), reservas del día, VIPs, alertas","60 seg","Jefe de turno","Reservas: SonTickets")
item(5,"Un momento de enseñanza: punto de coaching de 2 min","2 min","Jefe de turno","Ata la acción a la utilidad")
sched("Hora del Huddle (almuerzo):")
sched("Hora del Huddle (cena):")
gap()

# SYNC
block_title("SYNC SEMANAL","30 min · mismo día y hora cada semana · innegociable")
hdr()
item(1,"Presentación de M1: categorización del equipo + patrones de la semana + planes RAA","15 min","Jeffry (Cocina) y José (Sala) · 5-7 min c/u","Del scorecard de 7 pasos")
item(2,"Bloqueos y apoyo que necesitan","5 min","Ambos M1","")
item(3,"Focos de la próxima semana","5 min","CEO / GM","")
item(4,"Quick wins y reconocimiento","5 min","CEO / GM","Reconoce consistencia (equipo defensivo)")
sched("Día y hora del Sync:")
gap()

# MENSUAL
block_title("REVISIÓN MENSUAL","60 min · primera semana de cada mes")
hdr()
item(1,"Recorrido completo del P&L","15 min","Sistema Controlador + Econtrista","Fuente: maestro")
item(2,"Análisis de tendencias: qué mejoró, qué bajó y por qué","10 min","CEO","")
item(3,"Desarrollo del equipo: quién crece, quién necesita coaching, quién está en riesgo","10 min","CEO","Usa la clasificación: José crece; Jeffry/Marvin con potencial")
item(4,"Foco estratégico del próximo mes","10 min","CEO","")
item(5,"Piso abierto: qué no funciona que nadie ha dicho","15 min","Todos","")
sched("Día del mes para la Revisión:")

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Ritmos_de_Reunion.xlsx"
wb.save(out); print("guardado:",out)
