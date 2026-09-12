# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Flujo de Producción (BOH)",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.16 · Las 4 capas de la consistencia de cocina",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"LAS 4 CAPAS",12,True,COBRE); r+=1
wr(r,"Prep, montaje de estación (mise en place), ejecución de orden y control de calidad. Cada capa tiene su hoja. La prueba: que un cocinero nuevo siga el flujo y produzca un plato aceptable.",10,h=40); r+=1
r+=1
wr(r,"YA TIENES MEDIA CANCHA GANADA",12,True,VERDE); r+=1
wr(r,"Tu recetario con fichas técnicas (125 recetas costeadas) es el CÓMO de cada plato. Súmale una FOTO del plato como debe verse, pégala en el pase, y el cocinero nuevo referencia la foto, no su memoria. La foto en el pase es la herramienta de control de calidad más barata de la cocina.",10,h=44); r+=1
r+=1
wr(r,"CONEXIÓN CON EL CMV Y EL CONTROLADOR",12,True,VERDE); r+=1
for t in [
 "• Toda proteína se PESA (báscula en cada estación). El peso real vs. la porción teórica de la ficha técnica es el cruce anti-sobre-porcionado del Sistema Controlador.",
 "• Toda merma se registra en el waste log → alimenta el inventario semanal y el CMV cocina (meta 30%).",
 "• Toda preparación se etiqueta: nombre, fecha de prep, fecha de consumo, iniciales. FIFO: lo nuevo detrás de lo viejo.",
]:
    wr(r,t,10,bg=CREMA,h=32); r+=1
r+=1
wr(r,"TEMPERATURAS (en °C)",12,True,VERDE); r+=1
wr(r,"Caliente: 60 °C o más (140 °F). Frío: por debajo de 4 °C (40 °F). Chequeo continuo.",10,h=20)

# ===== PREP =====
pp=wb.create_sheet("📋 Estándares de prep"); pp.sheet_view.showGridLines=False
for col,w in [("A",4),("B",26),("C",22),("D",26),("E",14),("F",12),("G",26)]: pp.column_dimensions[col].width=w
pp.merge_cells("A1:G1")
t=pp.cell(1,1,"ROSANTA · Estándares de prep (etiquetado, porción, FIFO)"); t.font=F(13,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); pp.row_dimensions[1].height=24
for c,h in enumerate(["#","Ítem de prep","Receta / porción","Formato de etiqueta","FIFO","Par","Notas"],start=1):
    cell=pp.cell(3,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
pp.row_dimensions[3].height=28
prep=[
 ("Lomo / Rib Eye (porcionado)","Ficha técnica · por peso","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]","Pesar SIEMPRE (cruce anti-sobre-porción)"),
 ("Pesca del día / Róbalo","Ficha técnica · por peso","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]","Fresco; revisar inocuidad"),
 ("Pulpo (pre-cocido)","Ficha técnica","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]",""),
 ("Pork belly / pepián base","Ficha técnica","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]",""),
 ("Costillas café-cardamomo","Ficha técnica","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]",""),
 ("Salsas y purés (pre-elaborados)","Ficha técnica","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]","Taste check 2x/servicio"),
 ("Ensalada Rosanta (componentes)","Ficha técnica","Nombre / fecha / consumo / iniciales","Nuevo atrás","[par]","CMV ref 21%"),
]
rr=4
for i,(item,rec,lab,fifo,par,nota) in enumerate(prep,start=1):
    pp.cell(rr,1,i).font=F(9,True); pp.cell(rr,1).alignment=ctr; pp.cell(rr,1).border=border
    pp.cell(rr,2,item).font=F(9,True,VERDE); pp.cell(rr,2).alignment=wrap; pp.cell(rr,2).border=border
    for c,v in [(3,rec),(4,lab),(5,fifo)]:
        pp.cell(rr,c,v).font=F(9); pp.cell(rr,c).alignment=wrap; pp.cell(rr,c).border=border
    pc=pp.cell(rr,6,par); pc.font=F(9); pc.alignment=ctr; pc.border=border; pc.fill=fill(AMBAR)
    pp.cell(rr,7,nota).font=F(9); pp.cell(rr,7).alignment=wrap; pp.cell(rr,7).border=border
    pp.row_dimensions[rr].height=32; rr+=1
rr+=1
pp.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=7)
pp.cell(rr,1,"Regla dura: cada ítem etiquetado, cada proteína pesada. Sin excepciones. La lista de prep la genera Jeffry cada día (según pars + reservas SonTickets) y se publica antes de las 7 AM.").font=F(9,italic=True,color="6B5A45")
pp.cell(rr,1).alignment=wrap; pp.row_dimensions[rr].height=30

# ===== MISE EN PLACE =====
st=wb.create_sheet("📋 Mise en place"); st.sheet_view.showGridLines=False
for col,w in [("A",22),("B",30),("C",34),("D",24),("E",22)]: st.column_dimensions[col].width=w
st.merge_cells("A1:E1")
t=st.cell(1,1,"ROSANTA · Montaje de estación (antes de cada servicio)"); t.font=F(13,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); st.row_dimensions[1].height=24
for c,h in enumerate(["Estación","Chequeo de equipo","Mise en place (par al volumen)","Ubicación de respaldo","Notas"],start=1):
    cell=st.cell(3,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr; cell.border=border
st.row_dimensions[3].height=28
stations=["Parrilla / Grill","Salteado / Sauté","Frío / Ensaladas","Pase / Expo","Postres"]
for i,stn in enumerate(stations):
    rr=4+i
    st.cell(rr,1,stn).font=F(9,True,VERDE); st.cell(rr,1).alignment=wrap; st.cell(rr,1).border=border
    for c in range(2,6):
        cell=st.cell(rr,c); cell.border=border; cell.alignment=wrap; cell.font=F(9); cell.fill=fill(CREMA if c in (3,4) else "FFFFFF")
    st.row_dimensions[rr].height=34
rr=4+len(stations)+1
st.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
st.cell(rr,1,"Por estación antes de cada servicio: 1) equipo encendido y funcionando, 2) par surtido al volumen esperado, 3) herramientas en el MISMO lugar siempre, 4) plan de respaldo (dónde está el restock, cómo pedir ayuda).").font=F(9,italic=True,color="6B5A45")
st.cell(rr,1).alignment=wrap; st.row_dimensions[rr].height=32

# ===== EJECUCIÓN + QC =====
qc=wb.create_sheet("📋 Ejecución y calidad"); qc.sheet_view.showGridLines=False
for col,w in [("A",5),("B",30),("C",60)]: qc.column_dimensions[col].width=w
qc.merge_cells("A1:C1")
t=qc.cell(1,1,"ROSANTA · Ejecución de orden y control de calidad"); t.font=F(13,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); qc.row_dimensions[1].height=24
rr=3
qc.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=3)
qc.cell(rr,1,"FLUJO DE EJECUCIÓN DE ORDEN").font=F(10,True,COBRE); qc.cell(rr,1).fill=fill(CREMA); qc.row_dimensions[rr].height=18; rr+=1
flow=[
 "Imprime el ticket → el expo (Jeffry / responsable de turno) canta la orden.",
 "Cada estación arranca sus componentes.",
 "La estación canta 'listo' al emplatar.",
 "El expo arma la orden completa y la coteja contra el ticket.",
 "Control de calidad: platos correctos, estándar de emplatado (ficha técnica + foto), temperatura.",
 "Orden al pase → el mesero la recoge.",
]
for i,step in enumerate(flow,start=1):
    qc.cell(rr,1,i).font=F(9,True); qc.cell(rr,1).alignment=ctr; qc.cell(rr,1).border=border
    qc.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=3)
    c=qc.cell(rr,2,step); c.font=F(9); c.alignment=wrap; c.border=border
    qc.row_dimensions[rr].height=24; rr+=1
rr+=1
qc.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=3)
qc.cell(rr,1,"PUNTOS DE CONTROL DE CALIDAD").font=F(10,True,COBRE); qc.cell(rr,1).fill=fill(CREMA); qc.row_dimensions[rr].height=18; rr+=1
for c,h in enumerate(["Chequeo","Cuándo","Qué revisar"],start=1):
    cell=qc.cell(rr,c,h); cell.font=F(9,True,MARRON); cell.border=border; cell.alignment=wrapc
qc.row_dimensions[rr].height=16; rr+=1
checks=[
 ("Line check","Antes de cada servicio","Montaje, producto fresco, PESOS de porción (cruzar vs. ficha técnica)"),
 ("Plate check","Cada orden en el pase","Platos correctos, coincide con la foto/ficha, sin errores"),
 ("Taste check","2x por servicio","Salsas, sopas y especiales por consistencia"),
 ("Temp check","Continuo","Caliente 60 °C+ (140 °F); frío por debajo de 4 °C (40 °F)"),
]
for chk,when,what in checks:
    qc.cell(rr,1,chk).font=F(9,True,VERDE); qc.cell(rr,1).alignment=wrap; qc.cell(rr,1).border=border
    qc.cell(rr,2,when).font=F(9); qc.cell(rr,2).alignment=wrap; qc.cell(rr,2).border=border
    c=qc.cell(rr,3,what); c.font=F(9); c.alignment=wrap; c.border=border; c.fill=fill(CREMA)
    qc.row_dimensions[rr].height=30; rr+=1
rr+=1
qc.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=3)
qc.cell(rr,1,"Specs de emplatado: fotografía tus 10 platos top como deben verse y pégalos en el pase. Ya tienes la ficha técnica en el recetario; súmale la foto.").font=F(9,italic=True,color="6B5A45")
qc.cell(rr,1).alignment=wrap; qc.row_dimensions[rr].height=28

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Flujo_Produccion_BOH.xlsx"
wb.save(out); print("guardado:",out)
