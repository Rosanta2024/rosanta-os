# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center")
ctr=Alignment(horizontal="center",vertical="center")
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Escalera de Reemplazo (Replacement Ladder)",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.4 · Reemplázate a ti mismo, un sombrero a la vez",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"LA IDEA",12,True,COBRE); r+=1
wr(r,"No se trata de tapar huecos, sino de sacarte a TI de la operación en el orden que te da más libertad más rápido. Mira tu H.O.T. Canvas: ¿qué sombrero, si lo entregaras mañana, te devolvería más horas por semana? Ese es tu primer reemplazo.",10,h=44); r+=1
r+=1
wr(r,"TU PUNTO DE PARTIDA EN ROSANTA",12,True,VERDE); r+=1
for t in [
 "Ya tienes cubiertos los dos primeros peldaños del curso: Jefe de Cocina (Jeffry López) y Jefe de Sala/FOH (José Mazate). Ese es un gran avance.",
 "Por eso esta escalera se enfoca en los sombreros que TODAVÍA cargas tú como CEO: gestión diaria, finanzas/verificación, reservas, compras y eventos/crecimiento.",
 "Los artefactos ya bajaron parte de esa carga (bot, maestro, cotizador, paneles). La escalera te dice a quién contratar para soltar el resto.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"CÓMO USAR LA HOJA 📋",12,True,VERDE); r+=1
for t in [
 "1. Las celdas ámbar son estimaciones mías: cámbialas por tus números reales (horas/semana, impacto 1-5, costo mensual en Q).",
 "2. El 'Índice de prioridad' se calcula solo = horas ÷ (costo/1000) × impacto. Más alto = reemplázalo primero.",
 "3. La 'Prioridad' (ALTA/MEDIA/BAJA) se actualiza sola según el índice.",
 "4. Abajo, tus 3 contrataciones en orden, con puesto, costo y fecha objetivo.",
]:
    wr(r,t,10,bg=CREMA,h=26); r+=1
r+=1
wr(r,"SECUENCIA GENERAL DEL CURSO (referencia)",12,True,VERDE); r+=1
for t in [
 "1. Jefe de Cocina/Sous — te saca de la operación de comida.  [Rosanta: HECHO — Jeffry]",
 "2. Jefe de Sala/Gerente FOH — te saca del manejo de crisis con clientes.  [Rosanta: HECHO — José]",
 "3. Controlador — verificación independiente de resultados. En Rosanta se resuelve como SISTEMA (Claude + artefactos + cerebro + intranet), no como persona. Conecta con Lección 2.6 y el scorecard de 7 pasos.",
 "4. Sub-gerente/Gerente General — te saca de las decisiones diarias.",
 "5. Coordinador de Marketing/Catering — hace crecer los ingresos sin tu esfuerzo personal.",
]:
    wr(r,t,10,h=30); r+=1
r+=1
wr(r,"ERRORES COMUNES",12,True,VERDE); r+=1
for t in [
 "1. Contratar lo que QUIERES en vez de lo que necesitas (marketing suena emocionante, pero primero suelta lo que te tiene enterrado).",
 "2. Contratar demasiado senior demasiado pronto: no necesitas un GM caro antes de tener los sistemas para que lo administre.",
 "3. No considerar opciones en el extranjero: el Controlador/contabilidad no necesita estar en el edificio (Lecciones 2.7-2.8).",
]:
    wr(r,t,10,h=30); r+=1

# ===== ESCALERA =====
la=wb.create_sheet("📋 Escalera de Reemplazo"); la.sheet_view.showGridLines=False
for col,w in [("A",5),("B",34),("C",13),("D",13),("E",15),("F",16),("G",13),("H",12)]: la.column_dimensions[col].width=w
la.merge_cells("A1:H1")
t=la.cell(1,1,"ROSANTA · Escalera de Reemplazo — puntúa tus sombreros"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); la.row_dimensions[1].height=24
la.merge_cells("A2:H2")
la.cell(2,1,"Celdas ámbar = estimaciones editables. Índice y Prioridad se calculan solos. Fuente de sombreros: tu H.O.T. Canvas (CEO).").font=F(9,italic=True,color="6B5A45")
hdr=["#","Sombrero (del H.O.T. Canvas)","Horas/sem que le dedicas","Impacto ingresos (1-5)","Costo mensual estimado (Q)","Tiempo a productividad","Índice prioridad","Prioridad"]
for c,h in enumerate(hdr,start=1):
    cell=la.cell(4,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=wrapc; cell.border=border
la.row_dimensions[4].height=42
# fila ejemplo
ej=["Ej.","Operación de cocina (abrir, prep, línea)",25,4,4000,"4-6 semanas"]
for c,v in enumerate(ej,start=1):
    cell=la.cell(5,c,v); cell.font=F(9,italic=True,color="6B5A45"); cell.fill=fill(CREMA); cell.border=border; cell.alignment=wrap if c==2 else ctr
la.cell(5,7,"=IF(E5=0,0,ROUND(C5*D5/(E5/1000),1))").font=F(9,italic=True,color="6B5A45"); la.cell(5,7).fill=fill(CREMA); la.cell(5,7).border=border; la.cell(5,7).alignment=ctr
la.cell(5,8,'=IF(G5="","",IF(G5>=15,"ALTA",IF(G5>=7,"MEDIA","BAJA")))').font=F(9,italic=True,color="6B5A45"); la.cell(5,8).fill=fill(CREMA); la.cell(5,8).border=border; la.cell(5,8).alignment=ctr
la.row_dimensions[5].height=30
# filas Rosanta (estimaciones)
hats=[
 (1,"Gestión diaria / operaciones (abrir-cerrar, decisiones del día)",20,4,9000,"8-12 semanas"),
 (2,"Finanzas + verificación independiente (Controlador)",8,3,3500,"4-6 semanas"),
 (3,"Reservas y atención de leads",6,3,4000,"3-4 semanas"),
 (4,"Compras y proveedores",5,3,4500,"4 semanas"),
 (5,"Eventos y catering / crecimiento",8,5,6000,"6-8 semanas"),
]
row=6
for num,hat,hrs,imp,cost,ttp in hats:
    la.cell(row,1,num).font=F(9,True); la.cell(row,1).alignment=ctr; la.cell(row,1).border=border
    c2=la.cell(row,2,hat); c2.font=F(9); c2.alignment=wrap; c2.border=border
    for c,v in [(3,hrs),(4,imp),(5,cost),(6,ttp)]:
        cell=la.cell(row,c,v); cell.font=F(9); cell.alignment=ctr; cell.border=border; cell.fill=fill(AMBAR)
    g=la.cell(row,7,f"=IF(E{row}=0,0,ROUND(C{row}*D{row}/(E{row}/1000),1))"); g.font=F(9,True); g.alignment=ctr; g.border=border
    h=la.cell(row,8,f'=IF(G{row}="","",IF(G{row}>=15,"ALTA",IF(G{row}>=7,"MEDIA","BAJA")))'); h.font=F(9,True); h.alignment=ctr; h.border=border
    la.row_dimensions[row].height=30; row+=1
# validacion impacto 1-5
dv=DataValidation(type="list",formula1='"1,2,3,4,5"',allow_blank=True); dv.prompt="Impacto en ingresos 1 (bajo) a 5 (alto)"; la.add_data_validation(dv); dv.add("D5:D10")
# leyenda indice
lr=row+1
la.merge_cells(start_row=lr,start_column=1,end_row=lr,end_column=8)
la.cell(lr,1,"Índice = horas/sem ÷ (costo mensual/1000) × impacto. Prioridad: ALTA ≥15 · MEDIA ≥7 · BAJA <7. Ajusta los números ámbar a tu realidad y el orden se recalcula.").font=F(8,italic=True,color="6B5A45")
la.cell(lr,1).alignment=wrap; la.row_dimensions[lr].height=30

# ===== TOP 3 =====
tr=lr+2
la.merge_cells(start_row=tr,start_column=1,end_row=tr,end_column=8)
la.cell(tr,1,"TUS 3 CONTRATACIONES PRIORITARIAS").font=F(12,True,VERDE); la.row_dimensions[tr].height=22
tr+=1
h2=["Prioridad","Puesto","Qué reemplaza (sombreros)","Horas recuperadas/sem","Costo mensual (Q)","Fecha objetivo de inicio","Estado"]
# columnas: usar A..G (7)
for c,h in enumerate(h2,start=1):
    cell=la.cell(tr,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=wrapc; cell.border=border
la.row_dimensions[tr].height=34
top=[
 ("#1","Jefe de Operaciones / Gerente","Gestión diaria + supervisión de compras y reservas",20,9000,"[fecha]","No iniciado"),
 ("#2","Coordinador de Eventos y Catering","Eventos/crecimiento (tu palanca #1 de ingresos)",8,6000,"[fecha]","No iniciado"),
 ("#3","Sistema Controlador (Claude + artefactos + cerebro + intranet)","Finanzas y verificación independiente de resultados (sistema, no persona)",8,800,"En curso","Construyendo"),
]
row=tr+1
for pr,role,repl,hrs,cost,fecha,est in top:
    la.cell(row,1,pr).font=F(9,True,COBRE); la.cell(row,1).alignment=ctr; la.cell(row,1).border=border
    for c,v in [(2,role),(3,repl)]:
        cell=la.cell(row,c,v); cell.font=F(9,True,VERDE) if c==2 else F(9); cell.alignment=wrap; cell.border=border; cell.fill=fill(CREMA)
    for c,v in [(4,hrs),(5,cost),(6,fecha)]:
        cell=la.cell(row,c,v); cell.font=F(9); cell.alignment=ctr; cell.border=border; cell.fill=fill(AMBAR)
    e=la.cell(row,7,est); e.font=F(9); e.alignment=ctr; e.border=border
    la.row_dimensions[row].height=34; row+=1
dv2=DataValidation(type="list",formula1='"No iniciado,Buscando,Entrevistando,Contratado,Construyendo,En curso,Operando"',allow_blank=True); la.add_data_validation(dv2); dv2.add(f"G{tr+1}:G{tr+3}")
# nota
nr=row+1
la.merge_cells(start_row=nr,start_column=1,end_row=nr,end_column=7)
la.cell(nr,1,"Decisión de Juanma: el puesto #3 (Controlador) NO será una persona, sino un SISTEMA — Claude + los artefactos + el cerebro + la intranet — que verifica resultados de forma independiente (Lección 2.6). Costo = herramientas, no salario. Requisitos para que funcione como control real: (1) fuentes de dato inviolables por la persona medida (POS con reglas de anulación, conteos de inventario cruzados), (2) alguien que actúe sobre las alertas (hoy tú; mañana el GM), (3) mantener vivas las conexiones/automatizaciones. El GM sigue siendo #1 por horas liberadas; Eventos #2.").font=F(8,italic=True,color="6B5A45")
la.cell(nr,1).alignment=wrap; la.row_dimensions[nr].height=70

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Escalera_de_Reemplazo.xlsx"
wb.save(out); print("guardado:",out)
