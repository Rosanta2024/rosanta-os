# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF2CC"
FONT="Arial"
def F(sz=10,b=False,color=MARRON,italic=False): return Font(name=FONT,size=sz,bold=b,color=color,italic=italic)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="BFB4A0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); wrapc=Alignment(wrap_text=True,vertical="center"); ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
wb=openpyxl.Workbook()

# ===== INSTRUCCIONES =====
ws=wb.active; ws.title="📖 Instrucciones"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=108
def wr(row,text,sz=10,b=False,color=MARRON,bg=None,italic=False,h=None):
    c=ws.cell(row=row,column=2,value=text); c.font=F(sz,b,color,italic); c.alignment=wrap
    if bg: c.fill=fill(bg); ws.cell(row=row,column=1).fill=fill(bg)
    if h: ws.row_dimensions[row].height=h
r=1; ws.row_dimensions[1].height=30
wr(r,"ROSANTA · Sistema Controlador · Panel de 3 Señales",16,True,"FFFFFF",VERDE,h=30); r+=1
wr(r,"Cocina con Carisma · Lección 2.6 · Verificación independiente de resultados",9,False,"FFFFFF",VERDE,italic=True); r+=1
r+=1
wr(r,"EL PROBLEMA DE FONDO",12,True,COBRE); r+=1
wr(r,"Quien genera el resultado nunca debe ser el único que lo reporta. No es desconfianza: es naturaleza humana. Si tu jefe de cocina es responsable del CMV Y el único que te lo reporta, el dato llega maquillado. El Sistema Controlador crea una capa de verificación independiente que te reporta directo a ti.",10,h=44); r+=1
r+=1
wr(r,"TU CONTROLADOR ES UN SISTEMA, NO UNA PERSONA",12,True,VERDE); r+=1
wr(r,"En Rosanta el Controlador = Claude + los artefactos + el maestro + la intranet. Los datos se jalan de fuentes que la persona medida NO controla: POS (Posfile), FEL, bancos, panel de reseñas, inventarios físicos, SonTickets. El cierre semanal (domingos 18:00) y el dashboard del lunes ya son la 'corrida' del controlador. Costo: herramientas, no salario.",10,h=58); r+=1
r+=1
wr(r,"LAS 3 SEÑALES",12,True,VERDE); r+=1
for t in [
 "Daily Flash = Pulso. ¿El restaurante está sano HOY? Fin de cada día.",
 "Weekly Scorecard = Tendencias. ¿Vamos mejorando o empeorando? Semanal.",
 "Monthly Review = Estrategia. ¿Damos en los objetivos? Mensual.",
]:
    wr(r,t,10,h=20); r+=1
r+=1
wr(r,"REQUISITOS PARA QUE SEA CONTROL REAL",12,True,VERDE); r+=1
for t in [
 "1. Fuente inviolable: la persona medida no puede editar el dato. Reglas sobre anulaciones/cortesías en POS; inventario con conteo CRUZADO (no lo cuenta solo quien se mide).",
 "2. Verificar procesos, no solo dinero: ¿se hizo el line check? ¿se pesaron las porciones? ¿se llenó el log de merma? Son los indicadores que anticipan el resultado financiero.",
 "3. Alguien actúa sobre la alerta: el sistema marca la desviación; la decisión de corregir es tuya (mañana, del GM).",
]:
    wr(r,t,10,bg=CREMA,h=32); r+=1
r+=1
wr(r,"EJEMPLO (Founders Board) — aplicado a Rosanta",12,True,VERDE); r+=1
wr(r,"Un controlador overseas cruzó cámaras con el log de porciones y descubrió sobre-porcionado de lomo del 15% (invisible en el P&L porque el jefe de cocina era el único que reportaba). El CMV bajó 3% = +$50K/año. Tu analogía: cruzar la porción teórica del recetario vs. ventas del POS vs. merma/inventario detecta el mismo hueco. El waste log y el inventario semanal son ese control.",10,h=58); r+=1

# ===== PANEL 3 SEÑALES =====
cp=wb.create_sheet("📋 Panel de 3 señales"); cp.sheet_view.showGridLines=False
for col,w in [("A",5),("B",34),("C",40),("D",18),("E",26)]: cp.column_dimensions[col].width=w
cp.merge_cells("A1:E1")
t=cp.cell(1,1,"ROSANTA · Panel de Control de 3 Señales"); t.font=F(14,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); cp.row_dimensions[1].height=24
rr=3
def sig(title,sub):
    global rr
    cp.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
    c=cp.cell(rr,1,title); c.font=F(12,True,"FFFFFF"); c.fill=fill(TIERRA); c.alignment=Alignment(vertical="center"); cp.row_dimensions[rr].height=22; rr+=1
    cp.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)
    cp.cell(rr,1,sub).font=F(9,italic=True,color="6B5A45"); cp.row_dimensions[rr].height=16; rr+=1
    for cc,h in enumerate(["#","Métrica","Fuente independiente (artefacto)","Meta","Notas"],start=1):
        cell=cp.cell(rr,cc,h); cell.font=F(9,True,MARRON); cell.fill=fill(CREMA); cell.alignment=wrapc; cell.border=border
    cp.row_dimensions[rr].height=16; rr+=1
def m(n,met,src,tgt,note=""):
    global rr
    cp.cell(rr,1,n).font=F(9,True); cp.cell(rr,1).alignment=ctr; cp.cell(rr,1).border=border
    cp.cell(rr,2,met).font=F(9); cp.cell(rr,2).alignment=wrap; cp.cell(rr,2).border=border
    cp.cell(rr,3,src).font=F(9); cp.cell(rr,3).alignment=wrap; cp.cell(rr,3).border=border
    tc=cp.cell(rr,4,tgt); tc.font=F(9); tc.alignment=ctr; tc.border=border; tc.fill=fill(AMBAR)
    cp.cell(rr,5,note).font=F(9); cp.cell(rr,5).alignment=wrap; cp.cell(rr,5).border=border
    cp.row_dimensions[rr].height=26; rr+=1
def gap():
    global rr; rr+=1

sig("SEÑAL 1 · DAILY FLASH (pulso)","Fin de cada día · registra el jefe de turno · dato del POS (independiente)")
m(1,"Ventas netas del día","POS (Posfile) → maestro","[meta]")
m(2,"% mano de obra (labor)","Horas/planilla vs ventas POS","[meta]")
m(3,"# comensales y ticket promedio","POS","~Q259")
m(4,"Anulaciones / cortesías (voids/comps)","Reporte de anulaciones del POS","[límite]","Requiere autorización del jefe de turno")
m(5,"Propinas registradas","POS (passthrough, no es costo)","")
gap()
sig("SEÑAL 2 · WEEKLY SCORECARD (tendencias)","Semanal · lo corre el Sistema Controlador (cierre-semanal + dashboard)")
m(1,"Ventas netas promedio/día","Maestro / dashboard-semanal","[meta]")
m(2,"CMV cocina %","Recetario + inventario semanal","30%")
m(3,"CMV barra %","Recetario + inventario semanal","20%")
m(4,"% mano de obra","Planilla + POS","[meta]")
m(5,"Prime cost (CMV + labor)","Calculado (maestro)","[meta]")
m(6,"Ticket promedio","POS","~Q259")
m(7,"Horas extra","Planilla","[límite]")
m(8,"Reseñas (rating promedio)","panel-reseñas + Google","[meta]")
m(9,"Merma (Q y %)","Log de merma / inventarios","[límite]","Cruza vs. porción teórica del recetario")
m(10,"Rotación de mesa","POS / SonTickets","45 min")
gap()
sig("SEÑAL 3 · MONTHLY REVIEW (estrategia)","Mensual · dueño + sistema")
m(1,"P&L completo","Maestro (reporte fin de mes)","")
m(2,"Progreso de metas trimestrales","Ruta 2x3x5 / mandala","")
m(3,"ROI de marketing / pauta","Auditoría Meta ads + tablero","[meta]")
m(4,"Rotación de personal","Planilla","[límite]")
m(5,"Ocupación entre semana","Dashboard","[meta]")

# ===== DISEÑO DEL SISTEMA =====
ds=wb.create_sheet("🔧 Diseño del sistema"); ds.sheet_view.showGridLines=False
for col,w in [("A",3),("B",30),("C",76)]: ds.column_dimensions[col].width=w
ds.row_dimensions[1].height=24
ds.cell(1,2,"ROSANTA · Cómo opera el Sistema Controlador").font=F(14,True,VERDE)
rows=[
 ("Quién es el Controlador","Claude + artefactos + maestro + intranet. Reporta directo a ti (Juanma), no al jefe de área. Independiente de quien genera el resultado."),
 ("Corrida automática","Cierre-semanal (domingos 18:00) reconcilia la semana; dashboard-refresh (lunes) actualiza; reporte-semanal actualiza el maestro y, fin de mes, el reporte del contador; auditoría Meta ads (día 25)."),
 ("Control anti-manipulación 1","Anulaciones/cortesías del POS solo con autorización del jefe de turno; el reporte de voids se cruza en el Daily Flash."),
 ("Control anti-manipulación 2","Inventario con conteo CRUZADO: no lo cuenta solo la persona medida. Inventario semanal (artefacto) + log de merma."),
 ("Cruce anti-sobre-porcionado","Porción teórica del recetario vs. ventas del POS vs. consumo real (inventario/merma). Si no cuadra, hay sobre-porcionado o fuga."),
 ("Verificación de procesos","Marca si se hizo el line check, si se pesaron porciones, si se llenó el waste log, si se cumplió el mise en place. Son los indicadores que anticipan el CMV."),
 ("Quién actúa sobre la alerta","Hoy: tú. Cuando exista el GM (José): él, y tú revisas el resumen. El sistema marca; la persona decide."),
 ("Dónde vive","Intranet (Finanzas + Recetario) como vista única del equipo; artefactos vivos en Cowork para el detalle."),
]
rr=3
ds.cell(rr,2,"Componente").font=F(9,True,"FFFFFF"); ds.cell(rr,2).fill=fill(TIERRA); ds.cell(rr,2).border=border; ds.cell(rr,2).alignment=wrapc
ds.cell(rr,3,"Cómo funciona en Rosanta").font=F(9,True,"FFFFFF"); ds.cell(rr,3).fill=fill(TIERRA); ds.cell(rr,3).border=border; ds.cell(rr,3).alignment=wrapc
ds.row_dimensions[rr].height=18; rr+=1
for k,v in rows:
    a=ds.cell(rr,2,k); a.font=F(9,True,VERDE); a.alignment=wrap; a.border=border; a.fill=fill(CREMA)
    b=ds.cell(rr,3,v); b.font=F(9); b.alignment=wrap; b.border=border
    ds.row_dimensions[rr].height=44; rr+=1

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Sistema_Controlador_Panel.xlsx"
wb.save(out); print("guardado:",out)
