# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
VERDE="4A6741"; TIERRA="A0785A"; CREMA="F5EFE0"; COBRE="C9923F"; MARRON="3D2B1F"; AMBAR="FFF7E0"
def F(sz=10,b=False,color=MARRON,it=False): return Font(name="Arial",size=sz,bold=b,color=color,italic=it)
def fill(h): return PatternFill("solid",fgColor=h)
thin=Side(style="thin",color="D9D2C2"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top"); ctr=Alignment(horizontal="center",vertical="center")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Checklist SOP Sala"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=5; ws.column_dimensions["B"].width=60; ws.column_dimensions["C"].width=14; ws.column_dimensions["D"].width=42

ws.merge_cells("A1:D1")
t=ws.cell(1,1,"ROSANTA · Validación del SOP de Sala"); t.font=F(15,True,"FFFFFF"); t.fill=fill(VERDE); t.alignment=Alignment(vertical="center"); ws.row_dimensions[1].height=26
ws.merge_cells("A2:D2")
ws.cell(2,1,"Marca cada actividad como Sí (correcta) / No / Ajustar. Usa 'Comentarios' para escribir cambios. Las filas ámbar son para AGREGAR tareas nuevas.").font=F(9,it=True,color="6B5A45"); ws.row_dimensions[2].height=26

hdr=["#","Actividad","¿Correcta?","Comentarios / cambios / agregar"]
for c,h in enumerate(hdr,1):
    cell=ws.cell(3,c,h); cell.font=F(9,True,"FFFFFF"); cell.fill=fill(TIERRA); cell.alignment=ctr if c!=2 else Alignment(horizontal="left",vertical="center"); cell.border=border
ws.row_dimensions[3].height=20

data=[
 ("0 · Configuración del turno",[
   "Días bajos (lun–mié): coctelero (supervisor + barra) + 1 mesero (recepción + piso)",
   "Fin de semana: 1 mixólogo + 2 meseros (2 mixólogos + 2 meseros si hay mucha gente)",
   "Recepción: la cubre quien esté más cerca de la puerta (no hay host dedicado)",
   "El supervisor de turno revisa reservas en WIX, asigna secciones y vigila tiempos",
 ]),
 ("1 · Apertura de sala",[
   "Revisar reservas y notas de grupos del día en WIX (rosanta.rest)",
   "Montar y pulir mesas; revisar copas antes de sacarlas de la estantería",
   "Surtir estaciones: cubiertos, servilletas, pan, agua",
   "Verificar ambiente: música, iluminación, jardín/pérgola/chimenea, baños",
   "Repasar el feature del día y su maridaje (del Huddle)",
   "Uniforme y presentación personal",
 ]),
 ("2 · Recepción y bienvenida",[
   "Saludar en ≤15 segundos, tono cálido: 'Bienvenidos a Rosanta'",
   "Preguntar si tienen reservación (verificar en WIX)",
   "Acompañar a la mesa; ofrecer área: salón, pérgola o jardín",
   "En jardín/pérgola: ofrecer manta; cojín de respaldo si aplica",
   "Atender con prioridad a personas mayores o con capacidades diferentes",
 ]),
 ("3 · Servicio de mesa — Presentación",[
   "Presentarse por nombre; breve conversación cálida",
   "Explicar la gastrococtelería (maridaje plato + coctel); cócteles modificables (con/sin alcohol); bebidas de casa, vinos y cervezas",
   "Presentar el menú, el feature del día y opciones fuera de carta",
   "En grupos grandes: preguntar si separan cuenta y considerar los puestos",
 ]),
 ("3 · Servicio de mesa — Toma de pedido",[
   "Numerar puestos en contra de las agujas del reloj, empezando por la izquierda (lado cocina)",
   "Registrar el pedido en el POS; confirmar modificaciones, alergias y comunicarlas a cocina",
   "Recomendar el maridaje; sugerir 1 entrante por cada 2 personas",
 ]),
 ("3 · Servicio de mesa — Tiempos y servicio",[
   "Servir bebidas en 3–5 minutos; revisar copas de vino antes de servir",
   "Controlar tiempos de entrantes y de salida de fuertes con sus maridajes",
   "Sacar los platos preferiblemente juntos, a la temperatura correcta; pedir apoyo",
   "Ofrecer pan adicional; retirar entrantes; cambiar cubiertos si aplica",
 ]),
 ("3 · Servicio de mesa — Atención",[
   "Check-back poco después de los fuertes (satisfacción)",
   "Monitorear niveles de bebida; rellenar copas; retirar platos vacíos",
   "Ofrecer postres o un segundo coctel con recomendación",
 ]),
 ("3 · Servicio de mesa — Facturación y despedida",[
   "El mesero verifica su cuenta antes de cerrarla",
   "Entregar la cuenta a tiempo y correcta",
   "Preguntar si desean factura (llega por FEL); pago tarjeta/efectivo; cambio USD/EUR a Q7",
   "Agradecer, invitar a volver, resetear la mesa; invitar a dejar reseña con el QR",
 ]),
 ("4 · Supervisor de turno",[
   "Revisar reservas del día (WIX) y asignar secciones",
   "Vigilar tiempos de salida y maridajes durante el servicio",
   "Manejar quejas con el protocolo LAST (Escucha, Reconoce, Resuelve, Agradece)",
   "Autorizar cortesías/descuentos según política; registrarlas en el POS",
   "Cerrar el turno: cuadre y resumen al dueño/GM (Daily Flash)",
 ]),
 ("5 · Cierre de sala",[
   "Side work por estación: limpiar, resurtir y guardar",
   "Ordenar el ambiente: apagar chimenea, música y luces exteriores",
   "Cuadre de propinas",
   "Dejar todo montado para el siguiente turno",
 ]),
 ("6 · Excepciones",[
   "Queja en mesa → protocolo LAST",
   "Plato agotado (86) → avisar de inmediato al pase y ofrecer alternativa",
   "Cortesía o descuento → según política; registrar en el POS con motivo",
   "Espera excedida → supervisor se acerca, se disculpa, da tiempo; +10 min = cortesía",
 ]),
]
dv=DataValidation(type="list",formula1='"Sí,No,Ajustar"',allow_blank=True)
ws.add_data_validation(dv)
r=4; n=0
for bloque,acts in data:
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
    c=ws.cell(r,1,bloque); c.font=F(10,True,VERDE); c.fill=fill(CREMA); c.alignment=Alignment(vertical="center"); c.border=border
    ws.row_dimensions[r].height=18; r+=1
    for a in acts:
        n+=1
        ws.cell(r,1,n).font=F(9,True); ws.cell(r,1).alignment=ctr; ws.cell(r,1).border=border
        ws.cell(r,2,a).font=F(9); ws.cell(r,2).alignment=wrap; ws.cell(r,2).border=border
        cc=ws.cell(r,3); cc.border=border; cc.alignment=ctr; cc.fill=fill("FFFFFF"); dv.add(cc)
        ws.cell(r,4).border=border; ws.cell(r,4).alignment=wrap
        ws.row_dimensions[r].height=30; r+=1
    # 2 filas para agregar
    for _ in range(2):
        ws.cell(r,1,"+").font=F(9,True,COBRE); ws.cell(r,1).alignment=ctr; ws.cell(r,1).border=border; ws.cell(r,1).fill=fill(AMBAR)
        ac=ws.cell(r,2,""); ac.border=border; ac.fill=fill(AMBAR); ac.alignment=wrap
        cc=ws.cell(r,3); cc.border=border; cc.fill=fill(AMBAR); dv.add(cc)
        ws.cell(r,4).border=border; ws.cell(r,4).fill=fill(AMBAR)
        ws.row_dimensions[r].height=22; r+=1

out="/sessions/clever-hopeful-pascal/mnt/outputs/Rosanta_Checklist_SOP_Sala.xlsx"
wb.save(out); print("guardado:",out,"| actividades:",n)
