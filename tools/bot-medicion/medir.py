#!/usr/bin/env python3
"""Escaladas, reservas y eventos del bot, de cada 10 chats (16-oct-2026).

Uso:  python3 medir.py <Rosanta_Leads.xlsx | export.json> <desde AAAA-MM-DD> <hasta AAAA-MM-DD> [--detalle] [--salida informe.md]
      (hasta es exclusivo). El .xlsx se baja de la hoja "Rosanta Leads"
      (Archivo › Descargar › Microsoft Excel), o es el .json que guarda el conector de Drive
      (download_file_content con exportMimeType xlsx: {"content": base64, ...}).
      --detalle muestra la conversacion de cada chat que paso a persona; --salida agrega todo a un .md.

Criterio (el mismo que reprodujo la línea base del cerebro v44: 3.7 ≈ 3.8):
  chat      = un contacto (canal + ID destino, o teléfono) con al menos una fila en Leads
  a persona = alguna fila con ¿Humano? = Sí
  reserva   = alguna fila con resumen "Reserva creada en Wix"
  evento    = alguna fila con tipo = evento
Se excluyen las filas de prueba (mensaje con "PRUEBA") y el canal correo (la linea base es de chats;
--con-correo lo incluye).
--detalle lista los chats que pasaron a persona, para leer el motivo en Conversaciones.
"""
import sys, io, json, base64, datetime as dt, collections
import openpyxl

def fecha(x):
    if isinstance(x, dt.datetime): return x
    for f in ('%d/%m/%Y %H:%M:%S', '%Y-%m-%d %H:%M:%S'):
        try: return dt.datetime.strptime(str(x)[:19], f)
        except ValueError: pass
    return None

def ident(v):
    s = str(v or '').strip()
    return s[:-2] if s.endswith('.0') else s

def main():
    if len(sys.argv) < 4: sys.exit(__doc__)
    ruta, desde, hasta = sys.argv[1], dt.datetime.fromisoformat(sys.argv[2]), dt.datetime.fromisoformat(sys.argv[3])
    if ruta.endswith('.xlsx'):
        wb = openpyxl.load_workbook(ruta, read_only=True)
    else:   # export del conector de Drive: JSON con el xlsx en base64
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(json.load(open(ruta))['content'])), read_only=True)
    filas = [tuple(list(r) + [None] * 16)[:16] for r in list(wb['Leads'].iter_rows(values_only=True))[1:]]
    chats = collections.defaultdict(list)
    for r in filas:
        f = fecha(r[0])
        if not f or not (desde <= f < hasta) or 'PRUEBA' in str(r[10] or ''): continue
        if r[1] == 'correo' and '--con-correo' not in sys.argv: continue   # la linea base es de chats (WhatsApp, IG, Messenger)
        chats[(r[1], ident(r[11]) or ident(r[4]) or 'sin-id-' + str(f))].append(r)
    n = len(chats)
    if not n: sys.exit('Sin chats en ese rango.')
    persona = [k for k, v in chats.items() if any(r[9] == 'Sí' for r in v)]
    reservas = sum(1 for v in chats.values() if any('creada en wix' in str(r[7] or '').lower() for r in v))
    eventos = sum(1 for v in chats.values() if any(r[5] == 'evento' for r in v))
    out = []
    out.append(f'{desde.date()} → {hasta.date()}: {n} chats')
    out.append(f'  pasan a persona: {len(persona)}  = {10*len(persona)/n:.1f} de cada 10')
    out.append(f'  reservas Wix:    {reservas}  = {10*reservas/n:.1f} de cada 10')
    out.append(f'  piden evento:    {eventos}  = {10*eventos/n:.1f} de cada 10')
    if '--detalle' in sys.argv:
        conv = list(wb['Conversaciones'].iter_rows(values_only=True))[1:]
        for k in persona:
            r = next(r for r in chats[k] if r[9] == 'Sí')
            out.append(f'\n### {fecha(r[0]):%m-%d %H:%M} · {k[0]} · {k[1]} · [{r[8]}] {str(r[7] or "")[:140]}')
            for c in conv:
                fc = fecha(c[0])
                if fc and desde <= fc < hasta and c[1] == k[0] and ident(c[2]) == k[1]:
                    out.append(f'  {fc:%m-%d %H:%M} {c[5]}: {str(c[6] or "")[:300]}'.replace('\n', ' '))
    print('\n'.join(out))
    if '--salida' in sys.argv:
        dest = sys.argv[sys.argv.index('--salida') + 1]
        with open(dest, 'a') as fh:
            fh.write(f'\n\n## Medicion automatica ({dt.date.today()})\n\n```\n' + '\n'.join(out) + '\n```\n')

if __name__ == '__main__':
    main()
