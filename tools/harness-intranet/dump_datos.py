#!/usr/bin/env python3
"""Vuelca el espejo del maestro a maestro.json para harness.js (tablero global, 25-sep-2026).
Uso: python3 dump_datos.py [ruta al espejo .xlsx]  -> escribe maestro.json en esta carpeta.
Las hojas de inventario se leen aparte (inventarios.json): exportarlas del Sheet o
pegar el JSON con {cocina: {'2026-08': [[fila]...], PRODUCTOS: [...]}, barra: {...}}."""
import openpyxl, json, datetime, sys, os
ESPEJO = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser(
    '~/My Drive/Rosanta OS/03_Finance_Data_OS/Maestro/Rosanta_Maestro_ESPEJO.xlsx')
wb = openpyxl.load_workbook(ESPEJO, read_only=True, data_only=True)
def enc(v):
    if isinstance(v, datetime.datetime): return {"$d": v.strftime('%Y-%m-%dT%H:%M:%S')}
    if isinstance(v, datetime.date): return {"$d": v.strftime('%Y-%m-%dT00:00:00')}
    if isinstance(v, datetime.time): return v.strftime('%H:%M:%S')
    return v
out = {}
for name in ['02_Ventas_Maestro', '01_FEL_Maestro', '03_Banco_Industrial', '04_Banco_BAC',
             '05_Tarjeta_Credito_BAC', '02b_Ventas_2025', '00_Proveedores', '01b_FEL_Emitidas']:
    if name in wb.sheetnames:
        out[name] = [[enc(c) for c in r] for r in wb[name].iter_rows(values_only=True)]
json.dump(out, open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'maestro.json'), 'w'))
print('maestro.json', {k: len(v) for k, v in out.items()})
